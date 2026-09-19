# -*- coding: utf-8 -*-
"""
AEGIS - adds 4 curated evidence images to slides/GATE-2/Group 4_AEGIS_Gate2_Report.docx:
  - aegis_baseline_comparison.png under "State of the Art, Head-to-Head"
  - aegis_bayes_risk_threshold.png under "Why the Threshold Is Not Called 'Optimal'..."
  - aegis_eval.png under the headline metrics / per-threat numbers
  - dashboard-GATE2.jpg under the PoC live demo script

Run:  .venv/Scripts/python AEGIS/src/patch_gate2_docs_images.py
"""

import os
from docx import Document
from docx.shared import Inches, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH

HERE = os.path.dirname(os.path.abspath(__file__))
DOC = os.path.join(HERE, "..", "slides", "GATE-2", "Group 4_AEGIS_Gate2_Report.docx")
REPORTS = os.path.join(HERE, "..", "reports")

IMAGES = [
    # anchor is the NEXT heading/paragraph after where the image should land -
    # insert_image_before() places content immediately before its anchor, so
    # anchoring to "the following heading" puts the image at the END of the
    # section it illustrates, not before that section starts.
    (os.path.join(REPORTS, "GATE-2", "aegis_baseline_comparison.png"),
     "Why the Threshold Is Not Called \"Optimal\" Without Qualification",
     "Figure 2.2 - AEGIS vs. the SPC z-score baseline: ROC curves (left) and per-threat "
     "detection at matched false-positive rate (right).", 6.0),
    (os.path.join(REPORTS, "GATE-2", "aegis_bayes_risk_threshold.png"),
     "Additional Robustness Evidence",
     "Figure 2.3 - Bayes-risk-optimal threshold vs. assumed cost ratio (left) and expected-cost "
     "curves at illustrative cost ratios (right); dashed lines mark the chosen STEP-UP/BLOCK "
     "thresholds.", 6.0),
    (os.path.join(REPORTS, "aegis_eval.png"),
     "Detection rate per threat type, deliberately not uniform:",
     "Figure 2.4 - Risk-score separation between legitimate and attack windows (left) and "
     "detection rate per threat (right).", 6.0),
    (os.path.join(REPORTS, "dashboard-GATE2.jpg"),
     "Agent overview.",
     "Figure 2.5 - The live AEGIS zero-trust gate dashboard (src/aegis_dashboard.py).", 6.0),
]


def find_anchor(doc, snippet):
    for p in doc.paragraphs:
        if snippet in p.text:
            return p
    return None


def insert_image_before(anchor_p, image_path, caption, width_in):
    img_p = anchor_p.insert_paragraph_before("")
    img_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = img_p.add_run()
    run.add_picture(image_path, width=Inches(width_in))

    cap_p = anchor_p.insert_paragraph_before("")
    cap_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap_run = cap_p.add_run(caption)
    cap_run.italic = True
    cap_run.font.size = Pt(9)

    anchor_p.insert_paragraph_before("")  # spacer


def main():
    doc = Document(DOC)
    for image_path, anchor_text, caption, width_in in IMAGES:
        if not os.path.exists(image_path):
            print(f"SKIP (file not found): {image_path}")
            continue
        anchor = find_anchor(doc, anchor_text)
        if anchor is None:
            print(f"SKIP (anchor not found): {anchor_text!r}")
            continue
        insert_image_before(anchor, image_path, caption, width_in)
        print(f"Inserted {os.path.basename(image_path)} before {anchor_text!r}")

    doc.save(DOC)
    print("Saved", DOC)


if __name__ == "__main__":
    main()
