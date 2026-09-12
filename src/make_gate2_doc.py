# -*- coding: utf-8 -*-
"""Build a standalone, polished Gate 2 document from docs/GATE-2.md content.

Independent of the official 5G Academy template (see make_gate2_report.py for
that version); this is a self-contained, nicely formatted Word document
covering 2.1-2.5 with its own title page.

Run:  .venv/Scripts/python AEGIS/src/make_gate2_doc.py
"""

from pathlib import Path
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "docs" / "GATE-2.docx"
DIAGRAM = ROOT / "reports" / "aegis_architecture_diagram.png"

NAVY = RGBColor(0x0B, 0x1F, 0x3A)
ACCENT = RGBColor(0x1E, 0x5A, 0x8A)
GREY = RGBColor(0x55, 0x55, 0x55)


def set_cell_shading(cell, hex_color):
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), hex_color)
    cell._tc.get_or_add_tcPr().append(shd)


def style_table(table):
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for row_idx, row in enumerate(table.rows):
        for cell in row.cells:
            for p in cell.paragraphs:
                for run in p.runs:
                    run.font.size = Pt(9.5)
                if row_idx == 0:
                    for run in p.runs:
                        run.font.bold = True
                        run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
            if row_idx == 0:
                set_cell_shading(cell, "1E5A8A")
            elif row_idx % 2 == 0:
                set_cell_shading(cell, "F2F5F8")


def add_table(doc, headers, rows, widths=None):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    for i, h in enumerate(headers):
        table.rows[0].cells[i].text = h
    for r in rows:
        cells = table.add_row().cells
        for i, val in enumerate(r):
            cells[i].text = str(val)
    style_table(table)
    if widths:
        for row in table.rows:
            for i, w in enumerate(widths):
                row.cells[i].width = Inches(w)
    return table


def h1(doc, text, number=None):
    p = doc.add_heading(level=1)
    run = p.add_run(f"{number} {text}" if number else text)
    run.font.color.rgb = NAVY
    return p


def h2(doc, text):
    p = doc.add_heading(level=2)
    run = p.add_run(text)
    run.font.color.rgb = ACCENT
    return p


def body(doc, text, bold=False, italic=False):
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.font.size = Pt(11)
    run.font.bold = bold
    run.font.italic = italic
    return p


def bullet(doc, text):
    p = doc.add_paragraph(style="List Bullet")
    run = p.add_run(text)
    run.font.size = Pt(11)
    return p


def numbered(doc, text):
    p = doc.add_paragraph(style="List Number")
    run = p.add_run(text)
    run.font.size = Pt(11)
    return p


doc = Document()
style = doc.styles["Normal"]
style.font.name = "Calibri"
style.font.size = Pt(11)

# --- Title page ---
title = doc.add_paragraph()
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = title.add_run("AEGIS")
run.font.size = Pt(36)
run.font.bold = True
run.font.color.rgb = NAVY

sub = doc.add_paragraph()
sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = sub.add_run("GATE 2 - Technical Solution & Architecture")
run.font.size = Pt(16)
run.font.color.rgb = ACCENT

sub2 = doc.add_paragraph()
sub2.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = sub2.add_run("Zero-Trust Identity for AI Agents on the Network")
run.font.size = Pt(13)
run.font.italic = True
run.font.color.rgb = GREY

doc.add_paragraph()
meta = doc.add_paragraph()
meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = meta.add_run(
    "5G Academy 2026, Fastweb + Vodafone, Topic 2, Security\n"
    "Team 4, Fauzan Ejaz (Captain), Adithya Zacharia Valavi, "
    "Ghazanfar Anees Siddiqui, Llagami Tedi"
)
run.font.size = Pt(10)
run.font.color.rgb = GREY

doc.add_page_break()

# ============================================================ 2.1
h1(doc, "Proposed Technical Solution", "2.1")
body(doc,
     "As AI agents, orchestration bots, and closed-loop automations increasingly issue "
     "commands directly to 5G network control services, classical authentication is no "
     "longer sufficient to establish trust. A stolen token, a leaked client certificate, "
     "or a compromised automation account will still pass authentication, the credential "
     "is valid, even though the actor behind it is not the legitimate agent. Classic AAA "
     "answers “is the credential valid?” It cannot answer “is this really "
     "the agent we onboarded, behaving the way it always behaves?”")
body(doc,
     "AEGIS answers that second question. It is a zero-trust gate for machine identities: "
     "it builds a behavioural fingerprint of each legitimate agent from how it actually "
     "calls the network's APIs, and continuously scores every new request against that "
     "fingerprint, aligned with NIST SP 800-207 zero-trust principles (never trust, always "
     "verify; verify continuously; assume breach).")
p = body(doc, "One-line pitch: ", bold=True)
p.add_run("credentials prove what you have; AEGIS verifies how you behave.").italic = True
body(doc,
     "Scope. AEGIS targets service/automation agents (orchestration, closed-loop "
     "automation, OSS/BSS jobs, API clients) calling network control APIs, modelled on "
     "the 5G Service-Based Architecture (SBA) / Service-Based Interface (SBI): AMF, SMF, "
     "NRF, NEF, PCF, UDM. End-user devices, the data plane, and general “AI "
     "safety” are explicitly out of scope: AEGIS guards the control-plane API "
     "surface used by machine actors, which keeps the problem concrete and solvable "
     "within the PoC timeline rather than open-ended.")
h2(doc, "The solution in one pipeline")
body(doc,
     "OBSERVE request patterns and metadata (id, timing, call sequence, payload shape) "
     "-> FINGERPRINT a known-good behavioural baseline per agent -> SCORE with rules "
     "plus an ML anomaly/spoof score -> DECIDE PASS / STEP-UP / BLOCK (zero-trust gate "
     "dashboard).")
h2(doc, "Threats the solution is built to catch")
body(doc,
     "We assume an attacker who has already obtained a valid credential (the realistic "
     "case) or controls a legitimate agent, and must therefore be caught by behaviour, "
     "not by credential checks:")
threat_rows = [
    ("T1", "Identity spoofing / impersonation", "Fingerprint mismatch: wrong endpoint mix, timing, sequence for that identity"),
    ("T2", "Compromised agent", "Drift from its own baseline: new endpoints, new methods, elevated error rate"),
    ("T3", "Reconnaissance / enumeration", "High cardinality of endpoints/targets, many 403/404s, breadth over depth"),
    ("T4", "Volumetric abuse / DoS", "Request-rate spike far above the agent's normal envelope"),
    ("T5", "Low-and-slow exfiltration", "Abnormal read/GET ratio and payload sizes, off-hours persistence"),
    ("T6", "Replay / sequence anomaly", "Broken call-sequence pattern vs the agent's normal call graph"),
    ("T7", "Privilege / scope creep", "Calls to NFs never in the agent's baseline scope"),
]
add_table(doc, ["#", "Threat", "Behavioural signal AEGIS keys on"], threat_rows, widths=[0.4, 1.6, 4.3])
doc.add_paragraph()
body(doc,
     "This directly extends the Gate 1 feasibility case: the technology (behavioural "
     "fingerprinting and anomaly detection), the risks, and the economic rationale "
     "(enabling safe automation, NIS2/GDPR/EU-AI-Act audit evidence) were validated at "
     "Gate 1. Gate 2 delivers the concrete architecture (2.2, with a full system "
     "diagram), a working implementation (2.3), and numerical proof that the approach "
     "works (2.4), the three things Gate 1 could only describe as intent.")

# ============================================================ 2.2
doc.add_page_break()
h1(doc, "Architecture", "2.2")
body(doc,
     "AEGIS is structured as a four-stage pipeline sitting logically in front of the "
     "network's control-plane APIs. Every request an agent sends to a Network Function "
     "is observed by the gate before (in production) or in parallel with (in the PoC) "
     "reaching the NF.")
if DIAGRAM.exists():
    pic_p = doc.add_paragraph()
    pic_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    pic_p.add_run().add_picture(str(DIAGRAM), width=Inches(6.3))
    cap = doc.add_paragraph()
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = cap.add_run(
        "Figure 2.1. AEGIS system architecture: OBSERVE -> FINGERPRINT -> SCORE -> "
        "DECIDE, with the rules layer, ML layer, and known-good profile store feeding "
        "SCORE, and PASS / STEP-UP / BLOCK as the three decision outcomes."
    )
    run.font.size = Pt(9)
    run.font.italic = True
    run.font.color.rgb = GREY
h2(doc, "Component description")
comp_items = [
    ("OBSERVE.", "Every request is parsed into a normalized event: agent identity, HTTP "
     "method, target Network Function, endpoint, timestamp, and payload size. Events are "
     "grouped into a per-agent stream, the unit the rest of the pipeline operates on."),
    ("FINGERPRINT.", "Requests are aggregated into sliding 60-second windows and turned "
     "into a feature vector per agent, per window: volume/rate, surface (distinct "
     "NFs/endpoints/target IDs), method mix, sequence, errors, payload, temporal "
     "(hour-of-day), and cross-window breadth (a rolling union of endpoints/NFs and "
     "error count over the trailing 5 minutes, so reconnaissance spread thin enough to "
     "look quiet in any single window is still caught). These vectors are compared "
     "against a known-good profile store, the behavioural baseline built for each agent "
     "during onboarding."),
    ("SCORE.", "Two layers run in parallel and are fused. The rules layer (fast, "
     "explainable) applies hard caps and forbidden transitions, a session-state machine "
     "that tracks PDU-session lifecycle (catching orphaned update/release calls, T6), "
     "and rolling 5-minute accumulators (breadth, error count, peak payload) that catch "
     "cumulative anomalies too sparse for any single window, the mechanism behind T3 and "
     "T5 detection. The ML layer is an Isolation Forest trained only on legitimate "
     "traffic, engaging once an agent has at least 10 requests in the trailing 5 "
     "minutes. Fusion: risk equals the maximum of the rule score and the ML score. "
     "Thresholds are tuned for a realistic false-positive budget (about 1.5 percent) "
     "rather than maximum sensitivity, which is why detection is intentionally not "
     "uniform across threats (see 2.4)."),
    ("DECIDE.", "The fused risk score maps to PASS (low risk, behaves like its "
     "known-good self), STEP-UP (medium risk: re-authentication, approval, or "
     "throttling), or BLOCK (high risk: quarantine the identity, alert the SOC)."),
    ("Trust boundary.", "Every request from an agent to a Network Function crosses the "
     "AEGIS gate. In production this is envisioned as an inline policy check at the edge "
     "(FASTedge); for the PoC it runs as an offline/log-replay analysis over synthetic "
     "SBI-style traffic."),
    ("Dashboard / enforcement surface.", "A Streamlit application exposes, per agent, "
     "current risk, PASS/STEP-UP/BLOCK status, a breakdown of detections by threat type, "
     "and an incident inspector with a plain-language explanation of why a request was "
     "flagged."),
]
for label, text in comp_items:
    p = bullet(doc, "")
    r1 = p.add_run(label + " ")
    r1.bold = True
    r1.font.size = Pt(11)
    r2 = p.add_run(text)
    r2.font.size = Pt(11)

# ============================================================ 2.3
doc.add_page_break()
h1(doc, "Implementation Strategy", "2.3")
p = body(doc, "Methodology. ", bold=True)
p.add_run(
    "AEGIS is built bottom-up, PoC-first, and evaluation-driven: the ML/security "
    "equivalent of test-driven development. No rule or model change is kept unless it "
    "is immediately re-scored against the full labelled dataset and shown to improve "
    "(or at least not regress) the numbers in 2.4. Each of the four architectural "
    "stages is a separately runnable, separately verifiable script, so the pipeline can "
    "be inspected and debugged stage by stage rather than as one opaque system."
)
h2(doc, "Tools and platforms")
tool_rows = [
    ("Language/runtime", "Python 3.12", "Single language across generation, detection, and dashboard"),
    ("Data generation & processing", "pandas, NumPy", "Synthetic SBI-traffic generation, windowing, feature engineering"),
    ("Anomaly detection", "scikit-learn (Isolation Forest)", "Unsupervised ML layer, trained on legitimate traffic only"),
    ("Visualization / gate UI", "Streamlit, Matplotlib", "Live zero-trust gate dashboard and evaluation charts"),
    ("Deck/report tooling", "python-docx, python-pptx", "Reproducible generation of the Gate decks and reports"),
    ("Delivery", "Git/GitHub, GitHub Pages", "Version control and a public showcase site"),
]
add_table(doc, ["Layer", "Tool / platform", "Role"], tool_rows, widths=[1.6, 1.9, 2.8])
doc.add_paragraph()
h2(doc, "Implementation phases")
phase_rows = [
    ("1. Synthetic environment", "Labelled SBI-style traffic, 5 agents, T1-T7", "generate_synthetic_traffic.py"),
    ("2. Fingerprinting", "Per-agent windowed feature vectors", "aegis_detect.py (build_windows)"),
    ("3. Rules layer", "Deterministic scope/session/rolling checks", "aegis_detect.py (rule_score)"),
    ("4. ML layer", "Isolation Forest fit on legit-only traffic", "aegis_detect.py (ml_scores)"),
    ("5. Fusion & decision", "Risk scoring, PASS/STEP-UP/BLOCK", "aegis_detect.py (main, decision)"),
    ("6. Dashboard", "Live zero-trust gate UI", "aegis_dashboard.py"),
    ("7. Evaluation & iteration", "Metrics, charts, bug-fixing loop", "aegis_baseline_metrics.md"),
]
add_table(doc, ["Phase", "Output", "Primary file(s)"], phase_rows, widths=[1.7, 2.4, 2.2])
doc.add_paragraph()
h2(doc, "Step-by-step, and what the evaluation loop caught")
steps = [
    "Synthetic traffic generation produced 210,000 requests across 5 legitimate agent "
    "profiles with realistic PDU-session lifecycles, plus the seven threat scenarios "
    "injected as labelled anomalous segments, removing the dependency on real operator "
    "logs.",
    "Feature engineering turned the windowed request stream into the behavioural "
    "feature vectors described in 2.2.",
    "The rules layer implements deterministic scope, rate, and session-sequence "
    "checks. An earlier T6 window-count heuristic was replaced with a session-ID state "
    "machine after it was found to false-fire on legitimate bursty traffic.",
    "The ML layer, an Isolation Forest, is fit only on legitimate traffic with a "
    "held-out calibration split.",
    "Fusion and decision policy combine rule and ML scores into PASS / STEP-UP / "
    "BLOCK.",
    "The dashboard exposes per-agent risk, decisions, per-threat breakdown, and the "
    "incident inspector.",
    "Evaluation and iteration: re-scoring after every change surfaced two real bugs, "
    "not just tuning cosmetics. First, T3 and T5 are naturally sparse per window; "
    "rolling 5-minute accumulators fixed this generically. Second, the initial T2 and "
    "T5 attack scenarios both had the injected agent call a Network Function outside "
    "its own onboarded scope, indistinguishable from T7 to the rules layer, so both "
    "were trivially caught by the scope rule rather than the drift/payload logic they "
    "were meant to test. Redesigning both to stay within the attacking agent's own "
    "scope forced detection to rely on the intended statistical signal, which is the "
    "origin of the non-uniform per-threat numbers in 2.4.",
]
for s in steps:
    numbered(doc, s)

# ============================================================ 2.4
doc.add_page_break()
h1(doc, "Numerical Performance Analysis", "2.4")
p = body(doc, "Benchmarking against real-world detection systems, not just internal targets. ", bold=True)
p.add_run(
    "Before finalizing the operating point, we researched how comparable systems "
    "perform in practice, since a PoC that reports 100 percent detection on everything "
    "is a red flag to anyone who has looked at production security data, not a "
    "strength."
)
bullet(doc,
       "Published NIDS studies on CICIDS2017/UNSW-NB15 consistently show volumetric/DoS "
       "attacks detected at 97 to 99.8 percent (large, unambiguous signal), while "
       "reconnaissance and other low-frequency attack classes detect far lower unless "
       "heavy class-balancing is applied.")
bullet(doc,
       "Industry SOC data (Microsoft/Omdia State of the SOC 2025, SANS 2025 Detection "
       "and Response Survey) puts typical false-positive rates at 46 to 83 percent, "
       "with elite, well-tuned SOCs under 10 percent, and under 5 percent considered "
       "excellent.")
bullet(doc,
       "UEBA/insider-threat research shows a hard precision/recall tradeoff: models "
       "reporting near-perfect recall do so by sacrificing precision, sometimes to as "
       "low as 0.54; low-and-slow exfiltration is consistently the hardest category, "
       "consistent with the roughly 77-day average real-world dwell time cited for "
       "insider threats.")
body(doc,
     "We therefore tuned AEGIS's decision thresholds and rule sensitivities for a "
     "realistic false-positive budget (about 1.5 percent), in the "
     "well-tuned-not-superhuman range, rather than for maximum sensitivity, and "
     "treated a non-uniform, threat-dependent detection rate as the expected, correct "
     "outcome rather than something to eliminate.")
body(doc,
     "Test set: approximately 11,365 held-out windows (11,010 legitimate, 355 attack), "
     "60-second window size, decision thresholds STEP-UP at 0.6 or above and BLOCK at "
     "0.85 or above.")
h2(doc, "Headline metrics (rules + ML fused)")
metric_rows = [
    ("ROC-AUC", "0.965"),
    ("Precision", "0.672"),
    ("Recall (detection rate)", "0.907"),
    ("F1 score", "0.772"),
    ("False-positive rate (legit traffic flagged)", "1.43%"),
    ("ROC-AUC, ML only (no rules)", "0.721"),
]
add_table(doc, ["Metric", "Value"], metric_rows, widths=[3.5, 1.5])
doc.add_paragraph()
body(doc,
     "These sit comfortably inside the range reported in the literature above: an FPR "
     "under 2 percent is elite-SOC territory, not the near-zero rate that would raise "
     "suspicion of a leaky evaluation, and a precision of 0.67 is in line with, in fact "
     "better than, several published UEBA results that trade precision for recall.")
h2(doc, "Detection rate per threat type, deliberately not uniform")
detect_rows = [
    ("T4 - Volumetric abuse / DoS", "60", "100%", "Loud rate spike, the easiest signal in the literature"),
    ("T7 - Privilege / scope creep", "35", "100%", "Deterministic policy/ACL check, not statistical inference"),
    ("T1 - Identity spoofing / impersonation", "54", "96%", "Mostly a scope mismatch, with a stealthier in-scope variant mixed in"),
    ("T6 - Replay / sequence anomaly", "57", "93%", "An isolated single anomaly without context is treated as ambiguous"),
    ("T2 - Compromised agent", "43", "91%", "Drift confined to the agent's own scope, relies on genuine behavioural drift"),
    ("T3 - Reconnaissance / enumeration", "60", "80%", "Enumeration within an agent's own scope, hardest class in NIDS benchmarks"),
    ("T5 - Low-and-slow exfiltration", "46", "76%", "Hardest threat category industry-wide, caught only via rolling accumulation"),
]
add_table(doc, ["Threat", "Windows", "Detected", "Why this level"], detect_rows, widths=[1.9, 0.7, 0.8, 2.6])
doc.add_paragraph()
body(doc,
     "Two of these threats (T4, T7) are legitimately near 100 percent because they are "
     "not statistical detection problems, a rate cap and a scope/ACL check are "
     "deterministic lookups, and real production systems are expected to catch these "
     "essentially every time. The other five sit in a genuine 76 to 96 percent spread "
     "that tracks the strength of their underlying behavioural signal: loud beats "
     "subtle, and diffuse/low-and-slow attacks are always the hardest.")
h2(doc, "Gate decisions on the test set")
decision_rows = [
    ("PASS", "10,886", "95.8%"),
    ("STEP-UP", "148", "1.3%"),
    ("BLOCK", "331", "2.9%"),
]
add_table(doc, ["Decision", "Windows", "Share"], decision_rows, widths=[2.0, 1.5, 1.5])
doc.add_paragraph()
h2(doc, "Interpretation")
bullet(doc,
       "A 1.43 percent false-positive rate sits inside the well-tuned-SOC band from the "
       "industry data above (elite under 10 percent, excellent under 5 percent) without "
       "claiming an implausible near-zero rate.")
bullet(doc,
       "90.7 percent recall means roughly 9 in 10 attacks are caught before reaching "
       "PASS, strong but honestly short of perfect, consistent with recall figures "
       "reported for real behavioural/UEBA systems at comparable false-positive "
       "budgets.")
bullet(doc,
       "The precision of 0.67 reflects the same recall/precision tension documented "
       "across the UEBA literature: catching subtle threats (T3, T5) at a low "
       "false-positive rate necessarily means the gate occasionally flags legitimate "
       "but unusual behaviour, which the STEP-UP tier absorbs without disrupting the "
       "agent outright.")
bullet(doc,
       "These results still substantially exceed the two baselines identified in the "
       "threat model, a naive rate-threshold rule and a random gate.")

# ============================================================ 2.5
doc.add_page_break()
h1(doc, "Expected Proof of Concept (PoC)", "2.5")
body(doc,
     "Unlike a purely forward-looking PoC description, AEGIS already has a working, "
     "demonstrable baseline implementing the full pipeline end to end on synthetic "
     "data, matching the architecture in Figure 2.1 exactly, not a simplified stand-in "
     "for it.")
h2(doc, "Live demo script")
demo_steps = [
    ("Agent overview.", "Open the dashboard and show all 5 legitimate agents with "
     "their live risk scores, all sitting in PASS, matching their known-good "
     "baseline."),
    ("Trigger a threat.", "Select a scored window for one of T1-T7 (for example, T4 "
     "volumetric) and show the dashboard flip that agent to BLOCK in real time."),
    ("Open the incident inspector.", "Show the plain-language why: which feature "
     "tripped (rate spike, scope violation, rolling breadth), tracing directly back "
     "to the rules layer or ML layer that produced it."),
    ("Repeat across threat types.", "Cycle through a spoofing case (T1), a quiet "
     "within-scope drift case (T2), and a low-and-slow exfiltration case (T5), "
     "showing the full spread of detection difficulty from 2.4 live, not just the "
     "easy cases."),
    ("Show the numbers.", "Bring up the metrics report and evaluation chart alongside "
     "the dashboard: ROC-AUC, recall, false-positive rate, and the per-threat table, "
     "so the qualitative demo and the quantitative evidence are shown side by side."),
    ("Make the zero-trust gap concrete.", "Point to a flagged window where the "
     "request used a fully valid credential (authentication succeeded) but was still "
     "blocked on behaviour, the exact gap in classical AAA that AEGIS closes."),
]
for i, (label, text) in enumerate(demo_steps, 1):
    p = doc.add_paragraph(style="List Number")
    r1 = p.add_run(label + " ")
    r1.bold = True
    r1.font.size = Pt(11)
    r2 = p.add_run(text)
    r2.font.size = Pt(11)

h2(doc, "PoC readiness at Gate 2")
ready_rows = [
    ("Synthetic SBI-style traffic generator (5 agents, T1-T7)", "Done"),
    ("Fingerprinting, rules layer, ML layer, fusion, decision policy", "Done"),
    ("Live dashboard with incident inspector", "Done"),
    ("Numerical evaluation (ROC-AUC, recall, precision, FPR, per-threat)", "Done"),
    ("ROC operating-point / threshold sensitivity analysis", "Done"),
    ("Real M2M / API-gateway traffic", "Not yet, synthetic only, blocked on Fastweb/Vodafone data access"),
]
add_table(doc, ["Capability", "Status"], ready_rows, widths=[4.5, 1.8])
doc.add_paragraph()
body(doc,
     "What this validates: that a behavioural-fingerprint plus rules/ML gate is a "
     "practical way to add a continuous-verification layer on top of existing "
     "credential-based authentication for machine identities on 5G control-plane "
     "APIs, trainable without labelled attacks, and numerically effective at a "
     "realistic operating point, without requiring changes to the underlying Network "
     "Functions themselves.")
h2(doc, "Path to Gate 3")
bullet(doc,
       "Closing the gap on the hardest threats (T3 recon at 80 percent, T5 "
       "exfiltration at 76 percent) with richer features, for example per-target-ID "
       "cardinality, not just per-endpoint.")
bullet(doc,
       "Integrating real M2M / API-gateway logs if made available by "
       "Fastweb/Vodafone, the one open item outside our control.")
bullet(doc, "Recording the final video demonstration from the live demo script above.")
bullet(doc,
       "Expanding the number of agent and attack variants to stress-test "
       "generalization beyond the current 5-agent / 7-threat catalogue.")

doc.add_paragraph()
footer = doc.add_paragraph()
footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = footer.add_run("AEGIS, Zero-Trust Identity for AI Agents, Team 4, 5G Academy 2026")
run.font.size = Pt(9)
run.font.color.rgb = GREY

OUT.parent.mkdir(exist_ok=True)
doc.save(OUT)
print(f"Saved to {OUT}")
