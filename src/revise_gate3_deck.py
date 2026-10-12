# -*- coding: utf-8 -*-
"""Revision pass on the CURRENT AEGIS_Gate3_Official.pptx (as it sits on disk
right now). Four requested changes, nothing else touched:

  1. Strip "Gate 3" / "Gate 2" framing throughout - present capabilities
     standalone, not as a progress report relative to earlier gates.
  2. Move the "Beyond Synthetic: A Real 5G Core" (Open5GS) slide to right
     after the Isolation Forest slides (was in the PoC-readiness section).
  3. Reconcile the "seven threats" vs "10 attack variants" framing on the
     threats-table slide.
  4. Swap in four regenerated charts (SPC head-to-head, risk-score
     histogram, ROC operating-point, CV stability) that fix real legend/
     label-overlap rendering bugs, and refresh their two captions'
     numbers to match (the underlying scripts were re-run against the
     current, widened 10-variant catalogue - see git history for the
     before/after numbers).

Run:  .venv/Scripts/python AEGIS/src/revise_gate3_deck.py
In/Out: AEGIS/slides/GATE-3/AEGIS_Gate3_Official.pptx  (edited in place)
"""

import copy
import os
from pptx import Presentation

HERE = os.path.dirname(os.path.abspath(__file__))
DECK = os.path.join(HERE, "..", "slides", "GATE-3", "AEGIS_Gate3_Official.pptx")
REPORTS = os.path.join(HERE, "..", "reports")

SPC_PNG = os.path.join(REPORTS, "GATE-2", "aegis_baseline_comparison.png")
EVAL_PNG = os.path.join(REPORTS, "aegis_eval.png")
ROC_PNG = os.path.join(REPORTS, "GATE-2", "aegis_bayes_risk_threshold.png")
CV_PNG = os.path.join(REPORTS, "GATE-2", "aegis_cv_threshold_stability.png")


# ------------------------------------------------------------------ helpers
def find_shape_by_text(slide, snippet):
    for shp in slide.shapes:
        if shp.has_text_frame and snippet in shp.text_frame.text:
            return shp
    return None


def find_slide_by_text(prs, snippet):
    for s in prs.slides:
        if find_shape_by_text(s, snippet) is not None:
            return s
    raise ValueError(f"no slide contains {snippet!r}")


def set_simple_text(shape, new_text):
    tf = shape.text_frame
    p0 = tf.paragraphs[0]
    if not p0.runs:
        p0.add_run()
    r0 = p0.runs[0]
    r0.text = new_text
    for r in list(p0.runs[1:]):
        r._r.getparent().remove(r._r)
    for p in list(tf.paragraphs[1:]):
        p._p.getparent().remove(p._p)


def set_bullets(shape, bullets):
    """Rewrites a '-  <prefix run> <body run>' bulleted box, reusing the
    first paragraph's formatting as the template for every new paragraph."""
    tf = shape.text_frame
    template_p = tf.paragraphs[0]._p
    parent = template_p.getparent()
    new_ps = []
    for text in bullets:
        p_copy = copy.deepcopy(template_p)
        runs_xml = p_copy.findall(".//{*}r")
        for extra in runs_xml[2:]:
            extra.getparent().remove(extra)
        prefix_r, body_r = runs_xml[0], runs_xml[1]
        body_r.find("{*}t").text = text
        new_ps.append(p_copy)
    for p in list(tf.paragraphs):
        p._p.getparent().remove(p._p)
    for p in new_ps:
        parent.append(p)


def get_picture(slide):
    for shp in slide.shapes:
        if shp.shape_type == 13:
            return shp
    return None


def replace_picture(slide, image_path):
    old = get_picture(slide)
    left, top, width, height = old.left, old.top, old.width, old.height
    old._element.getparent().remove(old._element)
    return slide.shapes.add_picture(image_path, left, top, width, height)


def slide_index(prs, slide):
    for i, s in enumerate(prs.slides):
        if s is slide:
            return i
    raise ValueError("slide not found")


def move_slide(prs, old_index, new_index):
    xml_slides = prs.slides._sldIdLst
    slides = list(xml_slides)
    el = slides[old_index]
    xml_slides.remove(el)
    xml_slides.insert(new_index, el)


prs = Presentation(DECK)

# ============================================================ 1. de-gate the framing
# Cover/index slide: "AEGIS / Gate 3" subtitle -> capability subtitle
slide2 = find_slide_by_text(prs, "Gate 3")
gate3_shape = find_shape_by_text(slide2, "Gate 3")
for p in gate3_shape.text_frame.paragraphs:
    for r in p.runs:
        if r.text.strip() == "Gate 3":
            r.text = "Zero-Trust Identity for AI Agents"

# Section-kicker overline on slides 4/5/6: drop the "GATE 3: " prefix
for snippet in ["Proposed Technical Solution", "The Seven Threats We Watch For",
                "The Gap, Our Solution, and the State of the Art"]:
    s = find_slide_by_text(prs, snippet)
    kicker = find_shape_by_text(s, "GATE 3: TECHNICAL SOLUTION")
    if kicker is not None:
        for p in kicker.text_frame.paragraphs:
            for r in p.runs:
                if "GATE 3: " in r.text:
                    r.text = r.text.replace("GATE 3: ", "")

# "PoC Readiness at Gate 3" -> "PoC Readiness"
poc_slide = find_slide_by_text(prs, "PoC Readiness at Gate 3")
set_simple_text(find_shape_by_text(poc_slide, "PoC Readiness at Gate 3"), "PoC Readiness")

# Open5GS slide's gate-relative sentence -> standalone framing
open5gs_slide = find_slide_by_text(prs, "Beyond Synthetic: A Real 5G Core")
gate_sentence = find_shape_by_text(
    open5gs_slide, "Gate 2's roadmap item")
set_simple_text(
    gate_sentence,
    "A core validation question for any synthetic-first approach: does detection "
    "still hold up against real M2M / API-gateway traffic, not just the generator?")

print("1. De-gated cover, section kickers, PoC-readiness title, Open5GS sentence.")

# ============================================================ 2. move Open5GS slide after Isolation Forest
iforest_slide = find_slide_by_text(prs, "Training, Features & Calibration")
target_pos = slide_index(prs, iforest_slide) + 1
move_slide(prs, slide_index(prs, open5gs_slide), target_pos)

# its own kicker now belongs to the Implementation Strategy section, not PoC
kicker = find_shape_by_text(open5gs_slide, "PART 5 OF 6 - EXPECTED PROOF OF CONCEPT")
set_simple_text(kicker, "PART 3 OF 6 - IMPLEMENTATION STRATEGY")

print(f"2. Moved Open5GS slide to position {target_pos + 1} (after Isolation Forest).")

# ============================================================ 3. reconcile 7 threats vs 10 variants
threats_slide = find_slide_by_text(prs, "The Seven Threats We Watch For")
subtitle = find_shape_by_text(threats_slide, "Same four-stage pipeline.")
set_simple_text(
    subtitle,
    "Same four-stage pipeline. Seven different ways an identity can go wrong - "
    "three (T1, T4, T7) also have a harder v2/slow-ramp variant, 10 attack forms in total.")

print("3. Added 7-threats / 10-variants reconciliation on the threats slide.")

# ============================================================ 4a. swap in the fixed charts
replace_picture(find_slide_by_text(prs, "State of the Art, Head-to-Head"), SPC_PNG)
replace_picture(find_slide_by_text(prs, "Detection Rate Per Threat"), EVAL_PNG)
replace_picture(find_slide_by_text(prs, "ROC Operating-Point Analysis"), ROC_PNG)
replace_picture(find_slide_by_text(prs, "Appendix: Cross-Validated Threshold Stability"), CV_PNG)

print("4a. Swapped in 4 regenerated chart images.")

# ============================================================ 4b. refresh the two captions whose numbers changed
spc_slide = find_slide_by_text(prs, "State of the Art, Head-to-Head")
spc_caption = find_shape_by_text(spc_slide, "SPC, the classic control-chart baseline")
set_simple_text(
    spc_caption,
    "SPC, the classic control-chart baseline, run on the same split as AEGIS: strong on "
    "loud signals (T4 100%, T1 100%), blind to subtler ones (T2 34%, T3 45%, T6 31%) - "
    "the measured case for behavioural fingerprinting + fusion.")

cv_slide = find_slide_by_text(prs, "Appendix: Cross-Validated Threshold Stability")
cv_caption = find_shape_by_text(cv_slide, "The 0.60 threshold was picked")
set_simple_text(
    cv_caption,
    "The 0.60 threshold was picked on the same test set final numbers are reported on - "
    "a fair concern. 5-fold CV (threshold picked on 4 folds, evaluated on the 5th): "
    "threshold 0.604 +/- 0.005, recall 91.6% +/- 3.3%, FPR 1.43% +/- 0.16% - not "
    "cherry-picked to one slice of data.")

print("4b. Refreshed SPC and CV-stability caption numbers to match.")

prs.save(DECK)
print("Saved", DECK)
