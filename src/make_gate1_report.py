# -*- coding: utf-8 -*-
"""
Fill the official 5G Academy Gate Evaluation Report template for AEGIS — Gate 1.
Completes the cover page + Gate 1 sections (1.1–1.4), embeds the Gantt chart, and
leaves Gate 2 / Gate 3 as template placeholders (per the template instructions).

Run:  .venv/Scripts/python AEGIS/src/make_gate1_report.py
In :  AEGIS/slides/5G_Academy_Gate_Report_Template.docx
Out:  AEGIS/slides/AEGIS_Gate1_Report.docx
"""

import os
from docx import Document
from docx.oxml import OxmlElement
from docx.text.paragraph import Paragraph
from docx.shared import Inches, Pt, RGBColor

HERE = os.path.dirname(os.path.abspath(__file__))
TPL = os.path.join(HERE, "..", "slides", "5G_Academy_Gate_Report_Template.docx")
OUT = os.path.join(HERE, "..", "slides", "AEGIS_Gate1_Report.docx")
GANTT = os.path.join(HERE, "..", "reports", "aegis_gantt.png")

doc = Document(TPL)


# ------------------------------------------------------------------ helpers
def heading_index(text):
    for i, p in enumerate(doc.paragraphs):
        if p.text.strip() == text:
            return i
    return -1


def placeholder_after(idx, marker="["):
    for p in doc.paragraphs[idx + 1:]:
        if p.text.strip().startswith(marker):
            return p
    return None


def set_runs(p, parts):
    for r in list(p.runs):
        r._element.getparent().remove(r._element)
    for (t, b) in parts:
        run = p.add_run(t); run.bold = b
    return p


def style_or(name, fallback="List Paragraph"):
    try:
        _ = doc.styles[name]; return name
    except KeyError:
        return fallback


def add_after(anchor, parts, style=None):
    """Insert a new paragraph right after `anchor` (a Paragraph); return it."""
    new_p = OxmlElement('w:p')
    anchor._p.addnext(new_p)
    np = Paragraph(new_p, anchor._parent)
    if style:
        np.style = doc.styles[style]
    for (t, b) in parts:
        run = np.add_run(t); run.bold = b
    return np


def bullets_after(anchor, items):
    st = style_or("List Bullet")
    cur = anchor
    for parts in items:
        cur = add_after(cur, parts, style=st)
    return cur


# ------------------------------------------------------------------ cover page
tbl = doc.tables[0]
vals = ["4",
        "Fauzan Ejaz",
        "Fauzan Ejaz (Leader), Adithya Zacharia Valavi, Ghazanfar Anees Siddiqui, Llagami Tedi",
        "AEGIS — Zero-Trust Identity for AI Agents on the Network",
        "☑ Gate 1    ☐ Gate 2    ☐ Gate 3",
        "7 September 2026"]
for row, v in zip(tbl.rows, vals):
    cell = row.cells[1]
    cell.text = ""
    cell.paragraphs[0].add_run(v)

# ================================================================== 1.1 Technologies
ph = placeholder_after(heading_index("1.1 Technologies"))
set_runs(ph, [("AEGIS applies a Zero-Trust security model to the identities of the AI agents and "
               "automations that operate the 5G control plane. It combines a telecom control-plane "
               "surface, a zero-trust security paradigm, and behavioural machine learning:", False)])
bullets_after(ph, [
    [("5G control-plane surface — ", True),
     ("the 5G Service-Based Architecture (SBA) and Service-Based Interface (SBI): Network Functions "
      "(AMF, SMF, NRF, PCF, UDM, NEF) exposing HTTP/2 REST APIs — the same interfaces that "
      "orchestration, OSS and closed-loop-automation agents call.", False)],
    [("Zero-Trust security (NIST SP 800-207) — ", True),
     ("continuous, per-request verification of machine identities (“never trust, always verify”), "
      "applied to behaviour rather than to credentials alone.", False)],
    [("Behavioural fingerprinting & anomaly detection — ", True),
     ("unsupervised machine learning (Isolation Forest) trained only on legitimate traffic, combined "
      "with deterministic policy rules and a session-level state machine that validates PDU-session "
      "call sequences.", False)],
    [("Detection features — ", True),
     ("request rate and inter-arrival timing, distinct Network Functions / endpoints, method mix, "
      "call-sequence integrity, error rate, payload size, and time-of-day.", False)],
    [("Data & tooling — ", True),
     ("a synthetic 5G SBA traffic generator (Python) produces labelled data with injected attacks; "
      "scikit-learn and pandas for modelling; a Streamlit dashboard implements the zero-trust gate "
      "(PASS / STEP-UP / BLOCK). Real M2M / API-gateway logs can be ingested if provided.", False)],
    [("Deployment vision — ", True),
     ("in production the gate runs as an inline policy check at the network edge (FASTedge); the PoC "
      "runs offline on replayed logs.", False)],
    [("Standards & frameworks — ", True),
     ("NIST SP 800-207, 3GPP SBA, TM Forum closed-loop automation (OODA); aligned with NIS2, GDPR and "
      "the EU AI Act.", False)],
])

# ================================================================== 1.2 Project Schedule
ph_img = placeholder_after(heading_index("1.2 Project Schedule (Gantt Chart)"))
for r in list(ph_img.runs):
    r._element.getparent().remove(r._element)
if os.path.exists(GANTT):
    ph_img.add_run().add_picture(GANTT, width=Inches(6.3))

# phase deadline lines
for p in doc.paragraphs:
    tx = p.text.strip()
    if tx.startswith("Phase 1:"):
        set_runs(p, [("Phase 1 — Foundation & baseline PoC ", True),
                     ("(threat model, synthetic 5G SBA data, rules + ML detector)  —  Deadline: 10 Aug 2026 "
                      "(completed).", False)])
    elif tx.startswith("Phase 2:"):
        set_runs(p, [("Phase 2 — Technical solution & evaluation ", True),
                     ("(feature hardening, numerical performance analysis, gate dashboard, optional real-data "
                      "integration)  —  Deadline: Gate 2, 22 Sep 2026.", False)])
    elif tx.startswith("Phase 3:"):
        set_runs(p, [("Phase 3 — Final Proof of Concept & video ", True),
                     ("(integration, live zero-trust gate demo, storyboard & shooting)  —  Deadline: Gate 3, "
                      "15 Oct 2026.", False)])

# ================================================================== 1.3 Feasibility
ph = placeholder_after(heading_index("1.3 Feasibility Analysis"))
set_runs(ph, [("The solution is technically feasible and already de-risked by a working baseline "
               "prototype:", False)])
last = bullets_after(ph, [
    [("Resources — ", True),
     ("entirely open-source (Python, scikit-learn, Streamlit); no specialised hardware, it runs on a "
      "standard laptop. Production adds only lightweight edge compute at FASTedge.", False)],
    [("Data availability — ", True),
     ("the main external risk (access to real network traffic) is fully mitigated: the PoC is "
      "synthetic-first. A 3GPP-grounded generator produces realistic, labelled SBA traffic; real logs "
      "are an upgrade, not a dependency.", False)],
    [("Evidence of feasibility — ", True),
     ("the working baseline already detects all seven modelled threat classes on 11,343 held-out "
      "request windows, at ROC-AUC 0.99, 98% recall and a 0.8% false-positive rate.", False)],
])
# risk/mitigation table inserted after the feasibility bullets
risk_intro = add_after(last, [("Key risks and mitigations:", True)])
risks = [
    ("Risk", "Mitigation"),
    ("Synthetic data may not capture real agent behaviour",
     "Modelled on 3GPP SBA semantics; the pipeline ingests real logs when available."),
    ("False positives disrupt operations",
     "Held-out calibration + a soft STEP-UP action before BLOCK; thresholds are tunable."),
    ("Real-data access is not granted",
     "The full PoC is deliverable on synthetic data alone."),
    ("Adaptive / evasive attackers",
     "Layered rules + ML + session state machine; roadmap for periodic retraining."),
]
rtbl = doc.add_table(rows=0, cols=2)
try:
    rtbl.style = "Table Grid"
except KeyError:
    pass
for a, b in risks:
    cells = rtbl.add_row().cells
    cells[0].text = a; cells[1].text = b
# bold the header row
for c in rtbl.rows[0].cells:
    for r in c.paragraphs[0].runs:
        r.bold = True
risk_intro._p.addnext(rtbl._tbl)

# ================================================================== 1.4 Economic / Business
ph = placeholder_after(heading_index("1.4 Economic / Business Analysis (where applicable)"))
set_runs(ph, [("Illustrative value model, to be validated jointly with Fastweb + Vodafone:", False)])
bullets_after(ph, [
    [("Cost — ", True),
     ("development is engineering time only (open-source, existing compute, no licensing). Production "
      "cost is modest edge compute for the inline gate.", False)],
    [("Risk avoided — ", True),
     ("detecting a compromised or spoofed automation before it acts on the control plane cuts "
      "mean-time-to-detect from hours to seconds and contains blast-radius, avoiding outage and "
      "breach costs.", False)],
    [("Enabler of safe automation — ", True),
     ("zero-trust on machine identities is the guardrail that lets the operator safely expand network "
      "automation (closed-loop / self-healing), protecting the value of the wider AI-native roadmap.", False)],
    [("Compliance & audit — ", True),
     ("supports a NIS2 zero-trust posture; explainable, logged decisions provide audit evidence for "
      "GDPR and the EU AI Act.", False)],
    [("Operational efficiency — ", True),
     ("automates machine-identity monitoring that cannot be performed manually at API scale.", False)],
    [("Market relevance — ", True),
     ("securing AI agents on critical infrastructure is an emerging, board-level priority; AEGIS "
      "targets a gap that credential-based AAA does not cover.", False)],
])

doc.save(OUT)
print("Saved", OUT)
print("Tables:", len(doc.tables), "| paragraphs:", len(doc.paragraphs))
