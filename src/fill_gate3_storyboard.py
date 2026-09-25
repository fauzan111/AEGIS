# -*- coding: utf-8 -*-
"""
AEGIS - fills the official 5G Academy storyboard template (Boords-style,
slides/STORYBOARD Template/storyboard_template.pptx) with AEGIS's own
scene-by-scene content, mapped onto the template's own prescribed section
structure (Introduzione, problema/business, soluzione futura, parte test,
test, filosofia e regulation, finale, team) rather than inventing a new one.

Fills Title/Page fields, numbers each Scene No./Shot No. panel, and adds a
one-line shot description + script snippet in the gap below each panel -
the template's own sketch boxes are left untouched (this is a sketch
template; the visual frame in each box still needs to be drawn/pasted by
hand or by the videomaker).

Run:  .venv/Scripts/python AEGIS/src/fill_gate3_storyboard.py
In :  AEGIS/slides/STORYBOARD Template/storyboard_template.pptx (untouched)
Out:  AEGIS/slides/GATE-3/AEGIS_Gate3_Storyboard.pptx
"""

import copy
import os
from pptx import Presentation
from pptx.util import Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN

HERE = os.path.dirname(os.path.abspath(__file__))
TEMPLATE = os.path.join(HERE, "..", "slides", "STORYBOARD Template", "storyboard_template.pptx")
OUT_DIR = os.path.join(HERE, "..", "slides", "GATE-3")
OUT = os.path.join(OUT_DIR, "AEGIS_Gate3_Storyboard.pptx")

INK = RGBColor(0x02, 0x0C, 0x31)
MUTE = RGBColor(0x55, 0x55, 0x55)
FONT = "Calibri"

TITLE = "AEGIS - Zero-Trust Identity for AI Agents"

# 6 panels per page, in reading order (row1: L, R2col2, R2col3 / row2: same) -
# matches the template's own panel order top-left to bottom-right.
PANEL_LEFTS = [237490, 3729990, 7235190]
PANEL_TOPS = [696595, 3935095]
PANEL_W = 3187700
PANEL_H = 2070100
GAP_H = 1000000  # caption box height in the gap below each panel

# Page number -> (six shots: short visual description, script/voiceover line)
PAGES = {
    1: [  # Introduzione (0:00-0:30)
        ("Close-up: ID badge scan at an office door, green light", "\"A stolen ID badge still opens the door.\""),
        ("Door opens, person walks through - ordinary, unremarkable", "\"The badge is valid.\""),
        ("Match-cut: laptop screen, 5G network diagram appears", "\"But is the person behind it really who they claim to be?\""),
        ("AI/automation agent icon connects into the 5G core", "AI agents and automations now call the network directly."),
        ("Green checkmark overlay: \"credential: valid\"", "A stolen or misused credential still passes every check."),
        ("Cut to red alert icon / question mark", "\"AEGIS\" title card reveal."),
    ],
    2: [  # continuation of Introduzione / bridge into problem
        ("AEGIS logo settles center frame", "\"This is AEGIS.\""),
        ("Quick architecture teaser: 4 icons in a row (no labels yet)", "A zero-trust identity gate for AI agents on the network."),
        ("Team members at laptops / workstation B-roll", "Built by Team 4, 5G Academy 2026."),
        ("5G Academy + Fastweb + Vodafone logos, brief", "In partnership with Fastweb and Vodafone."),
        ("Transition wipe to network operations center visual", "Bridge into the problem."),
        ("Fade to black / section break", ""),
    ],
    3: [  # Presentazione problema / e business
        ("Graphic: orchestration tool calling AMF/SMF/UDM APIs directly", "Orchestration and automation tools now call 5G Network Functions directly."),
        ("Split screen: valid credential vs. abnormal behavior", "Authentication proves you have the right key."),
        ("Text overlay, large: \"It never asks if you're still you.\"", "It never asks whether you're still behaving like yourself."),
        ("NOC/SOC operator looking concerned at a dashboard", "A compromised or stolen credential can sit undetected for weeks."),
        ("Business-impact graphic: cost/risk of an undetected breach", "The business cost of a one-time-only identity check."),
        ("Transition: question mark resolves into AEGIS logo", "\"There's a better way.\""),
    ],
    4: [  # Presentazione soluzione futura
        ("AEGIS architecture diagram, full reveal", "AEGIS builds a behavioural fingerprint for every agent."),
        ("Animated: OBSERVE stage highlighted", "OBSERVE - every request is captured: who, what, when, how."),
        ("Animated: FINGERPRINT stage highlighted", "FINGERPRINT - that stream becomes a behavioural profile."),
        ("Animated: SCORE stage highlighted", "SCORE - live behaviour is compared to that profile, in real time."),
        ("Animated: DECIDE stage highlighted, PASS/STEP-UP/BLOCK gate", "DECIDE - pass it, challenge it, or shut it down, instantly."),
        ("Zero-trust tagline overlay (NIST SP 800-207 callout)", "Never trust. Always verify. Continuously."),
    ],
    5: [  # Parte test (live demo, part 1)
        ("Screen capture: AEGIS live dashboard home / topology view", "Here it is, live - not a mockup."),
        ("Click \"Launch Live Attack\" on the dashboard", "Watch what happens when an agent starts behaving like an attacker."),
        ("Risk score climbing in real time on screen", "Probing endpoints it's never touched, at a rate it's never used."),
        ("Incident inspector panel: plain-language explanation appears", "AEGIS explains exactly why, in language a security analyst can act on."),
        ("Decision badge flips to BLOCK, red highlight", "And shuts it down before real damage is done."),
        ("Before/After AEGIS comparison view on screen", "This is the difference AEGIS makes."),
    ],
    6: [  # Test (continuation - real-world validation)
        ("Terminal footage: Docker containers starting up", "We didn't stop at simulation."),
        ("Real 5G core topology diagram (Open5GS: AMF/SMF/UDM/etc.)", "We built a real 5G core."),
        ("Screen: a real device registering onto the real core", "Connected a real device to it."),
        ("Real captured network traffic on screen (terminal/Wireshark-style)", "And tested AEGIS against real network traffic."),
        ("AEGIS flags the real attack: BLOCK decision, real data", "The exact same detection code. No changes. It worked."),
        ("Stat card overlay: ROC-AUC 0.98, recall 91%, FPR 1.1%", "Benchmarked, not just claimed."),
    ],
    7: [  # Filosofia e regulation (has a "stacco"/cut marker mid-page)
        ("Text card: \"Explainable, not a black box\"", "Every decision AEGIS makes comes with a plain-language reason."),
        ("Incident inspector close-up: reasoning text on screen", "No guessing at why an alert fired."),
        ("Privacy/GDPR visual: subscriber ID being hashed/anonymised", "Subscriber identifiers are protected, never stored raw."),
        ("[STACCO - hard cut]", ""),
        ("Production vision: AEGIS gate sitting at the network edge (FASTedge)", "In production, AEGIS runs at the network edge - close to the core, low latency."),
        ("Closing philosophy line, text on screen", "Security that explains itself, built to be trusted, not just deployed."),
    ],
    8: [  # Finale
        ("Fast montage: badge scan, dashboard, real core, BLOCK decision", "Recap montage - the whole story in five seconds."),
        ("Tagline card, large text", "\"Trust can't be a one-time check.\""),
        ("Tagline continues", "\"It has to be continuous.\""),
        ("AEGIS logo, GitHub + live dashboard links on screen", "aegis5gacademy.streamlit.app - see it live."),
        ("5G Academy + Fastweb + Vodafone branding", "5G Academy 2026, in partnership with Fastweb and Vodafone."),
        ("Fade to black, thank you", "Thank you."),
    ],
}


def set_first_run_text(shape, text):
    tf = shape.text_frame
    p0 = tf.paragraphs[0]
    if not p0.runs:
        p0.add_run()
    p0.runs[0].text = text
    for r in list(p0.runs[1:]):
        r._r.getparent().remove(r._r)
    for p in list(tf.paragraphs[1:]):
        p._p.getparent().remove(p._p)


def find_all(slide, snippet):
    return [shp for shp in slide.shapes if shp.has_text_frame and snippet in shp.text_frame.text]


def add_caption(slide, left, top, text_desc, text_script):
    tb = slide.shapes.add_textbox(Emu(left), Emu(top), Emu(PANEL_W), Emu(GAP_H))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = 0; tf.margin_right = 0; tf.margin_top = 0; tf.margin_bottom = 0
    p1 = tf.paragraphs[0]
    p1.alignment = PP_ALIGN.LEFT
    r1 = p1.add_run(); r1.text = text_desc
    r1.font.size = Pt(9); r1.font.name = FONT; r1.font.bold = True; r1.font.color.rgb = INK
    if text_script:
        p2 = tf.add_paragraph()
        p2.alignment = PP_ALIGN.LEFT
        r2 = p2.add_run(); r2.text = text_script
        r2.font.size = Pt(8.5); r2.font.name = FONT; r2.font.italic = True; r2.font.color.rgb = MUTE


def main():
    prs = Presentation(TEMPLATE)

    for page_num, slide in enumerate(prs.slides, start=1):
        if page_num == 9:
            continue  # team slide handled separately below

        title_shape = find_all(slide, "Title:")
        if title_shape:
            set_first_run_text(title_shape[0], f"Title: {TITLE}")
        page_shape = find_all(slide, "Page:")
        if page_shape:
            set_first_run_text(page_shape[0], f"Page: {page_num}")

        scene_boxes = find_all(slide, "Scene No.")
        shot_boxes = find_all(slide, "Shot No.")
        # sort by position (top then left) to match reading order
        scene_boxes.sort(key=lambda s: (s.top, s.left))
        shot_boxes.sort(key=lambda s: (s.top, s.left))

        for i, shp in enumerate(scene_boxes):
            set_first_run_text(shp, f"Scene {i + 1}")
        for i, shp in enumerate(shot_boxes):
            set_first_run_text(shp, f"Shot {i + 1}")

        shots = PAGES.get(page_num, [])
        panel_index = 0
        for row_top in PANEL_TOPS:
            for col_left in PANEL_LEFTS:
                if panel_index >= len(shots):
                    break
                desc, script = shots[panel_index]
                add_caption(slide, col_left, row_top + PANEL_H + 40000, desc, script)
                panel_index += 1

    # ---- team slide (page 9) ----
    team_slide = prs.slides[8]
    group_shape = find_all(team_slide, "Group 1")
    if group_shape:
        set_first_run_text(group_shape[0], "Team 4 - AEGIS")
    team_tb = team_slide.shapes.add_textbox(Emu(600000), Emu(2800000), Emu(9000000), Emu(2000000))
    tf = team_tb.text_frame
    tf.word_wrap = True
    for i, name in enumerate(["Fauzan Ejaz (Captain)", "Adithya Zacharia Valavi",
                              "Ghazanfar Anees Siddiqui", "Llagami Tedi"]):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        r = p.add_run(); r.text = name
        r.font.size = Pt(20); r.font.name = FONT; r.font.color.rgb = INK

    os.makedirs(OUT_DIR, exist_ok=True)
    prs.save(OUT)
    print(f"Saved {OUT}")
    print(f"Total slides: {len(prs.slides)}")


if __name__ == "__main__":
    main()
