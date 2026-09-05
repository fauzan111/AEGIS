"""Remove em dashes from every AEGIS doc/deck/markdown file.

Replaces the em dash character with a plain hyphen, run by run, so
existing text formatting (bold, color, size) is untouched.
"""

from pathlib import Path
from docx import Document
from pptx import Presentation

ROOT = Path(__file__).resolve().parents[1]

DOCX_FILES = [
    ROOT / "docs" / "AEGIS_Tech_Overview.docx",
    ROOT / "docs" / "AEGIS_Prototype_Explained.docx",
    ROOT / "slides" / "AEGIS_Prototype_Explained.docx",
    ROOT / "slides" / "Group 4_AEGIS_Gate1_Report.docx",
]
PPTX_FILES = [
    ROOT / "slides" / "AEGIS_Gate1  -  NEW.pptx",
    ROOT / "slides" / "AEGIS_Gate1_Official.pptx",
]
MD_FILES = [
    ROOT / "README.md",
    ROOT / "docs" / "AEGIS_threat_model_and_architecture.md",
    ROOT / "reports" / "aegis_baseline_metrics.md",
]

EM_DASH = "—"


def clean_runs(paragraphs):
    count = 0
    for p in paragraphs:
        for run in p.runs:
            if EM_DASH in run.text:
                count += run.text.count(EM_DASH)
                run.text = run.text.replace(EM_DASH, "-")
    return count


def clean_docx(path):
    doc = Document(path)
    total = clean_runs(doc.paragraphs)
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                total += clean_runs(cell.paragraphs)
    for section in doc.sections:
        for part in (section.header, section.footer):
            total += clean_runs(part.paragraphs)
    if total:
        doc.save(path)
    print(f"{path.name}: replaced {total} em dash(es)")


def clean_text_frame(tf):
    return clean_runs(tf.paragraphs)


def clean_pptx(path):
    prs = Presentation(path)
    total = 0
    for slide in prs.slides:
        for shape in slide.shapes:
            if shape.has_text_frame:
                total += clean_text_frame(shape.text_frame)
            if shape.has_table:
                for row in shape.table.rows:
                    for cell in row.cells:
                        total += clean_text_frame(cell.text_frame)
        if slide.has_notes_slide:
            total += clean_text_frame(slide.notes_slide.notes_text_frame)
    if total:
        prs.save(path)
    print(f"{path.name}: replaced {total} em dash(es)")


def clean_md(path):
    text = path.read_text(encoding="utf-8")
    count = text.count(EM_DASH)
    if count:
        path.write_text(text.replace(EM_DASH, "-"), encoding="utf-8")
    print(f"{path.name}: replaced {count} em dash(es)")


for f in DOCX_FILES:
    if f.exists():
        clean_docx(f)
    else:
        print(f"MISSING: {f}")

for f in PPTX_FILES:
    if f.exists():
        clean_pptx(f)
    else:
        print(f"MISSING: {f}")

for f in MD_FILES:
    if f.exists():
        clean_md(f)
    else:
        print(f"MISSING: {f}")
