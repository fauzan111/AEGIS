# -*- coding: utf-8 -*-
"""One-off patch: add presenter notes (30-min, 4-speaker script) and
strengthen the business-case cell directly on the already-built
AEGIS_Gate1_Official.pptx, without needing the original template file
(which is temporarily missing from slides/). Safe to re-run.
"""
import os
from pptx import Presentation
from pptx.util import Pt

HERE = os.path.dirname(os.path.abspath(__file__))
SLIDES = os.path.join(HERE, "..", "slides")
OUT = os.path.join(SLIDES, "AEGIS_Gate1_Official.pptx")

prs = Presentation(OUT)


def set_cell(cell, text, size=9):
    cell.text = text
    cell.text_frame.word_wrap = True
    for p in cell.text_frame.paragraphs:
        for r in p.runs:
            r.font.size = Pt(size)


# ---- strengthen the business-case cell (slide 14, "Open points" table) ----
s14 = prs.slides[13]
for shape in s14.shapes:
    if shape.has_table:
        tbl = shape.table
        set_cell(tbl.cell(1, 1), (
            "Cost: no licensing, 100% open-source, runs on a laptop; production adds "
            "only light edge compute.\n"
            "Value: cuts detection time from hours to seconds, enables safer network "
            "automation, and gives audit-ready evidence for NIS2/GDPR/EU AI Act.\n"
            "To be jointly costed and validated with Fastweb + Vodafone."
        ), size=9)
        break

# ---- presenter notes ----
NOTES = [
"FAUZAN - 0:00-0:30\n"
"Good morning, we're Team 4, presenting AEGIS: a zero-trust identity check "
"for the AI agents and automations that now operate the 5G network.",

"FAUZAN - 0:30-1:00\n"
"Here's how we'll walk through this: first, we'll explain what AEGIS "
"actually is and how it works in plain language, no assumed background "
"needed. Then in the last section we'll cover the technologies we used, "
"our project schedule, feasibility, and the business case, which is "
"exactly what Gate 1 asks us to demonstrate.",

"FAUZAN - 1:00-2:30\n"
"One line to remember us by: credentials prove what you have, AEGIS "
"verifies how you behave. Today, an automated agent proves who it is with "
"a password or token, and that's the only check. A stolen credential "
"still passes that check every time. AEGIS adds a second, continuous "
"layer: it learns how each agent normally behaves and checks every "
"request against that pattern. It's built entirely on open-source tools, "
"runs on a laptop, and every decision it makes comes with a "
"plain-language reason, never a black box.",

"ADITHYA - 2:30-4:30\n"
"Before we go further, quick grounding: a 5G network isn't one big "
"machine, it's six specialised systems, each handling one job, like "
"departments in a company. AMF tracks where a device is. SMF manages "
"sessions. NRF is the internal directory. PCF enforces policy. UDM holds "
"subscriber data. NEF is the controlled front door for outside "
"automations. Every agent we model talks to these six systems, and "
"that's the surface AEGIS protects.",

"ADITHYA - 4:30-6:30\n"
"We modelled five distinct automated agents, each with its own job, "
"schedule, and normal routine: a session orchestrator running around the "
"clock, a nightly inventory job, a closed-loop monitoring agent, a simple "
"NRF heartbeat service, and a business-hours provisioning agent. This "
"matters because AEGIS doesn't look for 'bad behaviour' in the abstract, "
"it looks for behaviour that doesn't match what that specific agent "
"normally does, so each needed a genuinely distinct personality to prove "
"the idea works.",

"ADITHYA - 6:30-8:30\n"
"We built this in three parts. First, a traffic simulator we wrote "
"ourselves, not AI, a scripted program that generated about 210,000 "
"realistic requests for a virtual week. Second, we injected seven kinds "
"of attacks into that traffic, each labelled, so we know exactly which "
"requests are attacks, only possible because the data is simulated. "
"Third, the actual detection engine, this is where real machine learning, "
"an Isolation Forest, is used to spot the attacks we hid inside.",

"GHAZANFAR - 8:30-10:30\n"
"We assume the attacker already has a valid credential, the realistic, "
"hard case, so what changes is behaviour, not the password. We modelled "
"seven attacks: impersonation, a compromised agent drifting into new "
"behaviour, reconnaissance scanning, a volumetric flood, slow data "
"exfiltration, a broken call sequence, and scope creep. All seven are "
"built into our prototype, and as we'll show, all seven are currently "
"detected.",

"GHAZANFAR - 10:30-13:00\n"
"Here's how the gate actually decides, four steps, fully automatic, no "
"person watching in real time. Observe: every request gets logged. "
"Fingerprint: we compare the last 60 seconds of an agent's activity "
"against its own normal pattern. Score: hard rules plus a machine-"
"learning model, trained only on legitimate traffic, combine into one "
"risk number. Decide: fixed thresholds turn that number into pass, "
"step-up, or block. A human only gets involved afterwards, reviewing a "
"step-up or block once it's already been raised.",

"GHAZANFAR - 13:00-15:30\n"
"This isn't a design target, it's a working prototype with real numbers. "
"On over eleven thousand test windows the model never trained on: 0.99 "
"ROC-AUC, 98% of injected attacks caught, only 0.8% of legitimate traffic "
"wrongly flagged, all seven threats detected. The honest caveat: this "
"proves the method works on data we built ourselves, real traffic "
"validation is the next step, which is exactly what we're asking Fastweb "
"and Vodafone for.",

"LLAGAMI - 15:30-17:00\n"
"Now the part Gate 1 specifically asks us to cover: our technologies, "
"schedule, feasibility and business case. Here's our full timeline from "
"late July through to the final video in November. Everything up to and "
"including today's Gate 1 milestone is already complete, this isn't a "
"plan, it's work already done.",

"LLAGAMI - 17:00-18:30\n"
"Breaking that down by phase: the data foundation and baseline detector "
"you just saw results from are complete. From here, we harden the "
"detector, submit Gate 2 on 22 September, integrate real data if it's "
"granted, and deliver the full live demo at Gate 3 on 15 October.",

"LLAGAMI - 18:30-20:00\n"
"And here's that same schedule as a Gantt chart. Green is done, blue is "
"upcoming work, amber is the optional real-data integration track. We're "
"currently ahead of where we need to be for Gate 2.",

"LLAGAMI - 20:00-23:00\n"
"On feasibility: we identified five real risks and a mitigation for each, "
"from data not being shared in time, to FASTedge not exposing an inline "
"hook, to the detector not generalising to real traffic noise. The "
"biggest one, data access, is fully mitigated already: we're "
"synthetic-first, so no data approval process blocks us from having a "
"working system today. Real data only makes the result stronger, it "
"isn't a dependency.",

"LLAGAMI - 23:00-26:30\n"
"On technologies and cost: everything is open-source, Python, "
"scikit-learn, Streamlit, no licensing, runs on a laptop, production only "
"adds light edge compute. On business value: this cuts detection time "
"from hours to seconds, it's the guardrail that lets an operator safely "
"expand network automation, and every decision is logged and explainable, "
"which is direct audit evidence for NIS2, GDPR, and the EU AI Act. We're "
"tracking this work area by area, financial, engineering, communication, "
"and next steps, so nothing falls through the cracks before Gate 2.",

"FAUZAN - 26:30-28:30\n"
"So, week by week from here: hardening the detector, finishing the "
"numerical performance analysis for Gate 2, then integrating real data if "
"granted, and finally the live end-to-end dashboard demo for Gate 3 on 15 "
"October.",

"FAUZAN - 28:30-30:00\n"
"To close: the hardest part, proving the detection method actually works, "
"is already done, working, and in front of you today. We're confident in "
"this direction and looking forward to your questions. Thank you.",
]

assert len(NOTES) == len(prs.slides), (len(NOTES), len(prs.slides))
for slide, note in zip(prs.slides, NOTES):
    slide.notes_slide.notes_text_frame.text = note

prs.save(OUT)
print(f"Patched {OUT} - {len(prs.slides)} slides, notes added")
