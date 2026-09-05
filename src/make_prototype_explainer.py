# -*- coding: utf-8 -*-
"""Build a plain-language, part-by-part explainer of the AEGIS prototype.

Written for a non-technical reader (5G Academy internal head / stakeholders).
Walks through: the 5G network being modelled, the 5 simulated agents, how the
synthetic data + prototype were built, how attacks are simulated, who/what
decides something is an attack, and how the evaluation numbers (ROC-AUC,
precision, recall, F1, false-positive rate) are actually produced.

Run:  .venv/Scripts/python AEGIS/src/make_prototype_explainer.py
"""

from pathlib import Path
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "docs" / "AEGIS_Prototype_Explained.docx"

NAVY = RGBColor(0x0B, 0x1F, 0x3A)
ACCENT = RGBColor(0x1E, 0x5A, 0x8A)
GREY = RGBColor(0x55, 0x55, 0x55)


def set_cell_shading(cell, hex_color):
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), hex_color)
    cell._tc.get_or_add_tcPr().append(shd)


def style_table(table):
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for row_idx, row in enumerate(table.rows):
        for cell in row.cells:
            for p in cell.paragraphs:
                for run in p.runs:
                    run.font.size = Pt(10)
                if row_idx == 0:
                    for run in p.runs:
                        run.font.bold = True
                        run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
            if row_idx == 0:
                set_cell_shading(cell, "1E5A8A")
            elif row_idx % 2 == 0:
                set_cell_shading(cell, "F2F5F8")


def add_table(doc, headers, rows, widths=None):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    for i, h in enumerate(headers):
        table.rows[0].cells[i].text = h
    for r in rows:
        cells = table.add_row().cells
        for i, val in enumerate(r):
            cells[i].text = str(val)
    style_table(table)
    if widths:
        for row in table.rows:
            for i, w in enumerate(widths):
                row.cells[i].width = Inches(w)
    return table


def h1(doc, text, number=None):
    p = doc.add_heading(level=1)
    run = p.add_run(f"{number}. {text}" if number else text)
    run.font.color.rgb = NAVY
    return p


def h2(doc, text):
    p = doc.add_heading(level=2)
    run = p.add_run(text)
    run.font.color.rgb = ACCENT
    return p


def body(doc, text, bold=False, italic=False):
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.font.size = Pt(11)
    run.font.bold = bold
    run.font.italic = italic
    return p


def bullet(doc, text):
    p = doc.add_paragraph(style="List Bullet")
    run = p.add_run(text)
    run.font.size = Pt(11)
    return p


doc = Document()
style = doc.styles["Normal"]
style.font.name = "Calibri"
style.font.size = Pt(11)

# --- Title page ---
title = doc.add_paragraph()
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = title.add_run("AEGIS")
run.font.size = Pt(36)
run.font.bold = True
run.font.color.rgb = NAVY

sub = doc.add_paragraph()
sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = sub.add_run("How We Built the Prototype")
run.font.size = Pt(16)
run.font.color.rgb = ACCENT

sub2 = doc.add_paragraph()
sub2.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = sub2.add_run("A plain-language, part-by-part walkthrough")
run.font.size = Pt(13)
run.font.italic = True
run.font.color.rgb = GREY

sub3 = doc.add_paragraph()
sub3.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = sub3.add_run(
    "Written for a non-technical reader: what network we modelled, what "
    "agents we simulated, how the prototype was actually built, how attacks "
    "were generated, who decides something is an attack, and how we "
    "measured whether it works."
)
run.font.size = Pt(11)
run.font.color.rgb = GREY

doc.add_paragraph()
meta = doc.add_paragraph()
meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = meta.add_run(
    "5G Academy 2026, Fastweb + Vodafone, Topic 2, Security\n"
    "Team 4, Fauzan Ejaz (Captain), Adithya Zacharia Valavi, "
    "Ghazanfar Anees Siddiqui, Llagami Tedi"
)
run.font.size = Pt(10)
run.font.color.rgb = GREY

doc.add_page_break()

# ============================================================ PART 0
h1(doc, "The One-Minute Version", 0)
body(
    doc,
    "Before the detail: AEGIS is a security check that sits between "
    "automated software \"agents\" (bots, orchestration tools, automation "
    "scripts) and the 5G network's control systems. Today, those agents "
    "prove who they are with a password or token, and that is the only "
    "check. AEGIS adds a second check: it watches how each agent normally "
    "behaves, and if an agent starts acting strangely, even with a valid "
    "password, it gets challenged or blocked. To prove this idea works, we "
    "built a working, running prototype ourselves, using computer-generated "
    "(simulated) network traffic, because we do not yet have access to "
    "real network data. This document explains, step by step and without "
    "jargon, exactly what we built and how.",
)

# ============================================================ PART 1
h1(doc, "The 5G Network We Are Modelling (the topology)", 1)
body(
    doc,
    "A 5G network is not one big machine; it is a set of specialised "
    "software systems, called Network Functions, that each handle one job. "
    "Automated agents do not talk to \"the network\" in the abstract, they "
    "call these specific systems over the network's internal API layer, "
    "known as the Service-Based Interface. Think of each Network Function "
    "as a department in a company, and the API calls as the standard forms "
    "each department accepts.",
)
nf_rows = [
    ("AMF", "Access & Mobility Management", "Keeps track of where a device is and whether it is reachable, like a front desk that knows who is in the building."),
    ("SMF", "Session Management", "Opens, updates, and closes a device's data session, like the department that opens and closes a service ticket."),
    ("NRF", "Network Repository", "The internal directory: lets every other system register itself and find each other, like a company phone book."),
    ("PCF", "Policy Control", "Decides and enforces the rules a session must follow (speed, priority, permissions), like a rules and approvals desk."),
    ("UDM", "Unified Data Management", "Holds and serves subscriber data on request, like the records department."),
    ("NEF", "Network Exposure", "The controlled \"front door\" that lets outside applications and automations request network events or actions, like a reception window for external requests."),
]
add_table(doc, ["Network Function", "Full name", "In plain terms"], nf_rows, widths=[0.8, 1.9, 3.8])
doc.add_paragraph()
body(
    doc,
    "In our prototype, every one of these six Network Functions is "
    "represented, each with a handful of realistic operations (for example "
    "SMF has \"create a session\", \"update a session\", \"release a "
    "session\"). This is the topology, the map of departments and doors, "
    "that all of our simulated agents and attacks operate on.",
)

# ============================================================ PART 2
h1(doc, "The Five Agents We Modelled", 2)
body(
    doc,
    "Rather than simulate one generic \"agent\", we built five distinct "
    "automations, each with its own job, its own normal working hours, its "
    "own typical speed, and its own fixed routine of which departments it "
    "is allowed to visit and in what order. This matters because AEGIS "
    "does not look for \"bad behaviour\" in general, it looks for "
    "behaviour that does not match what that specific agent normally "
    "does, so each agent needed a genuinely distinct, realistic personality.",
)
agent_rows = [
    ("Session orchestrator", "Opens, updates, and closes user sessions across "
     "AMF, SMF, PCF", "Runs 24/7, steady pace", "Like an operations system "
     "that is constantly opening and closing service tickets"),
    ("Inventory job", "Discovers and reads records from NRF, UDM only", "Runs "
     "briefly overnight, 1am-5am", "Like a nightly stock-take that quietly "
     "checks what exists and then stops"),
    ("Closed-loop assurance", "Subscribes to live events from AMF, NEF, PCF",
     "Runs 24/7, moderate-to-high pace", "Like a monitoring system that "
     "watches for problems and reacts automatically"),
    ("NRF heartbeat service", "Registers and \"pings\" NRF only", "Runs "
     "24/7, very steady, low volume", "Like a simple check-in system saying "
     "\"I'm still alive\" at regular intervals"),
    ("Provisioning agent", "Updates subscriber parameters via NEF, UDM",
     "Runs business hours, 6am-10pm", "Like a customer-service system "
     "making account changes during the working day"),
]
add_table(doc, ["Agent", "What it does", "When / how fast", "Everyday analogy"], agent_rows, widths=[1.2, 2.1, 1.3, 1.9])
doc.add_paragraph()
body(
    doc,
    "Each agent's normal routine, which departments it visits, in what "
    "order, how often, at what times of day, and with what error rate, is "
    "learned by AEGIS as that agent's personal \"fingerprint\". Everything "
    "AEGIS later flags as suspicious is suspicious relative to one of these "
    "five personal fingerprints, not against some generic notion of "
    "normal.",
)

# ============================================================ PART 3
h1(doc, "How We Actually Built the Prototype", 3)
body(
    doc,
    "We could not use real Fastweb or Vodafone network data (we do not yet "
    "have access to it), so we built the prototype in three parts, all of "
    "which run end to end today on an ordinary laptop.",
)
build_rows = [
    ("1. A traffic simulator (not AI)", "A Python program we wrote "
     "ourselves that plays out a virtual week of activity for all five "
     "agents against the six Network Functions above, following each "
     "agent's own speed, hours, and routine. This is a scripted simulation, "
     "not a machine-learning model; it is closer to a very detailed, "
     "rule-based traffic generator than to an AI. It produced about "
     "210,000 individual requests."),
    ("2. Injected attacks", "Into that otherwise normal week, we inserted "
     "seven kinds of attack episodes (described in Part 4), each one "
     "labelled so we know afterwards exactly which requests were an "
     "attack and which were not. This labelling is only possible because "
     "the data is simulated; real traffic never comes pre-labelled."),
    ("3. A detection engine (this part is machine learning)", "A separate "
     "program that reads the traffic, builds a behavioural fingerprint per "
     "agent, and scores every short time window of activity for risk. This "
     "is where an actual machine-learning model, an Isolation Forest, is "
     "used, described in Part 5."),
]
add_table(doc, ["Step", "What it is"], build_rows, widths=[2.0, 4.4])
doc.add_paragraph()
body(
    doc,
    "So, to directly answer \"which model generated the data\": the data "
    "was not generated by an AI model at all, it was generated by a "
    "deterministic simulation program we wrote, designed to look like "
    "realistic 5G control-plane traffic. The only place a machine-learning "
    "model is actually used is afterwards, on the detection side, to spot "
    "the attacks we had hidden inside that simulated traffic.",
)
p = body(doc, "Why simulate at all, rather than wait for real data? ", bold=True)
p.add_run(
    "Because it lets us prove the detection approach works end to end, on "
    "a known, labelled answer key, before asking anyone for sensitive "
    "network data, and because our timeline (Gate 2 on 22 September, Gate "
    "3 on 15 October) does not allow us to wait for a data-sharing process "
    "to complete before we have something working to show."
).italic = True

# ============================================================ PART 4
h1(doc, "The Attacks: What They Are and How We Simulated Them", 4)
body(
    doc,
    "We modelled seven attack types. In every case, we assume the "
    "attacker already has a valid password or token for one of the five "
    "agents, the realistic and hardest case, because a stolen or misused "
    "credential still passes a normal login check. What changes is not the "
    "credential, it is the behaviour.",
)
attack_rows = [
    ("T1", "Identity spoofing", "We take one agent's identity and make it "
     "suddenly act like a completely different agent, visiting departments "
     "and following a routine that belongs to someone else."),
    ("T2", "Compromised agent", "We take a real agent and let it \"drift\", "
     "it starts visiting a department it has never been to before, and its "
     "error rate creeps up."),
    ("T3", "Reconnaissance / scanning", "We make an agent probe broadly "
     "across many departments and IDs it would not normally touch, "
     "generating lots of \"not found\" or \"forbidden\" responses, the "
     "digital equivalent of someone rattling every door handle in a "
     "building."),
    ("T4", "Volumetric flood", "We make an agent's request rate spike to "
     "several times its normal pace, flooding one department."),
    ("T5", "Slow data exfiltration", "We make an agent quietly pull "
     "unusually large amounts of data, off its normal hours, over an "
     "extended period, rather than in one obvious burst."),
    ("T6", "Replay / broken sequence", "We make an agent perform an "
     "operation out of the valid order, for example \"closing\" a session "
     "that was never \"opened\", which should never happen in normal use."),
    ("T7", "Scope creep", "We make an agent touch a department that is "
     "outside its assigned role entirely."),
]
add_table(doc, ["ID", "Attack", "How we simulated it"], attack_rows, widths=[0.5, 1.6, 4.3])
doc.add_paragraph()
body(
    doc,
    "Does the agent's own real work fail during an attack? No. This is an "
    "important point: we are not simulating a crash, a bug, or an outage. "
    "The agent's legitimate work continues exactly as normal throughout; "
    "the attack is an additional burst, or a change, of behaviour layered "
    "on top of, or instead of, a short slice of that agent's activity. "
    "Nothing about an attack request looks technically broken or "
    "unauthorised at the level of a simple login check, which is exactly "
    "why this is a hard, realistic problem: credential-based security "
    "alone would let every one of these seven attacks straight through.",
)

# ============================================================ PART 5
h1(doc, "Who, or What, Decides It Is an Attack", 5)
body(
    doc,
    "There is no person watching a screen deciding this in real time, the "
    "decision is fully automatic. Every 60 seconds, for every agent, the "
    "system looks at everything that agent did in that window and follows "
    "four steps.",
)
pipeline_rows = [
    ("1. Observe", "Every request is logged: which agent, which department, "
     "which operation, what time, how big the response was, what status "
     "code came back."),
    ("2. Fingerprint", "The last 60 seconds of that agent's activity is "
     "turned into a profile: how many requests, which departments, how "
     "varied, error rate, response sizes, time of day, and compared "
     "against that same agent's own known-normal profile."),
    ("3. Score", "Two independent checks run side by side and are combined "
     "into one risk number between 0 and 1. The first is a set of hard, "
     "clear-cut rules (for example, \"this agent just touched a department "
     "it has never touched before\", which is an unambiguous red flag). "
     "The second is a machine-learning model, an Isolation Forest, trained "
     "only on each agent's own past normal behaviour, which learns what "
     "\"normal for this agent\" statistically looks like and flags "
     "anything that does not fit, without ever needing to be shown an "
     "actual attack example to learn from."),
    ("4. Decide", "The combined risk number is compared against two fixed "
     "cut-off points. Below 0.5 is treated as normal and passed through. "
     "Between 0.5 and 0.8 triggers a step-up challenge (re-verify, "
     "approve, or slow down). Above 0.8 triggers a block, and the identity "
     "is flagged for review."),
]
add_table(doc, ["Step", "What happens"], pipeline_rows, widths=[1.3, 5.1])
doc.add_paragraph()
body(
    doc,
    "So the honest answer to \"who decides\" is: the combination of the "
    "rules and the machine-learning score decides, automatically, based "
    "purely on how far that moment of behaviour has drifted from that "
    "agent's own established normal. A human, for example a security "
    "analyst, only enters the picture afterwards, when a step-up or block "
    "decision surfaces on the dashboard for someone to review or act on. "
    "Nothing is decided by a person in the moment; the system's job is "
    "specifically to remove that need for constant manual watching.",
)

# ============================================================ PART 6
h1(doc, "How We Measured Whether It Actually Works", 6)
body(
    doc,
    "Because every request in our simulated data is labelled, we know in "
    "advance, from the generator itself, exactly which moments were "
    "genuinely an attack and which were not. This lets us grade the "
    "detector's decisions the same way a teacher grades an exam against an "
    "answer key, by comparing what the system decided against what we "
    "already know the truth to be.",
)
body(
    doc,
    "We split the simulated week into three parts: one part to teach the "
    "model each agent's normal behaviour, a second, separate part to "
    "calibrate its risk scale fairly, and a third part, 11,343 windows, "
    "that the model never saw during teaching or calibration, used purely "
    "for the final exam.",
)
metrics_explain_rows = [
    ("ROC-AUC (0.992)", "Answers: across every possible sensitivity "
     "setting, how well does the risk score separate real attacks from "
     "normal activity? 1.0 would be a perfect separation every time; 0.5 "
     "would be no better than a coin flip. 0.992 means the separation is "
     "very close to perfect."),
    ("Recall / detection rate (98.3%)", "Answers: of all the real attacks "
     "that were actually hidden in the test data, what percentage did the "
     "system actually catch, at the cut-off points we chose? We caught "
     "just over 98 out of every 100."),
    ("Precision (79%)", "Answers: of everything the system flagged as "
     "suspicious, what percentage really was an attack, rather than a "
     "false alarm? About 79 out of every 100 flags were genuine."),
    ("F1 score (0.876)", "A single number that balances recall and "
     "precision together, useful as a shorthand, but the two numbers above "
     "are the ones that actually matter operationally."),
    ("False-positive rate (0.83%)", "Answers: of all the completely normal "
     "activity in the test data, what percentage did the system wrongly "
     "flag? Fewer than 1 in 100 normal moments was wrongly flagged."),
]
add_table(doc, ["Metric", "What it actually tells you"], metrics_explain_rows, widths=[2.0, 4.4])
doc.add_paragraph()
body(
    doc,
    "In short: these numbers are not opinions or estimates, they come "
    "directly from comparing the system's automated decisions against a "
    "known, pre-labelled answer key on data the system never saw while it "
    "was learning.",
)

# ============================================================ PART 7
h1(doc, "What You Would Actually See on Screen", 7)
body(
    doc,
    "All of the above is not just a theory on paper, it runs as a live, "
    "clickable dashboard, styled like a security-operations screen. This "
    "is what someone watching it would actually see.",
)
dash_rows = [
    ("Top KPI strip", "Six live numbers at a glance: total windows seen, "
     "how many were PASS / STEP-UP / BLOCK, what percentage of real "
     "attacks were caught, and what percentage of false alarms occurred."),
    ("Risk-over-time chart", "Every window plotted as a dot over time, "
     "coloured by risk; two horizontal lines mark the STEP-UP and BLOCK "
     "cut-offs, so you can visually see risk crossing the line at the "
     "moment an attack happens."),
    ("Detection-per-threat chart", "A bar per attack type (T1-T7) showing "
     "what percentage of that specific attack was caught."),
    ("Incident inspector", "A searchable table of every flagged window, "
     "worst risk first. Clicking one shows a plain-English explanation, "
     "for example \"out-of-scope access - called a Network Function "
     "outside this agent's onboarded scope\", or \"request-rate spike - "
     "possible volumetric abuse\", so the decision is never a black box, "
     "it always comes with a reason a person can read and judge for "
     "themselves."),
    ("Live threshold sliders", "The STEP-UP and BLOCK cut-off points are "
     "adjustable sliders on screen, not hard-coded, so a reviewer can see "
     "in real time how many alerts go up or down as the sensitivity "
     "changes."),
]
add_table(doc, ["Dashboard panel", "What it shows"], dash_rows, widths=[1.6, 4.8])
doc.add_paragraph()
body(
    doc,
    "In other words, the prototype is not just a spreadsheet of numbers, "
    "it is a working tool that someone with no coding background could "
    "sit down at and understand within a couple of minutes: which agents "
    "are being watched, what got flagged, and, critically, why.",
)

# ============================================================ PART 8
h1(doc, "Why We Made These Specific Design Choices", 8)
body(
    doc,
    "A few numbers in this project (60 seconds, 0.5, 0.8, the choice of "
    "Isolation Forest) look arbitrary unless you know why they were "
    "picked. They were not guessed; each is a deliberate, explainable "
    "trade-off.",
)
choice_rows = [
    ("Why a 60-second window?", "Long enough to gather a meaningful "
     "handful of requests from an agent so a pattern can actually be "
     "seen, short enough that a fast attack, like a volumetric flood, "
     "is caught within roughly a minute rather than being averaged away "
     "over a much longer period."),
    ("Why two thresholds (0.5 and 0.8) instead of one?", "A single "
     "yes/no cut-off forces every borderline case into a hard block, "
     "which is exactly how legitimate traffic gets wrongly disrupted. "
     "Splitting risk into three bands lets the system respond "
     "proportionally, a mildly unusual window gets a lightweight "
     "challenge (step-up), while only a clearly dangerous window gets a "
     "hard block, mirroring how a human security team would actually "
     "triage."),
    ("Why Isolation Forest specifically?", "It is an unsupervised "
     "method, meaning it learns what \"normal\" looks like from "
     "legitimate traffic alone and does not need to be shown labelled "
     "attack examples to work. That matters because in the real world "
     "we will not have a large, labelled history of past attacks to "
     "train on, we will mostly have normal traffic, which is exactly the "
     "situation this type of model is built for."),
    ("Why combine rules with machine learning, instead of just one?",
     "Rules alone catch only the attacks someone thought to write a "
     "rule for, and machine learning alone can be slower to build "
     "confidence in and harder to explain. Together, the rules catch "
     "the clear-cut, explainable violations instantly, and the machine "
     "learning layer catches the subtler, statistical drift that no "
     "one thought to write an explicit rule for. Our own numbers show "
     "this matters: the ML model alone scores 0.726 on its own, versus "
     "0.992 once combined with the rules layer."),
]
add_table(doc, ["Choice", "Why we made it"], choice_rows, widths=[1.9, 4.5])

# ============================================================ PART 9
h1(doc, "What Happens After a STEP-UP or BLOCK", 9)
body(
    doc,
    "A decision from AEGIS is not the end of the story, it is the start "
    "of a response, and the response is proportional to the risk.",
)
response_rows = [
    ("PASS (low risk)", "Nothing happens, the request goes through as "
     "normal. This is the vast majority of traffic."),
    ("STEP-UP (medium risk)", "The request is not blocked outright, but "
     "the agent is challenged, for example asked to re-authenticate, "
     "requiring a secondary approval, or simply slowed down, giving a "
     "human or a secondary system a chance to double check without "
     "disrupting a legitimate agent that happened to do something "
     "slightly unusual."),
    ("BLOCK (high risk)", "The identity is quarantined and an alert is "
     "raised to a security team (a SOC, security operations centre), "
     "with the plain-English reason attached, exactly as shown in the "
     "incident inspector. A human then decides whether to fully revoke "
     "that agent's credentials, investigate further, or, if it turns "
     "out to be a false alarm, clear it and let the agent resume."),
]
add_table(doc, ["Decision", "What happens next"], response_rows, widths=[1.6, 4.8])
doc.add_paragraph()
body(
    doc,
    "This is also why the false-positive rate matters so much in this "
    "kind of system: every STEP-UP or BLOCK either interrupts a real "
    "automated process or takes up a real person's time, so a system "
    "that cries wolf too often would get switched off. Our 0.83% "
    "false-positive rate was optimised specifically with that cost in "
    "mind, not just to maximise detection at any cost.",
)

# ============================================================ PART 10
h1(doc, "The Tools We Used, in Plain Terms", 10)
tools_rows = [
    ("Python", "The programming language everything is written in, "
     "widely used, free, open-source, industry-standard."),
    ("pandas / numpy", "Standard libraries for handling large tables of "
     "data efficiently, used to store and process the 210,000 simulated "
     "requests."),
    ("scikit-learn", "A standard, well-established machine-learning "
     "library; this is where the Isolation Forest model comes from."),
    ("Streamlit", "A lightweight framework for turning a Python script "
     "into an interactive, browser-based dashboard, no separate web "
     "development needed."),
]
add_table(doc, ["Tool", "What it is / why we used it"], tools_rows, widths=[1.4, 5.0])
doc.add_paragraph()
body(
    doc,
    "Everything is free and open-source, runs entirely on a normal "
    "laptop, and needs no special hardware, cloud account, or licence. "
    "The only thing standing between this prototype and a production "
    "pilot is real data to validate against, and an integration point "
    "at the network edge (FASTedge) to act on a BLOCK decision inline, "
    "both of which are the two asks in our Tech Overview document.",
)

# ============================================================ PART 11
h1(doc, "Why This Actually Matters", 11)
body(
    doc,
    "As telecom operators automate more of the network, closed-loop "
    "healing, orchestration bots, AI copilots that can take real "
    "actions, those automations are given real credentials to real "
    "control systems. Today, security for those automations mostly "
    "stops at \"does this credential work\". But credentials get stolen, "
    "leaked, or misused by the very systems they were issued to, and a "
    "stolen credential sails straight through a login check every time. "
    "AEGIS closes that gap: it does not replace authentication, it adds "
    "a second, continuous layer that asks \"is this still really behaving "
    "like the agent we onboarded\", which is the question authentication "
    "was never designed to answer. The prototype proves that question "
    "can be answered automatically, in near real time, with high "
    "accuracy, using only ordinary request metadata, no sensitive "
    "subscriber data required.",
)

# ============================================================ PART 12
h1(doc, "The Honest Caveat", 12)
body(
    doc,
    "These strong numbers prove the detection method itself is sound and "
    "well-calibrated on the traffic we built. They do not yet prove it "
    "will hold up against the noise, scale, and unpredictability of real "
    "Fastweb or Vodafone network traffic, because it has never seen any. "
    "That is precisely the validation that even a small, anonymised real "
    "data sample from Fastweb or Vodafone would give us, and it is the "
    "central ask behind the Tech Overview document we shared separately.",
    italic=True,
)

# ============================================================ PART 13
doc.add_page_break()
h1(doc, "Questions We Expect, Answered in Advance", 13)
faq = [
    ("Is this real network traffic or made up?",
     "Fully simulated by us, on purpose, because we do not yet have "
     "access to real Fastweb/Vodafone traffic. Every request is "
     "labelled so we can grade our own detector honestly. See Part 3."),
    ("Could this same idea work on real traffic later?",
     "Yes, the detector does not care whether traffic is simulated or "
     "real, it only needs the same handful of fields (agent ID, target "
     "system, operation, timestamp, size, status), listed in the Tech "
     "Overview document. Swapping in real data does not require "
     "rebuilding the system, only retraining the baseline."),
    ("Does AEGIS replace passwords / authentication?",
     "No. It runs after authentication succeeds, as a second, "
     "continuous layer. Authentication proves the credential is valid; "
     "AEGIS proves the behaviour is still consistent with that "
     "identity."),
    ("What if a legitimate agent changes its normal behaviour on "
     "purpose (a planned update)?",
     "It would likely trigger a STEP-UP the first time, which is "
     "intentional, a real change of role should be verified once, not "
     "silently trusted. In a production rollout, the agent's baseline "
     "would then be deliberately retrained to reflect the new normal."),
    ("How long does it take to learn a new agent's baseline?",
     "In this prototype, we used five simulated days of normal "
     "behaviour to establish each agent's baseline. In production this "
     "would be tuned per agent, more predictable agents (like the NRF "
     "heartbeat service) need less time, more varied agents need more."),
    ("Does this slow down the network down?",
     "The scoring itself is lightweight, milliseconds per window, "
     "well within normal API latency budgets. The open question, which "
     "we have asked Fastweb/Vodafone directly, is whether FASTedge "
     "offers a hook to run this check inline without adding meaningful "
     "delay; that is covered in our Tech Overview document."),
    ("Is any of this using sensitive customer data?",
     "No. The detector only looks at request metadata, who, what, "
     "when, how big, what status, never subscriber identities (IMSI, "
     "MSISDN) and never message content. This was a deliberate scope "
     "decision, not an afterthought."),
    ("What is the actual next milestone?",
     "Gate 2 (22 September) is the technical solution and numerical "
     "performance write-up, building on exactly what is described in "
     "this document. Gate 3 (15 October) is the final working "
     "Proof-of-Concept demo and video."),
]
for q, a in faq:
    p = body(doc, f"Q: {q}", bold=True)
    body(doc, f"A: {a}")
    doc.add_paragraph()

doc.add_paragraph()
footer = doc.add_paragraph()
footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = footer.add_run(
    "AEGIS, Zero Trust Identity for AI Agents, Team 4, 5G Academy 2026"
)
run.font.size = Pt(9)
run.font.color.rgb = GREY

OUT.parent.mkdir(exist_ok=True)
doc.save(OUT)
print(f"Saved to {OUT}")
