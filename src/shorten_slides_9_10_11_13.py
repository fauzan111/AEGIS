# -*- coding: utf-8 -*-
"""Text-only tightening pass on slides 9, 10, 11, 13 of the CURRENT
AEGIS_Gate3_Official.pptx (as it sits on disk right now, including whatever
hand edits have been made in PowerPoint since the last script run). Changes
wording only - no font sizes, no positions, no other slides, no team names,
no alignment. Also corrects slide 13's caption, which still cited the old
7-threat / 76-96% range instead of the current 10-variant / 72-95% range.

Run:  .venv/Scripts/python AEGIS/src/shorten_slides_9_10_11_13.py
In/Out: AEGIS/slides/GATE-3/AEGIS_Gate3_Official.pptx  (edited in place)
"""

import copy
import os
from pptx import Presentation

HERE = os.path.dirname(os.path.abspath(__file__))
DECK = os.path.join(HERE, "..", "slides", "GATE-3", "AEGIS_Gate3_Official.pptx")


def find_shape_by_text(slide, snippet):
    for shp in slide.shapes:
        if shp.has_text_frame and snippet in shp.text_frame.text:
            return shp
    return None


def set_simple_text(shape, new_text):
    tf = shape.text_frame
    p0 = tf.paragraphs[0]
    if not p0.runs:
        p0.add_run()
    r0 = p0.runs[0]
    r0.text = new_text
    for r in list(p0.runs[1:]):
        r._r.getparent().remove(r._r)
    for p in list(tf.paragraphs[1:]):
        p._p.getparent().remove(p._p)


def set_bullets(shape, bullets):
    """Rewrites a '-  <prefix run> <body run>' bulleted box, reusing the
    first paragraph's formatting as the template for every new paragraph."""
    tf = shape.text_frame
    template_p = tf.paragraphs[0]._p
    parent = template_p.getparent()
    new_ps = []
    for text in bullets:
        p_copy = copy.deepcopy(template_p)
        runs_xml = p_copy.findall(".//{*}r")
        for extra in runs_xml[2:]:
            extra.getparent().remove(extra)
        prefix_r, body_r = runs_xml[0], runs_xml[1]
        body_r.find("{*}t").text = text
        new_ps.append(p_copy)
    for p in list(tf.paragraphs):
        p._p.getparent().remove(p._p)
    for p in new_ps:
        parent.append(p)


prs = Presentation(DECK)
slides = list(prs.slides)

# ------------------------------------------------------------------ slide 9: How Isolation Forest Works
slide9 = next(s for s in prs.slides if find_shape_by_text(s, "How Isolation Forest Works") is not None)
bullets9 = find_shape_by_text(slide9, "An Isolation Forest builds many random")
set_bullets(bullets9, [
    "Builds many random-split trees: each node picks a random feature and "
    "a random split value.",
    "Anomalies sit apart, so a few random splits isolate them. Normal "
    "points, surrounded by similar ones, take many more splits.",
    "Score = how short the average isolation path is, normalised to "
    "0..1. Short path -> high score -> looks like an attack.",
])
print("Shortened slide 9 bullets.")

# ------------------------------------------------------------------ slide 10: Training, Features & Calibration
slide10 = next(s for s in prs.slides if find_shape_by_text(s, "Training, Features & Calibration") is not None)
set_simple_text(find_shape_by_text(slide10, "scikit-learn IsolationForest"),
                "scikit-learn IsolationForest: 300 trees, contamination 0.02. Engages "
                "after 10+ requests in 5 minutes.")
set_simple_text(find_shape_by_text(slide10, "Scores are mapped to 0..1"),
                "Scores are mapped to 0..1 using a legit split the model never trained "
                "on - avoids understating the tail and inflating false positives.")
print("Shortened slide 10 card text.")

# ------------------------------------------------------------------ slide 11: Numerical Analysis: Headline Metrics
slide11 = next(s for s in prs.slides if find_shape_by_text(s, "Numerical Analysis: Headline Metrics") is not None)
bullets11 = find_shape_by_text(slide11, "We named a specific state of the art")
set_bullets(bullets11, [
    "We named a real state of the art and ran it ourselves: a Statistical "
    "Process Control (SPC) z-score detector, not a citation from elsewhere.",
    "SPC: 0.944 ROC-AUC / 63.4% recall. AEGIS: 0.965 / 90.7% at the same "
    "FPR - head-to-head detail on the next slide.",
    "The gap is statistically significant: paired bootstrap p=0.015, "
    "95% CI [0.004, 0.040].",
])
set_simple_text(find_shape_by_text(slide11, "These headline numbers are AEGIS's own"),
                "These numbers are AEGIS's own; the next slide shows the SPC "
                "comparison that grounds them, not just literature ranges.")
print("Shortened slide 11 bullets and footer caption.")

# ------------------------------------------------------------------ slide 13: Detection Rate Per Threat
slide13 = next(s for s in prs.slides if find_shape_by_text(s, "Detection Rate Per Threat") is not None
               and find_shape_by_text(s, "Deliberately Not Uniform") is not None)
set_simple_text(find_shape_by_text(slide13, "T4 (volumetric) and T7 (scope creep)"),
                "T1, T4 and T7 (plus their harder v2 / slow-ramp variants) sit at "
                "100% - deterministic policy or rate checks, not statistical "
                "inference. The other four span a genuine 72-95% range: loud "
                "beats subtle, and low-and-slow attacks are hardest.")
print("Shortened and corrected slide 13 caption (72-95% range, not the stale 76-96%).")

prs.save(DECK)
print(f"Saved {DECK}  ({len(prs.slides._sldIdLst)} slides)")
