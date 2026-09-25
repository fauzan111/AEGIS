# -*- coding: utf-8 -*-
"""
AEGIS - inserts real existing images into specific panels of the ALREADY
FILLED-IN storyboard (slides/GATE-3/AEGIS_Gate3_Storyboard.pptx), matching
last year's example pattern (a picture placed directly over each sketch
panel box, sized to fit without distortion).

Operates on the EXISTING file in place - does NOT regenerate from the
template, so the manually-fixed team slide (page 9) and any other manual
edits are preserved untouched. Only inserts images into named (page,
panel_index) pairs; every other panel is left as-is.

Run:  .venv/Scripts/python AEGIS/src/add_storyboard_images.py
In/Out: AEGIS/slides/GATE-3/AEGIS_Gate3_Storyboard.pptx (modified in place)
"""

import os
from pptx import Presentation
from pptx.util import Emu
from PIL import Image

HERE = os.path.dirname(os.path.abspath(__file__))
DECK = os.path.join(HERE, "..", "slides", "GATE-3", "AEGIS_Gate3_Storyboard.pptx")
ASSETS = os.path.join(HERE, "..", "assets")
REPORTS = os.path.join(HERE, "..", "reports")

PANEL_LEFTS = [237490, 3729990, 7235190]
PANEL_TOPS = [696595, 3935095]
PANEL_W = 3187700
PANEL_H = 2070100


def panel_box(index):
    """0-indexed panel position (0-5, reading order) -> (left, top)."""
    row, col = divmod(index, 3)
    return PANEL_LEFTS[col], PANEL_TOPS[row]


def add_picture_scaled(slide, image_path, left, top, max_w, max_h):
    with Image.open(image_path) as im:
        iw, ih = im.size
    scale = min(max_w / iw, max_h / ih)
    w, h = int(iw * scale), int(ih * scale)
    l = left + (max_w - w) // 2
    t = top + (max_h - h) // 2
    return slide.shapes.add_picture(image_path, Emu(l), Emu(t), width=Emu(w), height=Emu(h))


# (page_number, panel_index 0-5) -> image path. Page numbers match the
# storyboard's own pages (1-8); panel_index is reading order (row1: 0,1,2 /
# row2: 3,4,5), matching the same order the captions were written in.
IMAGE_PLACEMENTS = {
    (1, 2): os.path.join(ASSETS, "architecture.png"),                        # laptop/5G diagram reveal
    (1, 3): os.path.join(REPORTS, "GATE-2", "aegis_hook_analogy.png"),       # agent connects to network
    (2, 1): os.path.join(ASSETS, "architecture.png"),                        # architecture teaser
    (3, 1): os.path.join(REPORTS, "GATE-2", "aegis_hook_analogy.png"),      # valid credential vs abnormal behaviour
    (4, 0): os.path.join(ASSETS, "architecture.png"),                        # architecture full reveal
    (5, 0): os.path.join(ASSETS, "dashboard.jpg"),                           # dashboard home / live overview
    (5, 5): os.path.join(REPORTS, "aegis_eval.png"),                         # before/after-style detection chart
    (6, 5): os.path.join(REPORTS, "aegis_eval.png"),                         # stat card / results overlay
}


def main():
    prs = Presentation(DECK)
    inserted, missing = [], []

    for (page_num, panel_index), image_path in IMAGE_PLACEMENTS.items():
        if not os.path.exists(image_path):
            missing.append((page_num, panel_index, image_path))
            continue
        slide = prs.slides[page_num - 1]
        left, top = panel_box(panel_index)
        add_picture_scaled(slide, image_path, left, top, PANEL_W, PANEL_H)
        inserted.append((page_num, panel_index, os.path.basename(image_path)))

    prs.save(DECK)
    print(f"Saved {DECK}")
    print(f"Inserted {len(inserted)} images:")
    for p, i, name in inserted:
        print(f"  page {p}, panel {i}: {name}")
    if missing:
        print(f"Missing {len(missing)} source files (skipped):")
        for p, i, path in missing:
            print(f"  page {p}, panel {i}: {path}")


if __name__ == "__main__":
    main()
