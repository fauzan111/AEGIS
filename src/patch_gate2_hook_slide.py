# -*- coding: utf-8 -*-
"""AEGIS - inserts ONE new slide into AEGIS_Gate2_Official.pptx, right after
the Index slide: a plain-English "hook" slide for non-technical audience
members (HR, business stakeholders) before the deck goes technical. Uses the
office-badge / security-guard analogy, with a generated comparison graphic,
and previews the PASS/STEP-UP/BLOCK vocabulary used later in the deck.

Does not touch any other existing slide.

Run:  .venv/Scripts/python AEGIS/src/patch_gate2_hook_slide.py
"""

import os
from pptx import Presentation
from pptx.util import Pt, Inches, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from PIL import Image

HERE = os.path.dirname(os.path.abspath(__file__))
SLIDES_DIR = os.path.join(HERE, "..", "slides", "GATE-2")
REPORTS_G2 = os.path.join(HERE, "..", "reports", "GATE-2")
DECK = os.path.join(SLIDES_DIR, "AEGIS_Gate2_Official.pptx")
HOOK_IMG = os.path.join(REPORTS_G2, "aegis_hook_analogy.png")

NAVY = RGBColor(0x02, 0x0C, 0x31)
ACCENT = RGBColor(0x01, 0xAB, 0xDB)
MUTE_C = RGBColor(0x6B, 0x76, 0x88)
BODY_C = RGBColor(0x33, 0x41, 0x55)
PANEL_C = RGBColor(0xF2, 0xF7, 0xFA)
GREEN_C = RGBColor(0x0E, 0x93, 0x84)
AMBER_C = RGBColor(0xE0, 0xA1, 0x00)
RED_C = RGBColor(0xC0, 0x39, 0x4B)
FONT = "Montserrat"


def rect(slide, x, y, w, h, fill=None, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=None):
    sp = slide.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    if fill is None:
        sp.fill.background()
    else:
        sp.fill.solid(); sp.fill.fore_color.rgb = fill
    sp.line.fill.background()
    sp.shadow.inherit = False
    if radius is not None and shape == MSO_SHAPE.ROUNDED_RECTANGLE:
        try:
            sp.adjustments[0] = radius
        except Exception:
            pass
    return sp


def txt(slide, x, y, w, h, runs, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, line_spacing=1.1):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame; tf.word_wrap = True; tf.vertical_anchor = anchor
    tf.margin_left = 0; tf.margin_right = 0; tf.margin_top = 0; tf.margin_bottom = 0
    for i, para in enumerate(runs):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align; p.line_spacing = line_spacing
        for (t, sz, col, bold) in para:
            r = p.add_run(); r.text = t
            r.font.size = Pt(sz); r.font.name = FONT
            r.font.color.rgb = col; r.font.bold = bold
    return tb


def add_picture_scaled(slide, image_path, left, top, max_w, max_h):
    with Image.open(image_path) as im:
        iw, ih = im.size
    scale = min(max_w / iw, max_h / ih)
    w, h = int(iw * scale), int(ih * scale)
    l = left + (max_w - w) // 2
    t = top + (max_h - h) // 2
    return slide.shapes.add_picture(image_path, l, t, width=w, height=h)


def find_shape_by_text(slide, snippet):
    for shp in slide.shapes:
        if shp.has_text_frame and snippet in shp.text_frame.text:
            return shp
    return None


def move_slide(prs, old_index, new_index):
    xml_slides = prs.slides._sldIdLst
    slides = list(xml_slides)
    el = slides[old_index]
    xml_slides.remove(el)
    xml_slides.insert(new_index, el)


def pill(slide, x, bar_color, label, caption):
    w, h = 3.85, 1.0
    rect(slide, x, 4.5, w, h, fill=PANEL_C)
    rect(slide, x, 4.5, 0.09, h, fill=bar_color, shape=MSO_SHAPE.RECTANGLE)
    txt(slide, x + 0.25, 4.65, 3.4, 0.35, [[(label, 13, bar_color, True)]])
    txt(slide, x + 0.25, 5.0, 3.4, 0.4, [[(caption, 10, BODY_C, False)]])


def main():
    prs = Presentation(DECK)

    idx_index_slide = next(i for i, s in enumerate(prs.slides)
                            if find_shape_by_text(s, "Index") is not None)
    # "Vuota" is the layout every regular content slide uses (it carries the
    # recurring "5G Academy" header text + the two corner logo images as
    # layout-level background shapes) - the Index slide itself uses a
    # separate branding-free "Blank" layout, so we must not reuse its layout.
    content_layout = next(L for m in prs.slide_masters for L in m.slide_layouts
                           if L.name == "Vuota")

    hook = prs.slides.add_slide(content_layout)
    for shp in list(hook.placeholders):
        shp._element.getparent().remove(shp._element)

    txt(hook, 0.83, 0.14, 10.5, 0.3, [[("BEFORE WE GO TECHNICAL", 11, ACCENT, True)]])
    txt(hook, 0.83, 0.45, 11.6, 1.0,
        [[("A Stolen Badge Still Opens The Door.", 28, NAVY, True)]])
    txt(hook, 0.83, 1.35, 11.6, 0.55,
        [[("A good guard doesn't just check the badge - they notice when the person "
           "wearing it stops walking, working, and behaving like they always have. "
           "That's what AEGIS does for every AI agent talking to your network.",
           12.5, MUTE_C, False)]])

    add_picture_scaled(hook, HOOK_IMG, left=Emu(758952), top=Emu(1783080),
                        max_w=Emu(10607040), max_h=Emu(2423160))

    pill(hook, 0.83, GREEN_C, "PASS", "Acts like itself")
    pill(hook, 4.83, AMBER_C, "STEP-UP", "Something's off - verify")
    pill(hook, 8.83, RED_C, "BLOCK", "Impersonator - stop it")

    txt(hook, 0.83, 5.75, 11.6, 0.8,
        [[("One idea to carry through this whole talk: ", 15, NAVY, True),
          ("a badge proves who you're supposed to be. AEGIS proves whether you're "
           "still acting like them.", 15, NAVY, False)]],
        align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.TOP)

    move_slide(prs, len(prs.slides) - 1, idx_index_slide + 1)

    prs.save(DECK)
    print(f"Saved {DECK}")
    print(f"Total slides now: {len(prs.slides)}")
    print(f"Hook slide inserted at position: {idx_index_slide + 2}")


if __name__ == "__main__":
    main()
