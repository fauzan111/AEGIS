# -*- coding: utf-8 -*-
"""AEGIS - enhances the already-split Gate 2 "Proposed Technical Solution"
(pipeline) slide in place: adds a flat icon to each of the 4 phase cards, a
small ">" connector between them to read as a left-to-right flow, and a dark
"worked example" strip below the cards (column-aligned with them) that walks
one concrete request through OBSERVE -> FINGERPRINT -> SCORE -> DECIDE,
ending in a STEP-UP outcome - filling the empty space below the cards with
real content instead of just stretching them.

Only rebuilds the shapes below the pull-quote (top >= 2,600,000 EMU) on this
one slide; the kicker, title, intro line, and quote are untouched, as is
every other slide in the deck.

Run:  .venv/Scripts/python AEGIS/src/patch_gate2_pipeline_visuals.py
"""

import os
from pptx import Presentation
from pptx.util import Pt, Inches, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from PIL import Image

HERE = os.path.dirname(os.path.abspath(__file__))
SLIDES_DIR = os.path.join(HERE, "..", "slides", "GATE-2")
REPORTS_G2 = os.path.join(HERE, "..", "reports", "GATE-2")
DECK = os.path.join(SLIDES_DIR, "AEGIS_Gate2_Official.pptx")

NAVY = RGBColor(0x02, 0x0C, 0x31)
ACCENT = RGBColor(0x01, 0xAB, 0xDB)
ACCENT2 = RGBColor(0x3D, 0x67, 0xB1)
AMBER_C = RGBColor(0xE0, 0xA1, 0x00)
GREEN_C = RGBColor(0x0E, 0x93, 0x84)
BODY_C = RGBColor(0x33, 0x41, 0x55)
MUTE_C = RGBColor(0x6B, 0x76, 0x88)
LIGHTBLUE = RGBColor(0x6F, 0xCC, 0xDD)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
FONT = "Montserrat"


def rect(slide, x, y, w, h, fill=None, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=None):
    sp = slide.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    if fill is None:
        sp.fill.background()
    else:
        sp.fill.solid(); sp.fill.fore_color.rgb = fill
    sp.line.fill.background()
    sp.shadow.inherit = False
    if radius is not None and shape == MSO_SHAPE.ROUNDED_RECTANGLE:
        try:
            sp.adjustments[0] = radius
        except Exception:
            pass
    return sp


def txt(slide, x, y, w, h, runs, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, line_spacing=1.1):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame; tf.word_wrap = True; tf.vertical_anchor = anchor
    tf.margin_left = 0; tf.margin_right = 0; tf.margin_top = 0; tf.margin_bottom = 0
    for i, para in enumerate(runs):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align; p.line_spacing = line_spacing
        for (t, sz, col, bold) in para:
            r = p.add_run(); r.text = t
            r.font.size = Pt(sz); r.font.name = FONT
            r.font.color.rgb = col; r.font.bold = bold
    return tb


def add_picture_scaled(slide, image_path, left, top, max_w, max_h):
    with Image.open(image_path) as im:
        iw, ih = im.size
    scale = min(max_w / iw, max_h / ih)
    w, h = int(iw * scale), int(ih * scale)
    l = left + (max_w - w) // 2
    t = top + (max_h - h) // 2
    return slide.shapes.add_picture(image_path, l, t, width=w, height=h)


def find_shape_by_text(slide, snippet):
    for shp in slide.shapes:
        if shp.has_text_frame and snippet in shp.text_frame.text:
            return shp
    return None


def connector(slide, x_center_emu, y_center_emu, color):
    w = 200000
    txt(slide, (x_center_emu - w / 2) / 914400, (y_center_emu - 130000) / 914400,
        w / 914400, 260000 / 914400, [[(">", 15, color, True)]],
        align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)


COL_W = 2514600
GAP = 182880
LEFTS = [758952 + i * (COL_W + GAP) for i in range(4)]

CARD_TOP = 2697480
BAR_H = 73152
ICON_D = 500000
ICON_TOP = CARD_TOP + BAR_H + 120000
TITLE_TOP = ICON_TOP + ICON_D + 90000
TITLE_H = 300000
DESC_TOP = TITLE_TOP + TITLE_H + 60000
CARD_H = 2148840
CARD_BOTTOM = CARD_TOP + CARD_H
DESC_H = CARD_BOTTOM - DESC_TOP - 90000

PHASES = [
    ("OBSERVE", ACCENT, "icon_observe.png",
     "Every request an agent sends to the network is captured - who sent it, "
     "what it touched, when, and how."),
    ("FINGERPRINT", ACCENT2, "icon_fingerprint.png",
     "That stream becomes a behavioural profile per agent: its normal "
     "endpoints, pace, and call sequence."),
    ("SCORE", AMBER_C, "icon_score.png",
     "Live behaviour is compared to that profile in real time - fast rule "
     "checks fused with a machine-learning anomaly score."),
    ("DECIDE", GREEN_C, "icon_decide.png",
     "The score becomes an action instantly: let the request through, "
     "challenge it, or shut it down."),
]

EXAMPLE_ROW = [
    ("Agent-07 calls NRF\n40x per minute", LIGHTBLUE, 10, False),
    ("Normal endpoint mix,\nbut rate is elevated", LIGHTBLUE, 10, False),
    ("Risk score: 0.71", LIGHTBLUE, 12, True),
    ("STEP-UP:\nre-authenticate", AMBER_C, 12, True),
]

STRIP_TOP = CARD_BOTTOM + 160000
STRIP_H = 1450000
STRIP_BOTTOM = STRIP_TOP + STRIP_H


def main():
    prs = Presentation(DECK)

    pipeline_slide = next(s for s in prs.slides
                           if find_shape_by_text(s, "Proposed Technical Solution") is not None
                           and find_shape_by_text(s, "OBSERVE") is not None)

    for shp in list(pipeline_slide.shapes):
        if shp.top >= 2600000:
            shp._element.getparent().remove(shp._element)

    for i, (left, (name, color, icon, desc)) in enumerate(zip(LEFTS, PHASES)):
        x_in = left / 914400
        rect(pipeline_slide, x_in, CARD_TOP / 914400, COL_W / 914400, CARD_H / 914400,
             fill=WHITE, radius=0.1)
        rect(pipeline_slide, x_in, CARD_TOP / 914400, COL_W / 914400, BAR_H / 914400,
             fill=color)
        add_picture_scaled(pipeline_slide, os.path.join(REPORTS_G2, icon),
                            left=Emu(int(left + (COL_W - ICON_D) / 2)), top=Emu(ICON_TOP),
                            max_w=Emu(ICON_D), max_h=Emu(ICON_D))
        txt(pipeline_slide, (left + GAP) / 914400, TITLE_TOP / 914400,
            (COL_W - 2 * GAP) / 914400, TITLE_H / 914400,
            [[(name, 13, NAVY, True)]], align=PP_ALIGN.CENTER)
        txt(pipeline_slide, (left + GAP) / 914400, DESC_TOP / 914400,
            (COL_W - 2 * GAP) / 914400, DESC_H / 914400,
            [[(desc, 10, BODY_C, False)]], align=PP_ALIGN.CENTER)
        if i < 3:
            connector(pipeline_slide, left + COL_W + GAP / 2,
                      ICON_TOP + ICON_D / 2, MUTE_C)

    rect(pipeline_slide, 758952 / 914400, STRIP_TOP / 914400,
         10607040 / 914400, STRIP_H / 914400, fill=NAVY, radius=0.08)
    txt(pipeline_slide, (758952 + 228600) / 914400, (STRIP_TOP + 140000) / 914400,
        3000000 / 914400, 250000 / 914400,
        [[("WORKED EXAMPLE", 10.5, ACCENT, True)]])

    row_top = STRIP_TOP + 460000
    row_h = STRIP_BOTTOM - row_top - 140000
    for i, (left, (text, color, size, bold)) in enumerate(zip(LEFTS, EXAMPLE_ROW)):
        txt(pipeline_slide, (left + GAP) / 914400, row_top / 914400,
            (COL_W - 2 * GAP) / 914400, row_h / 914400,
            [[(line, size, color, bold)] for line in text.split("\n")],
            align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
        if i < 3:
            connector(pipeline_slide, left + COL_W + GAP / 2,
                      row_top + row_h / 2, LIGHTBLUE)

    prs.save(DECK)
    print(f"Saved {DECK}")


if __name__ == "__main__":
    main()
