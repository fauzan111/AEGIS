# -*- coding: utf-8 -*-
"""AEGIS - splits the Gate 2 "Proposed Technical Solution" slide (intro line +
pull-quote + pipeline arrow-line + all 7 T1-T7 threat cards crammed together)
into two slides, per tutor feedback (Vincenzo Merola, relaying Prof. Tulino /
Dr. Mauro): give the OBSERVE -> FINGERPRINT -> SCORE -> DECIDE pipeline its
own slide with each phase properly explained, and give the seven threats
their own slide with more room.

Both new slides are built by duplicating the original slide twice (so all
existing kicker/title/lead/quote styling is preserved byte-for-byte) and then
surgically removing/adding shapes on each copy - the same approach used by
patch_gate2_deck_evidence.py. Nothing else in the deck is touched.

Run:  .venv/Scripts/python AEGIS/src/patch_gate2_split_pipeline_threats.py
"""

import copy
import os
from pptx import Presentation
from pptx.util import Pt, Inches, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

HERE = os.path.dirname(os.path.abspath(__file__))
SLIDES_DIR = os.path.join(HERE, "..", "slides", "GATE-2")
DECK = os.path.join(SLIDES_DIR, "AEGIS_Gate2_Official.pptx")

NAVY = RGBColor(0x02, 0x0C, 0x31)
ACCENT = RGBColor(0x01, 0xAB, 0xDB)
ACCENT2 = RGBColor(0x3D, 0x67, 0xB1)
AMBER_C = RGBColor(0xE0, 0xA1, 0x00)
GREEN_C = RGBColor(0x0E, 0x93, 0x84)
BODY_C = RGBColor(0x33, 0x41, 0x55)
MUTE_C = RGBColor(0x6B, 0x76, 0x88)
FONT = "Montserrat"

# ------------------------------------------------------------------ helpers (matching existing patch scripts)
def duplicate_slide(prs, index):
    source = prs.slides[index]
    dest = prs.slides.add_slide(source.slide_layout)
    for shp in list(dest.shapes):
        shp._element.getparent().remove(shp._element)
    for shp in source.shapes:
        dest.shapes._spTree.append(copy.deepcopy(shp._element))
    return dest


def move_slide(prs, old_index, new_index):
    xml_slides = prs.slides._sldIdLst
    slides = list(xml_slides)
    el = slides[old_index]
    xml_slides.remove(el)
    xml_slides.insert(new_index, el)


def delete_slide(prs, index):
    xml_slides = prs.slides._sldIdLst
    slides = list(xml_slides)
    rId = slides[index].rId
    prs.part.drop_rel(rId)
    xml_slides.remove(slides[index])


def find_shape_by_text(slide, snippet):
    for shp in slide.shapes:
        if shp.has_text_frame and snippet in shp.text_frame.text:
            return shp
    return None


def set_simple_text(shape, new_text):
    tf = shape.text_frame
    p0 = tf.paragraphs[0]
    if not p0.runs:
        p0.add_run()
    r0 = p0.runs[0]
    r0.text = new_text
    for r in list(p0.runs[1:]):
        r._r.getparent().remove(r._r)
    for p in list(tf.paragraphs[1:]):
        p._p.getparent().remove(p._p)


def rect(slide, x, y, w, h, fill=None, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=None):
    sp = slide.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    if fill is None:
        sp.fill.background()
    else:
        sp.fill.solid(); sp.fill.fore_color.rgb = fill
    sp.line.fill.background()
    sp.shadow.inherit = False
    if radius is not None and shape == MSO_SHAPE.ROUNDED_RECTANGLE:
        try:
            sp.adjustments[0] = radius
        except Exception:
            pass
    return sp


def txt(slide, x, y, w, h, runs, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, line_spacing=1.1):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame; tf.word_wrap = True; tf.vertical_anchor = anchor
    tf.margin_left = 0; tf.margin_right = 0; tf.margin_top = 0; tf.margin_bottom = 0
    for i, para in enumerate(runs):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align; p.line_spacing = line_spacing
        for (t, sz, col, bold) in para:
            r = p.add_run(); r.text = t
            r.font.size = Pt(sz); r.font.name = FONT
            r.font.color.rgb = col; r.font.bold = bold
    return tb


def emu_top(shape):
    return shape.top


PHASES = [
    ("OBSERVE", ACCENT,
     "Every request an agent sends to the network is captured - who sent it, "
     "what it touched, when, and how."),
    ("FINGERPRINT", ACCENT2,
     "That stream becomes a behavioural profile per agent: its normal "
     "endpoints, pace, and call sequence."),
    ("SCORE", AMBER_C,
     "Live behaviour is compared to that profile in real time - fast rule "
     "checks fused with a machine-learning anomaly score."),
    ("DECIDE", GREEN_C,
     "The score becomes an action instantly: let the request through, "
     "challenge it, or shut it down."),
]


def main():
    prs = Presentation(DECK)

    idx_ptsol = next(i for i, s in enumerate(prs.slides)
                      if find_shape_by_text(s, "Proposed Technical Solution") is not None
                      and find_shape_by_text(s, "OBSERVE") is not None)

    # ---- duplicate twice from the pristine original ----------------------
    pipeline_slide = duplicate_slide(prs, idx_ptsol)
    move_slide(prs, len(prs.slides) - 1, idx_ptsol + 1)

    threats_slide = duplicate_slide(prs, idx_ptsol)
    move_slide(prs, len(prs.slides) - 1, idx_ptsol + 2)

    # ---- Pipeline slide: drop the arrow-line + T1-T7 grid (everything at
    #      or below the old pipeline-line's top), keep intro + quote -------
    for shp in list(pipeline_slide.shapes):
        if emu_top(shp) >= 2600000:
            shp._element.getparent().remove(shp._element)

    col_w = 2514600
    gap = 182880
    card_top = 2697480  # right where the old arrow-line sat, just below the quote box
    card_h = 2148840
    lefts = [758952 + i * (col_w + gap) for i in range(4)]

    for left, (name, color, desc) in zip(lefts, PHASES):
        x_in = left / 914400
        rect(pipeline_slide, x_in, card_top / 914400, col_w / 914400, card_h / 914400,
             fill=RGBColor(0xFF, 0xFF, 0xFF), radius=0.1)
        rect(pipeline_slide, x_in, card_top / 914400, col_w / 914400, 73152 / 914400,
             fill=color, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=None)
        txt(pipeline_slide, (left + 182880) / 914400, (card_top + 228600) / 914400,
            (col_w - 2 * 182880) / 914400, 320040 / 914400,
            [[(name, 13, NAVY, True)]])
        txt(pipeline_slide, (left + 182880) / 914400, (card_top + 594360) / 914400,
            (col_w - 2 * 182880) / 914400, 1463040 / 914400,
            [[(desc, 10, BODY_C, False)]])

    # ---- Threats slide: drop quote + arrow-line, retext title/lead, shift
    #      the T1-T7 grid up to fill the freed space ------------------------
    for shp in list(threats_slide.shapes):
        top = emu_top(shp)
        if 1600000 <= top < 3100000:
            shp._element.getparent().remove(shp._element)

    set_simple_text(find_shape_by_text(threats_slide, "Proposed Technical Solution"),
                     "The Seven Threats We Watch For")
    set_simple_text(find_shape_by_text(threats_slide, "A stolen credential"),
                     "Same four-stage pipeline. Seven different ways an identity can go wrong.")

    delta = 1554480
    for shp in list(threats_slide.shapes):
        if emu_top(shp) >= 3100000:
            shp.top = Emu(shp.top - delta)

    # ---- remove the original, now-split slide -----------------------------
    delete_slide(prs, idx_ptsol)

    prs.save(DECK)
    print(f"Saved {DECK}")
    print(f"Total slides now: {len(prs.slides)}")


if __name__ == "__main__":
    main()
