# -*- coding: utf-8 -*-
"""
AEGIS - replaces slide 13's 6-card text grid with the embedded live-demo
video (slides/GATE-2/AEGIS_Live_Demo.mp4), keeping the kicker/title/subtitle.

Run:  .venv/Scripts/python AEGIS/src/patch_gate2_video.py
"""

import os
from pptx import Presentation
from pptx.util import Emu

HERE = os.path.dirname(os.path.abspath(__file__))
SLIDES_G2 = os.path.join(HERE, "..", "slides", "GATE-2")
DECK = os.path.join(SLIDES_G2, "AEGIS_Gate2_Official.pptx")
VIDEO = os.path.join(SLIDES_G2, "AEGIS_Live_Demo.mp4")
POSTER = os.path.join(SLIDES_G2, "AEGIS_Live_Demo_poster.png")


def main():
    prs = Presentation(DECK)
    slide = prs.slides[12]  # slide 13, 0-indexed

    # find the subtitle (anchor - keep it) and every shape below it (the 6
    # card auto-shapes + their title/description text boxes - remove them all)
    subtitle = None
    to_remove = []
    for shp in slide.shapes:
        if shp.has_text_frame and "AEGIS has a working, demonstrable baseline" in shp.text_frame.text:
            subtitle = shp
            continue
    subtitle_bottom = subtitle.top + subtitle.height

    for shp in list(slide.shapes):
        if shp is subtitle:
            continue
        if shp.top is not None and shp.top >= subtitle_bottom - Emu(100000):
            to_remove.append(shp)

    print(f"Removing {len(to_remove)} shapes (the 6-card grid)")
    for shp in to_remove:
        shp._element.getparent().remove(shp._element)

    # video sized to the actual 1920x1020 aspect ratio, centred, filling the
    # space the card grid used to occupy
    import cv2
    cap = cv2.VideoCapture(VIDEO)
    vid_w = cap.get(cv2.CAP_PROP_FRAME_WIDTH)
    vid_h = cap.get(cv2.CAP_PROP_FRAME_HEIGHT)
    cap.release()
    aspect = vid_w / vid_h

    max_w, max_h = Emu(9906000), Emu(5100000)  # ~10.83in x 5.58in content box
    if max_w / aspect <= max_h:
        w = max_w
        h = Emu(int(max_w / aspect))
    else:
        h = max_h
        w = Emu(int(max_h * aspect))
    left = Emu(int((prs.slide_width - w) / 2))
    top = Emu(int(subtitle_bottom) + Emu(140000))

    slide.shapes.add_movie(VIDEO, left, top, w, h, poster_frame_image=POSTER,
                           mime_type="video/mp4")

    prs.save(DECK)
    print(f"Saved {DECK}")
    print(f"Video placed at L={left} T={top} W={w} H={h}")


if __name__ == "__main__":
    main()
