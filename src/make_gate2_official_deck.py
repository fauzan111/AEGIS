# -*- coding: utf-8 -*-
"""Build the official Gate 2 PowerPoint deck, reusing the exact 5G Academy
visual template (colours, fonts, cover, and closing slide) from the Gate 1
official deck, but generating the content slides fresh in code so the
30-minute presentation follows the 5 official Gate 2 headings:
2.1 Proposed Technical Solution, 2.2 Architecture, 2.3 Implementation
Strategy, 2.4 Numerical Performance Analysis, 2.5 Expected Proof of Concept,
plus a short Path-to-Gate-3 closer. Content mirrors docs/GATE-2.md.

Run:  .venv/Scripts/python AEGIS/src/make_gate2_official_deck.py
In :  AEGIS/slides/GATE-1/AEGIS_Gate1_Official.pptx  (template source)
Out:  AEGIS/slides/GATE-2/AEGIS_Gate2_Official.pptx
"""

import os
from pptx import Presentation
from pptx.util import Pt, Inches
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

# same brand colours as the Gate 1 deck, sampled from the official template
NAVY = RGBColor(0x02, 0x0C, 0x31)
ACCENT = RGBColor(0x01, 0xAB, 0xDB)
ACCENT2 = RGBColor(0x3D, 0x67, 0xB1)
LIGHTBLUE = RGBColor(0x6F, 0xCC, 0xDD)
BODY_C = RGBColor(0x33, 0x41, 0x55)
MUTE_C = RGBColor(0x6B, 0x76, 0x88)
CARD_C = RGBColor(0xFF, 0xFF, 0xFF)
HAIR_C = RGBColor(0xE2, 0xE8, 0xF0)
PANEL_C = RGBColor(0xF2, 0xF7, 0xFA)
GREEN_C = RGBColor(0x0E, 0x93, 0x84)
AMBER_C = RGBColor(0xE0, 0xA1, 0x00)
RED_C = RGBColor(0xC0, 0x39, 0x4B)
FONT = "Montserrat"

HERE = os.path.dirname(os.path.abspath(__file__))
SLIDES = os.path.join(HERE, "..", "slides")
TEMPLATE = os.path.join(SLIDES, "GATE-1", "AEGIS_Gate1_Official.pptx")
ARCH_PNG = os.path.join(HERE, "..", "reports", "aegis_architecture_diagram.png")
EVAL_PNG = os.path.join(HERE, "..", "reports", "aegis_eval.png")
ROC_PNG = os.path.join(HERE, "..", "reports", "GATE-1", "aegis_threshold_sensitivity.png")
OUT_DIR = os.path.join(SLIDES, "GATE-2")
OUT = os.path.join(OUT_DIR, "AEGIS_Gate2_Official.pptx")


# ------------------------------------------------------------------ helpers (same as Gate 1 script)
def rect(slide, x, y, w, h, fill=None, line=None, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=None):
    sp = slide.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    if fill is None:
        sp.fill.background()
    else:
        sp.fill.solid(); sp.fill.fore_color.rgb = fill
    if line is None:
        sp.line.fill.background()
    else:
        sp.line.color.rgb = line; sp.line.width = Pt(1.0)
    sp.shadow.inherit = False
    if radius is not None and shape == MSO_SHAPE.ROUNDED_RECTANGLE:
        try:
            sp.adjustments[0] = radius
        except Exception:
            pass
    return sp


def txt(slide, x, y, w, h, runs, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, line_spacing=1.1):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame; tf.word_wrap = True; tf.vertical_anchor = anchor
    tf.margin_left = 0; tf.margin_right = 0; tf.margin_top = 0; tf.margin_bottom = 0
    for i, para in enumerate(runs):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align; p.line_spacing = line_spacing
        for (t, sz, col, bold) in para:
            r = p.add_run(); r.text = t
            r.font.size = Pt(sz); r.font.name = FONT
            r.font.color.rgb = col; r.font.bold = bold
    return tb


def title(slide, text_):
    txt(slide, 0.83, 0.29, 10.5, 1.0, [[(text_, 28, NAVY, True)]])


def kicker(slide, text_):
    txt(slide, 0.83, 0.02, 10.5, 0.3, [[(text_, 11, ACCENT, True)]])


def lead(slide, text_, y=1.25):
    txt(slide, 0.83, y, 11.6, 0.5, [[(text_, 12.5, MUTE_C, False)]])


def delete_slide(prs, index):
    xml_slides = prs.slides._sldIdLst
    slides = list(xml_slides)
    rId = slides[index].rId
    prs.part.drop_rel(rId)
    xml_slides.remove(slides[index])


def by_name(slide, name):
    for shape in slide.shapes:
        if shape.name == name:
            return shape
    return None


# ------------------------------------------------------------------ load Gate 1 deck as the base
prs = Presentation(TEMPLATE)

BLANK = None
for layout in prs.slide_masters[0].slide_layouts:
    if layout.name in ("Vuota", "Blank"):
        BLANK = layout
        break


def new_slide():
    return prs.slides.add_slide(BLANK)


# ------------------------------------------------------------------ slide 1: cover (kept as-is)
# Shared 5G Academy programme cover - common to every team and every gate.

# ------------------------------------------------------------------ slide 2: index, rewritten for Gate 2
idx_slide = prs.slides[1]
heading = by_name(idx_slide, "TextBox 36")
if heading is not None:
    tf = heading.text_frame
    for i, ptext in enumerate(("AEGIS ", "- Gate 2 - Team 4")):
        if i >= len(tf.paragraphs):
            p = tf.add_paragraph()
        else:
            p = tf.paragraphs[i]
        if not p.runs:
            p.add_run()
        p.runs[0].text = ptext
        for extra in p.runs[1:]:
            extra.text = ""

# the template's own team-list box uses an en dash; replace it to keep the
# deck dash-free
for shape in idx_slide.shapes:
    if shape.has_text_frame and "Fauzan" in shape.text_frame.text:
        for p in shape.text_frame.paragraphs:
            for r in p.runs:
                if "–" in r.text or "—" in r.text:
                    r.text = r.text.replace("–", ":").replace("—", ":")

# remove every existing numbered topic textbox (Gate 1 had 8-9 of them) so we
# can lay out Gate 2's 6 topics fresh at the same proven coordinates
for shape in list(idx_slide.shapes):
    if shape.has_text_frame and shape.name.startswith("TextBox"):
        t = shape.text_frame.text.strip()
        is_number = t.isdigit() and len(t) == 2
        is_old_topic = t in (
            "What We're Building", "The 5G Network, In Plain Terms",
            "The Agents We Modelled", "How We Built The Prototype",
            "Technologies & Synthetic Data", "The Attacks We Simulated",
            "How AEGIS Decides", "Proof It Works",
            "Project Plan: Schedule, Feasibility & Business Case",
        )
        if is_number or is_old_topic:
            shape._element.getparent().remove(shape._element)

GATE2_TOPICS = [
    "Proposed Technical Solution",
    "Architecture",
    "Implementation Strategy",
    "Numerical Performance Analysis",
    "Expected Proof of Concept",
    "Path to Gate 3",
]
top0, step, num_w, num_l, txt_l, txt_w = 1150000, 600000, 460000, 6130000, 6700000, 5900000
for i, topicname in enumerate(GATE2_TOPICS):
    y = top0 + i * step
    numbox = idx_slide.shapes.add_textbox(num_l, y, num_w, 520000)
    tf = numbox.text_frame; tf.word_wrap = True
    r = tf.paragraphs[0].add_run(); r.text = f"{i+1:02d}"
    r.font.size = Pt(15); r.font.bold = True; r.font.name = FONT; r.font.color.rgb = ACCENT
    txtbox = idx_slide.shapes.add_textbox(txt_l, y, txt_w, 520000)
    tf = txtbox.text_frame; tf.word_wrap = True; tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    r = tf.paragraphs[0].add_run(); r.text = topicname
    r.font.size = Pt(14); r.font.bold = True; r.font.name = FONT; r.font.color.rgb = NAVY

# ------------------------------------------------------------------ delete all Gate-1-specific content slides
# Keep index 0 (cover), index 1 (index, just rewritten), and the last slide
# (thank you / closing). Delete everything in between (original slides 3-16
# of the Gate 1 deck: quote/intro, 5G network, agents, build process, tech &
# data, attacks, decision logic, results, roadmap, schedule table/picture,
# risk analysis, open points, next steps).
last_index = len(prs.slides) - 1
for i in range(last_index - 1, 1, -1):
    delete_slide(prs, i)

# ============================================================ PART 1/6 - Proposed Technical Solution
sl = new_slide()
kicker(sl, "PART 1 OF 6 - GATE 2: TECHNICAL SOLUTION & ARCHITECTURE")
title(sl, "Proposed Technical Solution")
lead(sl, "A stolen credential still passes authentication. AEGIS answers the question "
        "authentication never asks: is this really our agent, behaving the way it always does?")
rect(sl, 0.83, 1.85, 11.6, 0.85, fill=NAVY, radius=0.1)
txt(sl, 1.15, 2.05, 11.0, 0.5, [[
    ("“Credentials prove what you have, AEGIS verifies how you behave.”", 15, RGBColor(0xFF, 0xFF, 0xFF), True)]],
    align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
txt(sl, 0.83, 2.95, 11.6, 0.4, [[
    ("OBSERVE  ->  FINGERPRINT  ->  SCORE  ->  DECIDE   (zero-trust gate dashboard)", 13, ACCENT2, True)]],
    align=PP_ALIGN.CENTER)
threats = [
    ("T1", "Identity spoofing", "wrong endpoint mix, timing, sequence for that identity"),
    ("T2", "Compromised agent", "drift from its own baseline: new endpoints, elevated errors"),
    ("T3", "Reconnaissance", "high cardinality of endpoints/targets, many 403/404s"),
    ("T4", "Volumetric abuse", "request-rate spike far above the agent's normal envelope"),
    ("T5", "Slow exfiltration", "abnormal payload sizes, off-hours persistence"),
    ("T6", "Sequence anomaly", "broken call-sequence pattern vs the agent's normal graph"),
    ("T7", "Scope creep", "calls to NFs never in the agent's baseline scope"),
]
tw, th_, tgx, tgy = 3.75, 1.05, 0.2, 0.15
for i, (tid, name, desc) in enumerate(threats):
    r, c = divmod(i, 4)
    x = 0.83 + c * (tw + tgx); y = 3.55 + r * (th_ + tgy)
    rect(sl, x, y, tw, th_, fill=CARD_C, line=HAIR_C, radius=0.09)
    rect(sl, x, y, 0.5, th_, fill=RGBColor(0xFD, 0xE8, 0xEA))
    txt(sl, x, y, 0.5, th_, [[(tid, 13, RED_C, True)]], align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    txt(sl, x + 0.62, y + 0.1, tw - 0.75, 0.3, [[(name, 12, NAVY, True)]])
    txt(sl, x + 0.62, y + 0.45, tw - 0.75, 0.55, [[(desc, 9.5, BODY_C, False)]])

# ============================================================ PART 2/6 - Architecture (diagram)
sl = new_slide()
kicker(sl, "PART 2 OF 6 - ARCHITECTURE")
title(sl, "System Architecture")
lead(sl, "A four-stage pipeline sitting in front of the network's control-plane APIs. "
        "Every request an agent sends to a Network Function is observed by the gate.")
if os.path.exists(ARCH_PNG):
    from PIL import Image as PILImage
    im = PILImage.open(ARCH_PNG)
    ar = im.size[0] / im.size[1]
    pic_h = 5.3
    pic_w = pic_h * ar
    sl.shapes.add_picture(ARCH_PNG, Inches(0.9), Inches(1.85), height=Inches(pic_h))
    notes_x = 0.9 + pic_w + 0.4
else:
    notes_x = 0.9
notes_w = max(11.6 - (notes_x - 0.83), 3.0)
txt(sl, notes_x, 1.9, notes_w, 5.2, [
    [("OBSERVE", 12, ACCENT2, True)],
    [("Parse request + metadata; assemble a per-agent stream.", 10, BODY_C, False)],
    [("", 6, BODY_C, False)],
    [("FINGERPRINT", 12, ACCENT2, True)],
    [("Windowed feature vector: volume, surface, method mix, sequence, "
      "errors, payload, temporal, cross-window breadth.", 10, BODY_C, False)],
    [("", 6, BODY_C, False)],
    [("SCORE", 12, ACCENT2, True)],
    [("Rules (scope, session state, rolling accumulators) + an Isolation "
      "Forest, fused: risk = max(rule, ml).", 10, BODY_C, False)],
    [("", 6, BODY_C, False)],
    [("DECIDE", 12, ACCENT2, True)],
    [("Thresholds map risk to PASS / STEP-UP / BLOCK, automatically.", 10, BODY_C, False)],
], line_spacing=1.2)

# ============================================================ PART 2/6 (cont) - Scoring & Decision
sl = new_slide()
kicker(sl, "PART 2 OF 6 - ARCHITECTURE")
title(sl, "Scoring: Rules + ML, Fused")
score_cards = [
    ("RULES LAYER", "Fast, explainable. Hard caps and forbidden transitions: an out-of-scope "
     "NF call is flagged immediately. Includes a session-state machine and rolling "
     "5-minute accumulators for breadth, errors, and peak payload.", ACCENT),
    ("ML LAYER", "An Isolation Forest, trained only on legitimate traffic - no labelled "
     "attacks needed. Engages once an agent has at least 10 requests in the trailing "
     "5 minutes.", ACCENT2),
    ("FUSION", "risk = max(rule_score, ml_score). The rules layer has veto power on hard "
     "violations; the ML layer catches subtler drift the rules never anticipated.", GREEN_C),
]
cw = 3.78
for i, (h, d, c) in enumerate(score_cards):
    x = 0.83 + i * (cw + 0.15)
    rect(sl, x, 1.85, cw, 2.35, fill=CARD_C, line=HAIR_C, radius=0.07)
    rect(sl, x, 1.85, cw, 0.08, fill=c)
    txt(sl, x + 0.2, 2.1, cw - 0.4, 0.35, [[(h, 13, NAVY, True)]])
    txt(sl, x + 0.2, 2.5, cw - 0.4, 1.6, [[(d, 10, BODY_C, False)]], line_spacing=1.2)
legend = [("PASS", GREEN_C, "Behaves like its known-good self"),
          ("STEP-UP", AMBER_C, "Challenge: re-authenticate, approve, throttle"),
          ("BLOCK", RED_C, "Quarantine the identity, alert the SOC")]
for i, (lab, c, d) in enumerate(legend):
    x = 0.83 + i * (3.85 + 0.15)
    rect(sl, x, 4.5, 3.85, 1.0, fill=PANEL_C, radius=0.1)
    rect(sl, x, 4.5, 0.09, 1.0, fill=c)
    txt(sl, x + 0.25, 4.65, 3.4, 0.35, [[(lab, 13, c, True)]])
    txt(sl, x + 0.25, 5.0, 3.4, 0.4, [[(d, 10, BODY_C, False)]])
txt(sl, 0.83, 5.75, 11.6, 0.6, [[
    ("Operating point: thresholds are tuned for a realistic ~1.5% false-positive budget "
     "rather than maximum sensitivity, so detection is intentionally not uniform across "
     "threats (see Part 4).", 11, MUTE_C, False)]])

# ============================================================ PART 3/6 - Implementation Strategy: Tools & Phases
sl = new_slide()
kicker(sl, "PART 3 OF 6 - IMPLEMENTATION STRATEGY")
title(sl, "Tools, Platforms & Build Phases")
lead(sl, "Built bottom-up, PoC-first, and evaluation-driven: the ML/security equivalent "
        "of test-driven development.")
txt(sl, 0.83, 1.85, 5.6, 0.35, [[("TOOLS & PLATFORMS", 12.5, ACCENT2, True)]])
tools = [
    ("Python 3.12 + pandas/NumPy", "Traffic generation, windowing, feature engineering."),
    ("scikit-learn (Isolation Forest)", "Unsupervised ML layer, trained on legit traffic only."),
    ("Streamlit + Matplotlib", "The live zero-trust gate dashboard and evaluation charts."),
    ("python-docx / python-pptx", "Reproducible generation of every Gate report and deck."),
]
ty, th_ = 2.25, 0.72
for i, (name, desc) in enumerate(tools):
    y = ty + i * (th_ + 0.1)
    rect(sl, 0.83, y, 5.6, th_, fill=CARD_C, line=HAIR_C, radius=0.1)
    rect(sl, 0.83, y, 0.08, th_, fill=ACCENT2)
    txt(sl, 1.1, y + 0.07, 5.1, 0.3, [[(name, 11.5, NAVY, True)]])
    txt(sl, 1.1, y + 0.37, 5.1, 0.3, [[(desc, 9.5, BODY_C, False)]])
txt(sl, 6.75, 1.85, 5.7, 0.35, [[("SEVEN BUILD PHASES", 12.5, ACCENT, True)]])
phases = [
    "1. Synthetic environment - 210,000 labelled requests, 5 agents, T1-T7.",
    "2. Fingerprinting - windowed behavioural feature vectors.",
    "3. Rules layer - deterministic scope/session/rolling checks.",
    "4. ML layer - Isolation Forest fit on legit-only traffic.",
    "5. Fusion & decision policy - PASS / STEP-UP / BLOCK.",
    "6. Dashboard - live zero-trust gate UI.",
    "7. Evaluation & iteration - re-score after every change.",
]
runs = [[(p, 11, BODY_C, False)] for p in phases]
txt(sl, 6.75, 2.25, 5.7, 4.0, runs, line_spacing=1.35)

# ============================================================ PART 3/6 (cont) - What We Found & Fixed
sl = new_slide()
kicker(sl, "PART 3 OF 6 - IMPLEMENTATION STRATEGY")
title(sl, "What The Evaluation Loop Found & Fixed")
lead(sl, "Real engineering evidence, not just tuning cosmetics - two concrete bugs were "
        "found and fixed by re-scoring after every change.")
findings = [
    ("Bug 1: T6 window-count heuristic false-fired", "An earlier heuristic flagged legitimate "
     "bursty traffic as sequence anomalies. Replaced with a proper session-ID state machine "
     "that tracks PDU-session lifecycle."),
    ("Bug 2: T2 and T5 attacks were trivially caught by the wrong rule", "Both initially had "
     "the agent call a Network Function outside its own scope - indistinguishable from T7 to "
     "the rules layer. Redesigned both to stay within the agent's own scope, a more realistic "
     "model, which forced detection to rely on the intended statistical signal."),
]
fy, fh = 2.0, 2.0
for i, (h, d) in enumerate(findings):
    y = fy + i * (fh + 0.2)
    rect(sl, 0.83, y, 11.6, fh, fill=CARD_C, line=HAIR_C, radius=0.06)
    rect(sl, 0.83, y, 0.09, fh, fill=RED_C if i == 0 else AMBER_C)
    txt(sl, 1.15, y + 0.18, 10.8, 0.4, [[(h, 13.5, NAVY, True)]])
    txt(sl, 1.15, y + 0.65, 10.8, 1.2, [[(d, 11, BODY_C, False)]], line_spacing=1.25)
txt(sl, 0.83, 6.3, 11.6, 0.5, [[
    ("This is the origin of the non-uniform per-threat detection numbers in Part 4 - a "
     "feature, not a flaw.", 11, MUTE_C, True)]])

# ============================================================ PART 4/6 - Numerical Analysis: Benchmarking
sl = new_slide()
kicker(sl, "PART 4 OF 6 - NUMERICAL PERFORMANCE ANALYSIS")
title(sl, "Benchmarked Against Real Detection Systems")
lead(sl, "A PoC that reports 100% detection on everything is a red flag, not a strength. "
        "We researched how comparable systems perform in practice.")
research = [
    "Published NIDS studies show volumetric/DoS attacks detected at 97-99.8% - the "
    "easiest signal - while reconnaissance and other rare classes detect far lower.",
    "Industry SOC data (Microsoft/Omdia, SANS 2025) puts typical false-positive rates "
    "at 46-83%; even elite, well-tuned SOCs sit under 10%.",
    "UEBA research shows a hard precision/recall tradeoff: near-perfect recall models "
    "sacrifice precision to as low as 0.54. Low-and-slow exfiltration is consistently "
    "the hardest category.",
]
runs = [[("-  ", 11.5, ACCENT, True), (r, 11.5, BODY_C, False)] for r in research]
txt(sl, 0.83, 1.85, 11.6, 1.8, runs, line_spacing=1.3)
metrics = [("ROC-AUC", "0.965"), ("Recall", "90.7%"), ("False positives", "1.4%"), ("Precision", "0.67")]
for i, (lab, val) in enumerate(metrics):
    x = 0.83 + i * (2.85 + 0.15)
    rect(sl, x, 4.0, 2.85, 1.3, fill=PANEL_C, radius=0.12)
    txt(sl, x, 4.2, 2.85, 0.55, [[(val, 26, ACCENT2, True)]], align=PP_ALIGN.CENTER)
    txt(sl, x, 4.8, 2.85, 0.35, [[(lab.upper(), 10, MUTE_C, True)]], align=PP_ALIGN.CENTER)
txt(sl, 0.83, 5.55, 11.6, 0.9, [[
    ("These sit comfortably inside the literature above: FPR under 2% is elite-SOC "
     "territory, and precision of 0.67 is in line with (in fact better than) published "
     "UEBA results that trade precision for recall.", 11, BODY_C, False)]], line_spacing=1.25)

# ============================================================ PART 4/6 (cont) - Per-threat detection
sl = new_slide()
kicker(sl, "PART 4 OF 6 - NUMERICAL PERFORMANCE ANALYSIS")
title(sl, "Detection Rate Per Threat - Deliberately Not Uniform")
if os.path.exists(EVAL_PNG):
    from PIL import Image as PILImage
    im = PILImage.open(EVAL_PNG)
    ar = im.size[0] / im.size[1]
    pic_w = 11.6
    pic_h = pic_w / ar
    sl.shapes.add_picture(EVAL_PNG, Inches(0.83), Inches(1.5), width=Inches(pic_w))
    next_y = 1.5 + pic_h + 0.25
else:
    next_y = 1.9
txt(sl, 0.83, next_y, 11.6, 1.0, [[
    ("T4 (volumetric) and T7 (scope creep) sit at 100% because they are deterministic "
     "policy/rate checks, not statistical inference. The other five span a genuine 76-96% "
     "range that tracks the strength of their underlying behavioural signal: loud beats "
     "subtle, and diffuse, low-and-slow attacks are always the hardest.", 11.5, BODY_C, False)]],
    line_spacing=1.25)

# ============================================================ PART 4/6 (cont) - ROC operating point
sl = new_slide()
kicker(sl, "PART 4 OF 6 - NUMERICAL PERFORMANCE ANALYSIS")
title(sl, "ROC Operating-Point Analysis")
lead(sl, "The chosen threshold is justified, not just picked - a full sensitivity sweep "
        "against the statistically optimal point.")
if os.path.exists(ROC_PNG):
    from PIL import Image as PILImage
    im = PILImage.open(ROC_PNG)
    ar = im.size[0] / im.size[1]
    pic_w = 11.0
    pic_h = pic_w / ar
    sl.shapes.add_picture(ROC_PNG, Inches(1.15), Inches(1.9), width=Inches(pic_w))
    next_y = 1.9 + pic_h + 0.2
else:
    next_y = 2.2
txt(sl, 0.83, next_y, 11.6, 0.8, [[
    ("Youden's J optimum (threshold ~0.01) would catch 94% of attacks at 3.2% false "
     "positives. Our chosen threshold (0.6) trades a little recall for 1.4% false "
     "positives - a deliberate, defensible choice, not an accident.", 11, BODY_C, False)]],
    line_spacing=1.25)

# ============================================================ PART 5/6 - Expected PoC: Live demo script
sl = new_slide()
kicker(sl, "PART 5 OF 6 - EXPECTED PROOF OF CONCEPT")
title(sl, "Live Demo Script")
lead(sl, "AEGIS already has a working, demonstrable baseline - this is not a forward-looking "
        "description, it's a walkthrough of what exists today.")
steps = [
    ("1. Agent overview", "All 5 agents, live risk scores, all sitting in PASS."),
    ("2. Trigger a threat", "Select a scored window; watch the dashboard flip to BLOCK live."),
    ("3. Incident inspector", "Plain-language why: which feature tripped, and which layer caught it."),
    ("4. Cycle through difficulty", "A spoofing case, a quiet drift case, a low-and-slow case."),
    ("5. Show the numbers", "ROC-AUC, recall, false-positive rate, per-threat table, live."),
    ("6. Make zero-trust concrete", "A valid credential, still blocked on behaviour."),
]
sw, sh_ = 3.78, 1.55
for i, (h, d) in enumerate(steps):
    r, c = divmod(i, 3)
    x = 0.83 + c * (sw + 0.15); y = 1.95 + r * (sh_ + 0.15)
    rect(sl, x, y, sw, sh_, fill=CARD_C, line=HAIR_C, radius=0.07)
    rect(sl, x, y, sw, 0.07, fill=ACCENT)
    txt(sl, x + 0.18, y + 0.16, sw - 0.36, 0.35, [[(h, 12, NAVY, True)]])
    txt(sl, x + 0.18, y + 0.55, sw - 0.36, 0.9, [[(d, 10, BODY_C, False)]], line_spacing=1.2)

# ============================================================ PART 5/6 (cont) - PoC readiness
sl = new_slide()
kicker(sl, "PART 5 OF 6 - EXPECTED PROOF OF CONCEPT")
title(sl, "PoC Readiness at Gate 2")
ready = [
    ("Synthetic traffic generator (5 agents, T1-T7)", True),
    ("Fingerprinting, rules, ML, fusion, decision policy", True),
    ("Live dashboard with incident inspector", True),
    ("Numerical evaluation (ROC-AUC, recall, precision, FPR)", True),
    ("ROC operating-point / threshold sensitivity analysis", True),
    ("Real M2M / API-gateway traffic", False),
]
ry, rh = 1.9, 0.68
for i, (label, done) in enumerate(ready):
    y = ry + i * (rh + 0.08)
    rect(sl, 0.83, y, 11.6, rh, fill=CARD_C, line=HAIR_C, radius=0.08)
    c = GREEN_C if done else AMBER_C
    stat = "DONE" if done else "NOT YET"
    rect(sl, 0.83, y, 1.5, rh, fill=c, radius=0.08)
    txt(sl, 0.83, y, 1.5, rh, [[(stat, 11, RGBColor(0xFF, 0xFF, 0xFF), True)]],
        align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    txt(sl, 2.55, y, 9.6, rh, [[(label, 11.5, NAVY, False)]], anchor=MSO_ANCHOR.MIDDLE)
txt(sl, 0.83, 6.55, 11.6, 0.6, [[
    ("Real traffic access is the one open item outside our control - a direct ask for "
     "Fastweb and Vodafone.", 10.5, MUTE_C, True)]])

# ============================================================ PART 6/6 - Path to Gate 3
sl = new_slide()
kicker(sl, "PART 6 OF 6 - PATH TO GATE 3")
title(sl, "What's Next")
next_items = [
    "Close the gap on the hardest threats (T3 recon 80%, T5 exfiltration 76%) with "
    "richer features, such as per-target-ID cardinality.",
    "Integrate real M2M / API-gateway logs if Fastweb/Vodafone are able to provide "
    "them - the one open item outside our control.",
    "Record the final video demonstration from the live demo script in Part 5.",
    "Expand the number of agent and attack variants to stress-test generalisation "
    "beyond the current 5-agent / 7-threat catalogue.",
]
ny, nh = 2.0, 1.05
for i, item in enumerate(next_items):
    y = ny + i * (nh + 0.12)
    rect(sl, 0.83, y, 11.6, nh, fill=PANEL_C, radius=0.08)
    rect(sl, 0.83, y, 0.09, nh, fill=ACCENT2)
    txt(sl, 1.15, y + 0.12, 10.9, nh - 0.24, [[(item, 12, BODY_C, False)]],
        anchor=MSO_ANCHOR.MIDDLE, line_spacing=1.2)

# ------------------------------------------------------------------ reorder: new slides go right after the index
sldIdLst = prs.slides._sldIdLst
all_ids = list(sldIdLst)
n_new = 11
new_ids = all_ids[-n_new:]
for el in new_ids:
    sldIdLst.remove(el)
for i, el in enumerate(new_ids):
    sldIdLst.insert(2 + i, el)

# ============================================================ presenter notes (30-min talk, 4 speakers)
NOTES = [
# 1 cover
"FAUZAN - 0:00-0:30\n"
"Good morning, we're Team 4. Gate 1 established the case for AEGIS - today, Gate 2, "
"we show the concrete architecture, a working implementation, and real numbers.",

# 2 index
"FAUZAN - 0:30-1:00\n"
"Six parts today, matching the official Gate 2 ask exactly: the proposed solution, "
"the architecture, how we implemented it, a full numerical performance analysis, "
"what the proof of concept demonstrates, and where we go next toward Gate 3.",

# 3 Part 1 - proposed solution
"FAUZAN - 1:00-4:00\n"
"One line to remember us by: credentials prove what you have, AEGIS verifies how "
"you behave. A stolen token still passes normal authentication - the credential is "
"valid, the actor isn't. AEGIS answers the question authentication never asks. Our "
"pipeline in one line: observe, fingerprint, score, decide. And here are the seven "
"threats we catch by behaviour, not by credential checks - impersonation, a "
"compromised agent, reconnaissance, volumetric abuse, slow exfiltration, sequence "
"anomalies, and scope creep. All seven are built and, as we'll show, all seven are "
"detected.",

# 4 Part 2 - architecture diagram
"ADITHYA - 4:00-7:00\n"
"Here's the architecture end to end. Observe parses every request into a normalised "
"event. Fingerprint turns a sliding window of that agent's activity into a "
"behavioural feature vector - volume, surface, method mix, sequence, errors, "
"payload, time of day, and a rolling cross-window breadth signal that catches "
"threats too sparse for any single window. Score combines deterministic rules with "
"an unsupervised ML model. Decide turns that risk number into a gate action.",

# 5 Part 2 cont - scoring
"ADITHYA - 7:00-9:30\n"
"Two layers, fused. Rules are fast and explainable - a scope violation is flagged "
"immediately, a session-state machine tracks PDU-session lifecycle. The ML layer, "
"an Isolation Forest, is trained only on legitimate traffic, so it needs no "
"labelled attacks to learn from. Fusion takes the maximum of the two, so the rules "
"layer has veto power on hard violations while ML catches subtler drift. The fused "
"result maps to pass, step-up, or block.",

# 6 Part 3 - tools & phases
"ADITHYA - 9:30-12:00\n"
"On implementation: we built this bottom-up and evaluation-driven, the ML "
"equivalent of test-driven development - no change is kept unless it's re-scored "
"against the full labelled dataset first. Everything is open-source Python: "
"pandas and scikit-learn for the pipeline, Streamlit for the live dashboard. Seven "
"build phases, from the synthetic environment through to evaluation and iteration.",

# 7 Part 3 cont - bugs found
"GHAZANFAR - 12:00-14:30\n"
"This evaluation loop is not just a formality - it caught two real bugs. First, an "
"early sequence-detection heuristic was false-firing on legitimate bursty traffic; "
"we replaced it with a proper session-state machine. Second, and more subtly, our "
"first attempt at two of the seven threats accidentally had the attacker reach "
"outside its own scope, which is a different threat entirely - so we were "
"measuring the wrong detector. Fixing that is the direct origin of the realistic, "
"non-uniform numbers you're about to see.",

# 8 Part 4 - benchmarking
"GHAZANFAR - 14:30-17:30\n"
"Before finalising our numbers, we researched how real detection systems perform. "
"Published intrusion-detection studies show volumetric attacks caught at "
"97 to 99 percent, but rare classes far lower. Industry SOC data puts typical "
"false-positive rates between 46 and 83 percent, with even elite SOCs under 10. "
"So we tuned for a realistic false-positive budget, not the highest possible "
"number: 0.965 ROC-AUC, 90.7% recall, 1.4% false positives.",

# 9 Part 4 cont - per-threat
"GHAZANFAR - 17:30-19:30\n"
"Here's the detection rate per threat, and it's deliberately not uniform. "
"Volumetric abuse and scope creep sit at 100%, because those are loud, "
"deterministic signals. The other five span 76 to 96%, tracking exactly how "
"subtle each threat's behavioural signature actually is. Low-and-slow "
"exfiltration, at 76%, is the hardest case, matching the literature.",

# 10 Part 4 cont - ROC
"LLAGAMI - 19:30-21:30\n"
"We didn't just pick a threshold, we justified it. The statistically optimal "
"point would catch slightly more attacks, but at more than double the "
"false-positive rate. Our chosen threshold trades a small amount of recall for a "
"materially lower false-positive rate, which is the right tradeoff for a system "
"that has to run continuously without alert fatigue.",

# 11 Part 5 - demo script
"LLAGAMI - 21:30-24:00\n"
"Here's what the proof of concept actually demonstrates, live, today, not a future "
"promise. We show all five agents at a glance, trigger a threat and watch the gate "
"react in real time, open the incident inspector for a plain-language reason, "
"cycle through easy and hard cases, show the numbers side by side with the live "
"view, and make the zero-trust gap concrete: a valid credential, still blocked on "
"behaviour.",

# 12 Part 5 cont - readiness
"LLAGAMI - 24:00-26:00\n"
"Where we actually stand: the generator, the full detection pipeline, the live "
"dashboard, the numerical evaluation, and the threshold analysis are all done "
"today. The one item outside our control is real network traffic - that's a "
"direct ask to Fastweb and Vodafone, not a gap in our own work.",

# 13 Part 6 - path to Gate 3
"FAUZAN - 26:00-28:30\n"
"Looking ahead to Gate 3: harden the two hardest threats with richer features, "
"integrate real traffic if it's granted, record the final demo video, and expand "
"our agent and attack coverage to stress-test how well this generalises.",

# 14 thank you
"FAUZAN - 28:30-30:00\n"
"To close: Gate 1 was the plan, Gate 2 is the working system with real, honestly "
"benchmarked numbers behind it. We're confident in this direction and looking "
"forward to your questions. Thank you.",
]

for slide, note in zip(prs.slides, NOTES):
    slide.notes_slide.notes_text_frame.text = note

os.makedirs(OUT_DIR, exist_ok=True)
prs.save(OUT)
print(f"Saved {OUT}  ({len(prs.slides._sldIdLst)} slides)")
