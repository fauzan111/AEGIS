# -*- coding: utf-8 -*-
"""Inserts one new slide into the CURRENT AEGIS_Gate3_Official.pptx, right
after "Beyond Synthetic: A Real 5G Core" - a before/after metrics comparison
answering "how did Gate 3 improve the results". Operates on the deck as it
sits on disk right now (which already has hand edits made in PowerPoint
since the last script run - team names, a slide-11 title tweak) and touches
nothing else: no relabelling pass, no font-bump pass, no other slide edits.

Deliberately honest about causality: the real Open5GS work was validation
(same unmodified detect code correctly flagged real captured traffic), not
retraining - the Isolation Forest is still fit only on synthetic data. The
metric movement shown here came from a separate Gate-3 effort (catalogue
widening + T3/T5 rolling-signal hardening), so the two are presented as
related but distinct advances rather than implying the real traffic improved
the numbers.

Run:  .venv/Scripts/python AEGIS/src/add_gate3_results_slide.py
In/Out: AEGIS/slides/GATE-3/AEGIS_Gate3_Official.pptx  (edited in place)
"""

import os
from pptx import Presentation
from pptx.util import Pt, Inches
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR

NAVY = RGBColor(0x02, 0x0C, 0x31)
ACCENT = RGBColor(0x01, 0xAB, 0xDB)
ACCENT2 = RGBColor(0x3D, 0x67, 0xB1)
BODY_C = RGBColor(0x33, 0x41, 0x55)
MUTE_C = RGBColor(0x6B, 0x76, 0x88)
CARD_C = RGBColor(0xFF, 0xFF, 0xFF)
HAIR_C = RGBColor(0xE2, 0xE8, 0xF0)
PANEL_C = RGBColor(0xF2, 0xF7, 0xFA)
GREEN_C = RGBColor(0x0E, 0x93, 0x84)
RED_C = RGBColor(0xC0, 0x39, 0x4B)
FONT = "Montserrat"

HERE = os.path.dirname(os.path.abspath(__file__))
DECK = os.path.join(HERE, "..", "slides", "GATE-3", "AEGIS_Gate3_Official.pptx")


def find_shape_by_text(slide, snippet):
    for shp in slide.shapes:
        if shp.has_text_frame and snippet in shp.text_frame.text:
            return shp
    return None


def slide_index(prs, slide):
    for i, s in enumerate(prs.slides):
        if s is slide:
            return i
    raise ValueError("slide not found")


def move_slide(prs, old_index, new_index):
    xml_slides = prs.slides._sldIdLst
    slides = list(xml_slides)
    el = slides[old_index]
    xml_slides.remove(el)
    xml_slides.insert(new_index, el)


def rect(slide, x, y, w, h, fill=None, line=None, radius=None):
    from pptx.enum.shapes import MSO_SHAPE
    sp = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    if fill is None:
        sp.fill.background()
    else:
        sp.fill.solid(); sp.fill.fore_color.rgb = fill
    if line is None:
        sp.line.fill.background()
    else:
        sp.line.color.rgb = line; sp.line.width = Pt(1.0)
    sp.shadow.inherit = False
    if radius is not None:
        try:
            sp.adjustments[0] = radius
        except Exception:
            pass
    return sp


def txt(slide, x, y, w, h, runs, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, line_spacing=1.1):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame; tf.word_wrap = True; tf.vertical_anchor = anchor
    tf.margin_left = 0; tf.margin_right = 0; tf.margin_top = 0; tf.margin_bottom = 0
    for i, para in enumerate(runs):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align; p.line_spacing = line_spacing
        for (t, sz, col, bold) in para:
            r = p.add_run(); r.text = t
            r.font.size = Pt(sz); r.font.name = FONT
            r.font.color.rgb = col; r.font.bold = bold
    return tb


def kicker(slide, text_):
    txt(slide, 0.83, 0.02, 10.5, 0.3, [[(text_, 12, ACCENT, True)]])


def title(slide, text_):
    txt(slide, 0.83, 0.29, 11.0, 1.0, [[(text_, 28, NAVY, True)]])


def lead(slide, text_, y=1.25):
    txt(slide, 0.83, y, 11.6, 0.5, [[(text_, 13, MUTE_C, False)]])


prs = Presentation(DECK)

anchor_slide = next(s for s in prs.slides if find_shape_by_text(s, "Beyond Synthetic") is not None)
anchor_idx = slide_index(prs, anchor_slide)
kicker_text = find_shape_by_text(anchor_slide, "EXPECTED PROOF OF CONCEPT").text_frame.text

new_slide = prs.slides.add_slide(anchor_slide.slide_layout)
move_slide(prs, slide_index(prs, new_slide), anchor_idx + 1)

kicker(new_slide, kicker_text)
title(new_slide, "Gate 3 Results: What Actually Changed")
lead(new_slide, "Two separate advances, not one: real-world validation on the previous "
                "slide, and a measurable detection improvement from this slide - kept apart "
                "so neither claim overstates the other.")

# before/after metric cards
metrics = [
    ("ROC-AUC", "0.965", "0.979", True),
    ("RECALL", "90.7%", "91.6%", True),
    ("FALSE POSITIVES", "1.4%", "1.44%", False),
    ("PRECISION", "0.67", "0.64", False),
]
mw = 2.75
for i, (label, before, after, improved) in enumerate(metrics):
    x = 0.83 + i * (mw + 0.17)
    rect(new_slide, x, 2.0, mw, 1.65, fill=PANEL_C, radius=0.1)
    txt(new_slide, x, 2.13, mw, 0.3, [[(label, 10.5, MUTE_C, True)]], align=PP_ALIGN.CENTER)
    txt(new_slide, x, 2.45, mw, 0.5, [[
        (before, 17, MUTE_C, False), ("  ->  ", 14, MUTE_C, False),
        (after, 20, GREEN_C if improved else RED_C, True),
    ]], align=PP_ALIGN.CENTER)
    txt(new_slide, x, 3.0, mw, 0.5,
        [[("Gate 2 -> Gate 3" if improved else "Gate 2 -> Gate 3 (harder test set)", 9, MUTE_C, False)]],
        align=PP_ALIGN.CENTER)

# honest causal explanation
txt(new_slide, 0.83, 4.05, 11.6, 0.35, [[("WHAT DROVE THE CHANGE", 13, ACCENT2, True)]])
causes = [
    "Catalogue widened: 5 to 7 agents, 7 to 10 attack variants - including harder "
    "v2 and slow-ramp variants designed to be tougher to catch, not easier.",
    "T3 recon and T5 exfiltration hardened with a rolling per-target-ID cardinality "
    "feature, closing the two weakest spots named at Gate 2.",
    "Precision and false-positive rate moved slightly the other way - expected on a "
    "harder, more varied test set, and reported honestly rather than hidden.",
]
runs = [[("-  ", 12.5, ACCENT, True), (c, 12.5, BODY_C, False)] for c in causes]
txt(new_slide, 0.83, 4.45, 11.6, 1.8, runs, line_spacing=1.3)

txt(new_slide, 0.83, 6.3, 11.6, 0.6, [[
    ("Not from the real Open5GS capture: the detector is still trained only on "
     "synthetic data. Real traffic (previous slide) validates the same code works "
     "on real signalling - it did not change these numbers.", 11, MUTE_C, False)]],
    line_spacing=1.2)

prs.save(DECK)
print(f"Saved {DECK}  ({len(prs.slides._sldIdLst)} slides) - inserted at position {anchor_idx + 2}")
