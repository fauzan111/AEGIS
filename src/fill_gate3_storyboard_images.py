"""Fill every remaining empty shot card in AEGIS_Gate3_Storyboard.pptx with a
real or animated image: rendered motion graphics for conceptual shots, cropped
real dashboard screenshots for the two remaining dashboard-content gaps.

Live-action-only shots (badge scan, door opens, team B-roll, NOC operator,
the fast montage recap) are intentionally left untouched - they need an
actual camera and a person, not a graphic.
"""

from pptx import Presentation
from pptx.util import Emu

PPTX_PATH = "slides/GATE-3/AEGIS_Gate3_Storyboard.pptx"
ANIM_DIR = "reports/GATE-3/storyboard-animations"
SHOTS_DIR = "reports/GATE-3/storyboard-shots"

COL = [Emu(237490), Emu(3729990), Emu(7235190)]
ROW_TOP = [Emu(886241), Emu(4124741)]
VID_W, VID_H = Emu(3187700), Emu(1793831)

prs = Presentation(PPTX_PATH)


def zone(col_idx, row_idx):
    return COL[col_idx], ROW_TOP[row_idx], VID_W, VID_H


def add_video(page_idx, shot_id, col_idx, row_idx, video_file, poster_file=None):
    slide = prs.slides[page_idx]
    l, t, w, h = zone(col_idx, row_idx)
    caption = next(
        (shp.text_frame.text.split("\n")[0] for shp in slide.shapes
         if shp.shape_id == shot_id and shp.has_text_frame),
        "??",
    )
    print(f"  page{page_idx+1} shot{shot_id} ({caption!r}) -> {video_file}")
    kwargs = {}
    if poster_file:
        kwargs["poster_frame_image"] = f"{ANIM_DIR}/{poster_file}"
    slide.shapes.add_movie(f"{ANIM_DIR}/{video_file}", l, t, w, h, mime_type="video/mp4", **kwargs)


def add_picture_centered(page_idx, shot_id, col_idx, row_idx, image_path):
    from PIL import Image

    slide = prs.slides[page_idx]
    l, t, w, h = zone(col_idx, row_idx)
    im = Image.open(image_path)
    aspect = im.width / im.height
    pic_w = w
    pic_h = Emu(int(w / aspect))
    if pic_h > h:
        pic_h = h
        pic_w = Emu(int(h * aspect))
    pic_t = Emu(int(t + (h - pic_h) / 2))
    pic_l = Emu(int(l + (w - pic_w) / 2))
    caption = next(
        (shp.text_frame.text.split("\n")[0] for shp in slide.shapes
         if shp.shape_id == shot_id and shp.has_text_frame),
        "??",
    )
    print(f"  page{page_idx+1} shot{shot_id} ({caption!r}) -> {image_path} (picture)")
    slide.shapes.add_picture(image_path, pic_l, pic_t, pic_w, pic_h)


print("Inserting...")

# Page 1 (index 0)
add_video(0, 69, 1, 1, "badge-credential-valid.mp4")
add_video(0, 70, 2, 1, "logo-reveal-hook.mp4")

# Page 2 (index 1)
add_video(1, 65, 0, 0, "logo-reveal-plain.mp4")
add_video(1, 68, 0, 1, "textcard-branding.mp4")
add_video(1, 69, 1, 1, "transition-noc-bridge.mp4")

# Page 3 (index 2)
add_video(2, 67, 2, 0, "textcard-neveraskyou.mp4")
add_video(2, 69, 1, 1, "textcard-businessimpact.mp4")
add_video(2, 70, 2, 1, "transition-question-to-logo.mp4")

# Page 4 (index 3)
add_video(3, 70, 2, 1, "textcard-zerotrust-nist.mp4")

# Page 5 (index 4) - real screenshot crop, not a render
add_picture_centered(4, 69, 1, 1, f"{SHOTS_DIR}/dash_block_decision_closeup.png")

# Page 6 (index 5)
add_video(5, 66, 1, 0, "open5gs-core-topology.mp4")
add_video(5, 68, 0, 1, "real-sbi-capture-terminal.mp4")
add_video(5, 69, 1, 1, "block-real-attack-result.mp4")

# Page 7 (index 6)
add_video(6, 72, 0, 0, "textcard-explainable.mp4")
add_picture_centered(6, 73, 1, 0, f"{SHOTS_DIR}/dash_incident_reasoning_closeup.png")
add_video(6, 74, 2, 0, "badge-gdpr-hash.mp4")
add_video(6, 76, 1, 1, "badge-network-edge.mp4")
add_video(6, 77, 2, 1, "textcard-closing-philosophy.mp4")

# Page 8 (index 7)
add_video(7, 66, 1, 0, "textcard-tagline1.mp4")
add_video(7, 67, 2, 0, "textcard-tagline2.mp4")
add_video(7, 68, 0, 1, "logo-reveal-links.mp4")
add_video(7, 69, 1, 1, "textcard-branding.mp4")  # intentional bookend reuse of the branding card

prs.save(PPTX_PATH)
print("Saved.")
