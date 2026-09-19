# -*- coding: utf-8 -*-
"""
AEGIS - patches the two Gate 2 Word docs IN PLACE with the same content
updates already made to docs/GATE-2/GATE-2.md and the pptx: a named SPC
baseline head-to-head, corrected (non-"optimal") threshold language, the
Gap/Solution/State-of-the-Art framing, the 5 extra statistical-evidence
findings, and removal of the "blocked on Fastweb/Vodafone" language.

Edits both:
  slides/GATE-2/GATE-2-Tech Report.docx
  slides/GATE-2/Group 4_AEGIS_Gate2_Report.docx

Run:  .venv/Scripts/python AEGIS/src/patch_gate2_docs.py
"""

import os
from docx import Document
from docx.shared import Pt, RGBColor
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

HERE = os.path.dirname(os.path.abspath(__file__))
SLIDES_G2 = os.path.join(HERE, "..", "slides", "GATE-2")
TARGETS = [
    os.path.join(SLIDES_G2, "GATE-2-Tech Report.docx"),
    os.path.join(SLIDES_G2, "Group 4_AEGIS_Gate2_Report.docx"),
]

ACCENT = "1E5A8A"


def find_heading(doc, text, exact=False):
    for p in doc.paragraphs:
        t = p.text.strip()
        if (t == text) if exact else (text in t):
            return p
    return None


def set_cell_shading(cell, hex_color):
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), hex_color)
    cell._tc.get_or_add_tcPr().append(shd)


def insert_para_before(anchor_p, text, style=None, bold=False):
    new_p = anchor_p.insert_paragraph_before("", style=style)
    run = new_p.add_run(text)
    run.bold = bold
    return new_p


def insert_table_before(doc, anchor_p, header, rows):
    tbl = doc.add_table(rows=1 + len(rows), cols=len(header))
    if doc.tables:
        tbl.style = doc.tables[0].style  # reuse whatever table style this doc already uses
    for j, h in enumerate(header):
        cell = tbl.rows[0].cells[j]
        cell.text = h
        for run in cell.paragraphs[0].runs:
            run.font.bold = True
            run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        set_cell_shading(cell, ACCENT)
    for i, row in enumerate(rows, start=1):
        for j, val in enumerate(row):
            tbl.rows[i].cells[j].text = val
    anchor_p._p.addprevious(tbl._tbl)
    # spacer paragraph after the table, before the anchor
    anchor_p.insert_paragraph_before("")
    return tbl


def replace_text_in_doc(doc, old_snippet, new_text):
    """Handles both loose paragraphs and table-cell paragraphs."""
    changed = 0
    for p in doc.paragraphs:
        if old_snippet in p.text:
            for run in list(p.runs[1:]):
                run.text = ""
            if p.runs:
                p.runs[0].text = new_text
            else:
                p.add_run(new_text)
            changed += 1
    for t in doc.tables:
        for row in t.rows:
            for cell in row.cells:
                for p in cell.paragraphs:
                    if old_snippet in p.text:
                        for run in list(p.runs[1:]):
                            run.text = ""
                        if p.runs:
                            p.runs[0].text = new_text
                        else:
                            p.add_run(new_text)
                        changed += 1
    return changed


def pick_bullet_style(doc):
    names = {s.name for s in doc.styles}
    for candidate in ("List Bullet", "List Paragraph", "ListParagraph"):
        if candidate in names:
            return candidate
    return None  # let Word use its default rather than crash


def patch_one(path):
    doc = Document(path)
    bullet_style = pick_bullet_style(doc)
    report = []

    # ---- 1. Fastweb-blame language fixes (each doc has slightly different exact wording) ----
    n = 0
    n += replace_text_in_doc(doc, "Not yet, synthetic only, blocked on Fastweb/Vodafone data access",
                             "Not yet - the pipeline is built to take it in without rework; the next validation step")
    n += replace_text_in_doc(doc, "Not yet, synthetic only",
                             "Not yet - the next validation step, already designed in")
    report.append(f"PoC-readiness Fastweb fix: {n} occurrence(s)")

    n = 0
    n += replace_text_in_doc(
        doc,
        "Integrating real M2M / API-gateway logs if made available by Fastweb/Vodafone, the one open "
        "item outside our control.",
        "Extending the pipeline to real M2M / API-gateway logs as the next validation step - the "
        "architecture is already designed to take this in without rework.")
    n += replace_text_in_doc(
        doc,
        "Integrating real M2M / API-gateway logs if made available by Fastweb/Vodafone, as an "
        "upgrade path from the fully synthetic PoC.",
        "Extending the pipeline to real M2M / API-gateway logs as the next validation step from the "
        "fully synthetic PoC - already designed in, pursued as it becomes available.")
    report.append(f"Path-to-Gate3 Fastweb fix: {n} occurrence(s)")

    # ---- 2. Gap framing paragraph in 2.1, before the "directly extends Gate 1" paragraph ----
    anchor = None
    for p in doc.paragraphs:
        if "directly extends the Gate 1 feasibility case" in p.text:
            anchor = p
            break
    if anchor is not None:
        insert_para_before(
            anchor,
            "The gap, the solution, and the state of the art, in one place. The gap: a credential "
            "check proves you have the right token, once, at the door - it never asks whether the "
            "traffic behind that token still behaves like the agent it claims to be, so a stolen or "
            "misused identity passes every time. Our solution: a behavioural fingerprint per agent, "
            "scored continuously through the observe / fingerprint / score / decide pipeline above. "
            "Against the state of the art: a named, reimplemented Statistical Process Control "
            "baseline scores 0.944 ROC-AUC / 63.4% recall on our own data; AEGIS scores 0.965 / "
            "90.7% - a statistically significant gap (p = 0.015), not an assumed one. Full "
            "comparison in section 2.4.")
        report.append("Gap framing paragraph: inserted")
    else:
        report.append("Gap framing paragraph: ANCHOR NOT FOUND - skipped")

    # ---- 3. New content block at the end of 2.4, before "2.5 Expected Proof of Concept" ----
    anchor25 = find_heading(doc, "2.5 Expected Proof of Concept")
    if anchor25 is not None:
        h_style = None
        for p in doc.paragraphs:
            if p.text.strip() == "Interpretation":
                h_style = p.style
                break

        insert_para_before(anchor25, "State of the Art, Head-to-Head", style=h_style)
        insert_para_before(anchor25,
            "\"State of the art\" can mean anything unless you name a specific system, describe how "
            "it works, and test it on your own data. We did that: Statistical Process Control (SPC), "
            "a z-score control-chart detector - the classic baseline this kind of behavioural "
            "monitoring descends from. We implemented it ourselves and ran it on the identical "
            "train/test split AEGIS is evaluated on - a real head-to-head, not two numbers from two "
            "different papers.")
        insert_table_before(doc, anchor25,
            ["Detector", "ROC-AUC", "Recall @ matched 1.43% FPR"],
            [["SPC z-score baseline", "0.944", "63.4%"],
             ["AEGIS (rules + ML fusion)", "0.965", "90.7%"]])
        insert_para_before(anchor25,
            "The improvement is statistically significant: a paired bootstrap on the AUC difference "
            "gives p = 0.015 (95% CI on the difference: [0.004, 0.040], excluding zero). Per-threat, "
            "SPC catches loud, single-feature deviations well (T4 100%, T1 94%) but is structurally "
            "blind to threats with no single-feature signature - T2 compromise (21%), T3 recon "
            "(27%), T6 sequence (51%) - because it has no notion of scope, sequence, or cross-window "
            "behaviour. That gap is the measured case for AEGIS's fingerprinting-plus-fusion "
            "approach, not an assumed architectural preference.")

        insert_para_before(anchor25, "Why the Threshold Is Not Called \"Optimal\" Without Qualification", style=h_style)
        insert_para_before(anchor25,
            "A threshold can only rigorously be called optimal with respect to a stated objective. "
            "The Neyman-Pearson lemma defines the true optimum as the likelihood-ratio test that "
            "maximises detection subject to a false-alarm cap, which requires knowing the real "
            "class-conditional distributions - not available, in production or here. What we can "
            "check honestly, for our chosen STEP-UP threshold of 0.60:")
        for bullet in [
            "Pareto-efficiency: no single alternative threshold beats it on both recall and "
            "false-positive rate simultaneously - zero dominators found by direct search.",
            "Distance to the achievable frontier: 0.60 sits within 0.2 percentage points of FPR of "
            "the empirical Neyman-Pearson-achievable frontier for this risk score (the ROC curve's "
            "concave hull) - a small, measured gap attributable to the finite ~11,000-window test "
            "set, not a meaningful inefficiency.",
            "Cost-ratio framing: minimising Cost(t) = r*(1-recall(t)) + FPR(t) for a stated cost "
            "ratio r = C(missed attack)/C(false alarm), our operating point is consistent with a "
            "cost ratio in roughly the 0.16-0.34 range - a checkable, named claim, not an "
            "unqualified \"optimal.\"",
        ]:
            insert_para_before(anchor25, bullet, style=bullet_style)

        insert_para_before(anchor25, "Additional Robustness Evidence", style=h_style)
        insert_para_before(anchor25,
            "Full detail for each check below is in reports/GATE-2/, one file per analysis.")
        for bullet in [
            "Bootstrap confidence intervals (3,000 replicates): ROC-AUC 0.965 [0.951, 0.977], "
            "recall 90.7% [87.6%, 93.5%], FPR 1.43% [1.21%, 1.64%].",
            "Cross-validated threshold stability: 5-fold CV threshold 0.600 +/- 0.033, recall "
            "90.4% +/- 5.0% - confirms 0.60 was not cherry-picked to one test set.",
            "Temporal split robustness: a strict chronological (not random) train/test split gives "
            "ROC-AUC 0.969, recall 90.1% - held up despite far less training data.",
            "Concept-drift check: KS-tests (FDR-corrected) find 6 of 20 agent/feature pairs drift "
            "even within 5 synthetic days - motivating a periodic re-baselining cadence as an "
            "explicit Gate 3 requirement.",
            "Adversarial evasion stress test: pacing the same T4 flood volume over 24 hours instead "
            "of 1 drops detection 100% to 0% - a real, demonstrated gap, with the fix (a rolling "
            "volumetric accumulator, mirroring T3/T5) now a named Gate 3 priority.",
        ]:
            insert_para_before(anchor25, bullet, style=bullet_style)
        insert_para_before(anchor25, "")
        report.append("2.4 new content block: inserted (SoA table, threshold section, evidence bullets)")
    else:
        report.append("2.4 new content block: ANCHOR '2.5 Expected Proof of Concept' NOT FOUND - skipped")

    doc.save(path)
    return report


def main():
    for path in TARGETS:
        print("=" * 10, os.path.basename(path), "=" * 10)
        for line in patch_one(path):
            print(" -", line)
        print()


if __name__ == "__main__":
    main()
