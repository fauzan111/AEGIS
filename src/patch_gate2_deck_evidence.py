# -*- coding: utf-8 -*-
"""
AEGIS - patches AEGIS_Gate2_Official.pptx IN PLACE with the professor's Gate 2
rehearsal feedback: a new Gap -> Solution -> State-of-the-Art slide, a named
head-to-head baseline comparison slide, corrected (non-"optimal") threshold
language + chart, and 5 appendix slides for the new statistical evidence.

This edits the existing, hand-tweaked deck directly (duplicating and
re-texting existing slides to preserve all styling) rather than regenerating
from make_gate2_official_deck.py, which would overwrite manual edits made
after that script last ran.

Run:  .venv/Scripts/python AEGIS/src/patch_gate2_deck_evidence.py
"""

import copy
import os
from pptx import Presentation
from pptx.util import Emu

HERE = os.path.dirname(os.path.abspath(__file__))
SLIDES_DIR = os.path.join(HERE, "..", "slides", "GATE-2")
REPORTS_G2 = os.path.join(HERE, "..", "reports", "GATE-2")
DECK = os.path.join(SLIDES_DIR, "AEGIS_Gate2_Official.pptx")

ACCENT = "01ABDB"
INK = "334155"


# ------------------------------------------------------------------ helpers
def duplicate_slide(prs, index, skip_pictures=True):
    """Deep-copies all shapes from prs.slides[index] onto a freshly appended
    slide (same layout), optionally skipping PICTURE shapes - callers that
    want a picture on the new slide add it fresh via add_picture(), which is
    far simpler and safer than copying image parts + relationship IDs."""
    source = prs.slides[index]
    dest = prs.slides.add_slide(source.slide_layout)
    for shp in list(dest.shapes):
        shp._element.getparent().remove(shp._element)
    for shp in source.shapes:
        if skip_pictures and shp.shape_type == 13:  # MSO_SHAPE_TYPE.PICTURE
            continue
        newel = copy.deepcopy(shp._element)
        dest.shapes._spTree.append(newel)
    return dest


def move_slide(prs, old_index, new_index):
    xml_slides = prs.slides._sldIdLst
    slides = list(xml_slides)
    el = slides[old_index]
    xml_slides.remove(el)
    xml_slides.insert(new_index, el)


def find_shape_by_text(slide, snippet):
    for shp in slide.shapes:
        if shp.has_text_frame and snippet in shp.text_frame.text:
            return shp
    return None


def set_simple_text(shape, new_text):
    """Replaces text in a text box that has one paragraph / one run style
    (titles, kickers, subtitles, single stat numbers) - keeps the first
    run's formatting, drops any extra runs/paragraphs."""
    tf = shape.text_frame
    p0 = tf.paragraphs[0]
    if not p0.runs:
        p0.add_run()
    r0 = p0.runs[0]
    r0.text = new_text
    for r in list(p0.runs[1:]):
        r._r.getparent().remove(r._r)
    for p in list(tf.paragraphs[1:]):
        p._p.getparent().remove(p._p)


def set_bullets(shape, bullets):
    """Rewrites a bulleted body text box (pattern: bold accent '-  ' prefix
    run + regular dark-ink run(s) per paragraph), reusing the first original
    paragraph's pPr/run formatting as the template for every new paragraph."""
    tf = shape.text_frame
    template_p = tf.paragraphs[0]._p
    parent = template_p.getparent()
    new_ps = []
    for text in bullets:
        p_copy = copy.deepcopy(template_p)
        runs_xml = p_copy.findall(".//{*}r")
        # keep exactly 2 runs: bold accent prefix, then the body text run
        for extra in runs_xml[2:]:
            extra.getparent().remove(extra)
        prefix_r, body_r = runs_xml[0], runs_xml[1]
        body_t = body_r.find("{*}t")
        body_t.text = text
        new_ps.append(p_copy)
    for p in list(tf.paragraphs):
        p._p.getparent().remove(p._p)
    for p in new_ps:
        parent.append(p)


def set_caption(shape, new_text):
    """Single-paragraph caption text boxes at the bottom of picture slides -
    same simple single-run replace as set_simple_text."""
    set_simple_text(shape, new_text)


def replace_picture(slide, old_pic_shape, image_path):
    left, top, width, height = (old_pic_shape.left, old_pic_shape.top,
                                 old_pic_shape.width, old_pic_shape.height)
    old_pic_shape._element.getparent().remove(old_pic_shape._element)
    return slide.shapes.add_picture(image_path, left, top, width, height)


def get_picture(slide):
    for shp in slide.shapes:
        if shp.shape_type == 13:
            return shp
    return None


# ------------------------------------------------------------------ main
def main():
    prs = Presentation(DECK)

    # ---- 1. New slide: Gap -> Solution -> State of the Art -------------
    # duplicate slide index 4 (Scoring: Rules+ML, Fused - the 3-card layout)
    gap_slide = duplicate_slide(prs, 4)
    move_slide(prs, len(prs.slides) - 1, 3)  # right after "Proposed Technical Solution"

    set_simple_text(find_shape_by_text(gap_slide, "PART 2 OF 6"),
                    "PART 1 OF 6 - GATE 2: TECHNICAL SOLUTION & ARCHITECTURE")
    set_simple_text(find_shape_by_text(gap_slide, "Scoring: Rules"),
                    "The Gap, Our Solution, and the State of the Art")

    set_simple_text(find_shape_by_text(gap_slide, "RULES LAYER"), "THE GAP")
    set_simple_text(find_shape_by_text(gap_slide, "Fast, explainable"),
                    "A credential check proves you have the right token, once, "
                    "at the door. It never asks whether the traffic behind that "
                    "token still behaves like the agent it claims to be - a "
                    "stolen or misused identity passes every time.")

    set_simple_text(find_shape_by_text(gap_slide, "ML LAYER"), "OUR SOLUTION")
    set_simple_text(find_shape_by_text(gap_slide, "An Isolation Forest"),
                    "A behavioural fingerprint per agent, scored continuously: "
                    "observe, fingerprint, score, decide. Rules (fast, "
                    "explainable) fused with an Isolation Forest (trained on "
                    "legit traffic only) - gated PASS / STEP-UP / BLOCK.")

    set_simple_text(find_shape_by_text(gap_slide, "FUSION"), "VS. STATE OF THE ART")
    set_simple_text(find_shape_by_text(gap_slide, "risk = max"),
                    "Named, reimplemented, and run on our own data - not cited: "
                    "a Statistical Process Control z-score baseline scores 0.944 "
                    "ROC-AUC / 63.4% recall. AEGIS scores 0.965 / 90.7% - a "
                    "statistically significant gap (paired bootstrap, p=0.015).")

    # bottom "Operating point:" caption -> repurpose as the one-line takeaway
    set_simple_text(find_shape_by_text(gap_slide, "Operating point:"),
                    "One line: credentials prove what you have, AEGIS verifies how you behave - "
                    "and we can show, not just claim, that this beats the classic baseline.")

    # ---- 2. New slide: State of the Art, Head-to-Head (after slide 8) --
    # find current index of "Numerical Analysis: State of the Art" slide
    idx_soa = next(i for i, s in enumerate(prs.slides)
                   if find_shape_by_text(s, "State of the Art & Headline") is not None)
    idx_perthreat_template = next(i for i, s in enumerate(prs.slides)
                                   if find_shape_by_text(s, "Detection Rate Per Threat") is not None)

    baseline_slide = duplicate_slide(prs, idx_perthreat_template, skip_pictures=True)
    move_slide(prs, len(prs.slides) - 1, idx_soa + 1)

    set_simple_text(find_shape_by_text(baseline_slide, "PART 4 OF 6"),
                    "PART 4 OF 6 - NUMERICAL PERFORMANCE ANALYSIS")
    set_simple_text(find_shape_by_text(baseline_slide, "Detection Rate Per Threat"),
                    "State of the Art, Head-to-Head")
    add_picture_scaled(baseline_slide, os.path.join(REPORTS_G2, "aegis_baseline_comparison.png"),
                       left=Emu(409818), top=Emu(902335), max_w=Emu(10607040), max_h=Emu(4049961))
    set_simple_text(find_shape_by_text(baseline_slide, "T4 (volumetric)"),
                    "Statistical Process Control (SPC): the classic control-chart baseline, "
                    "implemented and run by us on the identical train/test split AEGIS uses. It "
                    "catches loud, single-feature deviations well (T4 100%, T1 94%) but is "
                    "structurally blind to threats with no single-feature signature - T2 "
                    "compromise (21%), T3 recon (27%), T6 sequence (51%). That gap is the "
                    "measured case for fingerprinting + fusion, not an architectural preference.")

    # ---- 3. Update slide 8 body bullets (SoA framing) -------------------
    soa_slide = prs.slides[idx_soa]
    set_bullets(find_shape_by_text(soa_slide, "Published NIDS studies"), [
        "We named a specific state of the art and ran it ourselves: a Statistical "
        "Process Control (SPC) z-score control-chart detector, the classic baseline "
        "this kind of behavioural monitoring descends from - not a citation from a "
        "different dataset.",
        "SPC on our own data: 0.944 ROC-AUC, 63.4% recall at a matched 1.43% FPR. "
        "AEGIS: 0.965 ROC-AUC, 90.7% recall at the same FPR - a real head-to-head, "
        "detailed on the next slide.",
        "The improvement is statistically significant, not just a bigger number: a "
        "paired bootstrap on the AUC difference gives p=0.015, 95% CI [0.004, 0.040], "
        "excluding zero.",
    ])
    set_caption(find_shape_by_text(soa_slide, "These sit comfortably"),
               "These headline numbers are AEGIS's own; the next slide shows the SPC "
               "baseline comparison that grounds them against a concrete alternative, "
               "not just literature ranges.")

    # ---- 4. Update ROC Operating-Point slide (corrected threshold language) --
    roc_slide = next(s for s in prs.slides if find_shape_by_text(s, "ROC Operating-Point") is not None)
    set_simple_text(find_shape_by_text(roc_slide, "The chosen threshold is justified"),
                    "We don't call 0.60 'optimal' without saying optimal with respect to what - "
                    "here's the statistically precise version of that claim.")
    old_pic = get_picture(roc_slide)
    if old_pic is not None:
        replace_picture(roc_slide, old_pic, os.path.join(REPORTS_G2, "aegis_bayes_risk_threshold.png"))
    set_caption(find_shape_by_text(roc_slide, "Youden's J optimum") or
               find_shape_by_text(roc_slide, "trades a little recall"),
               "0.60 is Pareto-efficient (no single alternative threshold beats it on both "
               "recall and FPR) and sits within 0.2 points of FPR of the Neyman-Pearson-"
               "achievable frontier for this score - a checkable claim, not an assertion. "
               "We minimise Cost(t)=r*(1-recall)+FPR for a stated cost ratio r, not an "
               "unlabelled 'optimal'.")

    # ---- 5. Appendix slides (after the closing "Thank you" slide) -------
    idx_thankyou = next(i for i, s in enumerate(prs.slides)
                        if find_shape_by_text(s, "Thank you") is not None)
    idx_template = next(i for i, s in enumerate(prs.slides)
                        if find_shape_by_text(s, "Detection Rate Per Threat") is not None
                        or find_shape_by_text(s, "State of the Art, Head-to-Head") is not None)

    appendix_specs = [
        ("aegis_bootstrap_analysis.png", "Appendix: Bootstrap Confidence Intervals",
         "Point estimates alone overstate precision on a ~350-window attack sample. "
         "Stratified bootstrap (3,000 replicates): ROC-AUC 0.965 [0.951, 0.977], recall "
         "90.7% [87.6%, 93.5%], FPR 1.43% [1.21%, 1.64%] - and the AEGIS-vs-SPC AUC gap "
         "is significant at p=0.015."),
        ("aegis_cv_threshold_stability.png", "Appendix: Cross-Validated Threshold Stability",
         "The 0.60 threshold was picked on the same test set final numbers are reported "
         "on - a fair concern. 5-fold CV (threshold picked on 4 folds, evaluated on the "
         "5th): threshold 0.600 +/- 0.033, recall 90.4% +/- 5.0%, FPR 1.41% +/- 0.55% - "
         "not cherry-picked to one slice of data."),
        ("aegis_temporal_split_check.png", "Appendix: Temporal Split Robustness",
         "The main pipeline splits legit windows randomly across the week; real "
         "deployment fits on the past and scores the future. Rerun with a strict "
         "chronological split (1 day training, 4 days test): ROC-AUC 0.969, recall "
         "90.1%, FPR 1.24% - held up, even with far less training data."),
        ("aegis_concept_drift_check.png", "Appendix: Concept-Drift / Stationarity Check",
         "Does an agent's 'normal' stay stable long enough for a one-time baseline to "
         "hold? KS-tests (day 1 vs. last day, FDR-corrected): 6 of 20 agent/feature "
         "pairs show real drift even within 5 synthetic days - motivating a periodic "
         "re-baselining cadence as an explicit Gate 3 requirement, not an assumption."),
        ("aegis_evasion_stress_test.png", "Appendix: Adversarial Evasion Stress Test",
         "All seven threats in the main evaluation are naive - none paced against our "
         "thresholds deliberately. Pacing the same T4 flood volume over 24 hours instead "
         "of 1 drops detection 100% to 0%. A real, demonstrated gap - the direct fix (a "
         "rolling volumetric accumulator, mirroring T3/T5) is now a named Gate 3 priority."),
    ]

    insert_at = idx_thankyou + 1
    for fname, title, caption in appendix_specs:
        new_slide = duplicate_slide(prs, idx_template, skip_pictures=True)
        move_slide(prs, len(prs.slides) - 1, insert_at)
        insert_at += 1
        set_simple_text(find_shape_by_text(new_slide, "PART 4 OF 6") or
                        find_shape_by_text(new_slide, "PART 4 of 6"),
                        "APPENDIX - EXTRA STATISTICAL EVIDENCE")
        set_simple_text(find_shape_by_text(new_slide, "Detection Rate Per Threat") or
                        find_shape_by_text(new_slide, "State of the Art, Head-to-Head"),
                        title)
        add_picture_scaled(new_slide, os.path.join(REPORTS_G2, fname),
                           left=Emu(409818), top=Emu(902335), max_w=Emu(10607040), max_h=Emu(4049961))
        cap_shape = find_shape_by_text(new_slide, "T4 (volumetric)") or \
            find_shape_by_text(new_slide, "Statistical Process Control")
        if cap_shape is not None:
            set_caption(cap_shape, caption)

    prs.save(DECK)
    print(f"Saved {DECK}")
    print(f"Total slides now: {len(prs.slides)}")


def add_picture_scaled(slide, image_path, left, top, max_w, max_h):
    """Adds a picture sized to fit within (max_w, max_h) preserving aspect
    ratio, centred in that box - avoids stretching/distorting charts that
    don't share the exact aspect ratio of the slot they're dropped into."""
    from PIL import Image
    with Image.open(image_path) as im:
        iw, ih = im.size
    scale = min(max_w / iw, max_h / ih)
    w, h = int(iw * scale), int(ih * scale)
    l = left + (max_w - w) // 2
    t = top + (max_h - h) // 2
    return slide.shapes.add_picture(image_path, l, t, width=w, height=h)


if __name__ == "__main__":
    main()
