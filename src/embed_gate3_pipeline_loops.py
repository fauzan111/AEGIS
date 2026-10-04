"""Embed the 4 animated pipeline-stage loops into the empty page-4 shot cards
of AEGIS_Gate3_Storyboard.pptx (OBSERVE / FINGERPRINT / SCORE / DECIDE)."""

from pptx import Presentation
from pptx.util import Emu

PPTX_PATH = "slides/GATE-3/AEGIS_Gate3_Storyboard.pptx"
ANIM_DIR = "reports/GATE-3/storyboard-animations"

# (shot text-box id on page 4, video file, poster file, placement)
PLACEMENTS = [
    (66, "pipeline-observe-loop.mp4", "poster-observe.jpg", Emu(3729990), Emu(886241), Emu(3187700), Emu(1793831)),
    (67, "pipeline-fingerprint-loop.mp4", "poster-fingerprint.jpg", Emu(7235190), Emu(886241), Emu(3187700), Emu(1793831)),
    (68, "pipeline-score-loop.mp4", "poster-score.jpg", Emu(237490), Emu(4124741), Emu(3187700), Emu(1793831)),
    (69, "pipeline-decide-loop.mp4", "poster-decide.jpg", Emu(3729990), Emu(4124741), Emu(3187700), Emu(1793831)),
]

prs = Presentation(PPTX_PATH)
page4 = prs.slides[3]  # 0-indexed, page 4

# sanity: confirm the captions we expect are actually there
captions = {}
for shp in page4.shapes:
    if shp.name.startswith("TextBox") and shp.has_text_frame:
        captions[shp.shape_id] = shp.text_frame.text.split("\n")[0]

for shot_id, video, poster, l, t, w, h in PLACEMENTS:
    print(f"Shot {shot_id}: {captions.get(shot_id, '??')!r} -> {video}")
    page4.shapes.add_movie(
        f"{ANIM_DIR}/{video}",
        l, t, w, h,
        poster_frame_image=f"{ANIM_DIR}/{poster}",
        mime_type="video/mp4",
    )

prs.save(PPTX_PATH)
print("Saved.")
