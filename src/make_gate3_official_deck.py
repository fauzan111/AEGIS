# -*- coding: utf-8 -*-
"""Build the official Gate 3 PowerPoint deck. Continues directly from the
Gate 2 deck (same template, same visual identity, same 6-part structure) with
the tutor's rehearsal feedback folded in:

  - Slide "Scoring: Rules + ML, Fused" trimmed and simplified.
  - The old generic tools/build-phases slide replaced with two slides
    dedicated to explaining the Isolation Forest ML layer (how it works,
    what it's trained on, how its score is calibrated) - the model that
    rehearsal feedback said needed clearer explanation.
  - The "what the evaluation loop found and fixed" slide removed.
  - A new slide on the Gate-3 real Open5GS core validation work.
  - Larger body/title font sizes throughout for readability.
  - The per-threat detection chart swapped for the clearer regenerated
    aegis_eval.png, and all headline numbers updated to the latest run
    (wider agent/attack catalogue: ROC-AUC 0.979, recall 91.6%, FPR 1.44%).
  - Closing slide reframed as remaining work before final submission
    (no "Gate 4" framing - Gate 3 is the final stage). AWS/production
    deployment is intentionally left out - not finalised yet.

Run:  .venv/Scripts/python AEGIS/src/make_gate3_official_deck.py
In :  AEGIS/slides/GATE-1/AEGIS_Gate1_Official.pptx  (template source)
Out:  AEGIS/slides/GATE-3/AEGIS_Gate3_Official.pptx
"""

import os
from pptx import Presentation
from pptx.util import Pt, Inches
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

# same brand colours as the Gate 1 / Gate 2 decks, sampled from the official template
NAVY = RGBColor(0x02, 0x0C, 0x31)
ACCENT = RGBColor(0x01, 0xAB, 0xDB)
ACCENT2 = RGBColor(0x3D, 0x67, 0xB1)
LIGHTBLUE = RGBColor(0x6F, 0xCC, 0xDD)
BODY_C = RGBColor(0x33, 0x41, 0x55)
MUTE_C = RGBColor(0x6B, 0x76, 0x88)
CARD_C = RGBColor(0xFF, 0xFF, 0xFF)
HAIR_C = RGBColor(0xE2, 0xE8, 0xF0)
PANEL_C = RGBColor(0xF2, 0xF7, 0xFA)
GREEN_C = RGBColor(0x0E, 0x93, 0x84)
AMBER_C = RGBColor(0xE0, 0xA1, 0x00)
RED_C = RGBColor(0xC0, 0x39, 0x4B)
FONT = "Montserrat"

HERE = os.path.dirname(os.path.abspath(__file__))
SLIDES = os.path.join(HERE, "..", "slides")
TEMPLATE = os.path.join(SLIDES, "GATE-1", "AEGIS_Gate1_Official.pptx")
ARCH_PNG = os.path.join(HERE, "..", "reports", "aegis_architecture_diagram.png")
EVAL_PNG = os.path.join(HERE, "..", "reports", "aegis_eval.png")
ROC_PNG = os.path.join(HERE, "..", "reports", "GATE-1", "aegis_threshold_sensitivity.png")
IFOREST_PNG = os.path.join(HERE, "..", "reports", "aegis_isolation_forest_concept.png")
OPEN5GS_PNG = os.path.join(HERE, "..", "reports", "GATE-3", "storyboard-shots", "open5gs_terminal.png")
OUT_DIR = os.path.join(SLIDES, "GATE-3")
OUT = os.path.join(OUT_DIR, "AEGIS_Gate3_Official.pptx")


# ------------------------------------------------------------------ helpers (same as Gate 1/2 scripts)
def rect(slide, x, y, w, h, fill=None, line=None, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=None):
    sp = slide.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    if fill is None:
        sp.fill.background()
    else:
        sp.fill.solid(); sp.fill.fore_color.rgb = fill
    if line is None:
        sp.line.fill.background()
    else:
        sp.line.color.rgb = line; sp.line.width = Pt(1.0)
    sp.shadow.inherit = False
    if radius is not None and shape == MSO_SHAPE.ROUNDED_RECTANGLE:
        try:
            sp.adjustments[0] = radius
        except Exception:
            pass
    return sp


def txt(slide, x, y, w, h, runs, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, line_spacing=1.1):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame; tf.word_wrap = True; tf.vertical_anchor = anchor
    tf.margin_left = 0; tf.margin_right = 0; tf.margin_top = 0; tf.margin_bottom = 0
    for i, para in enumerate(runs):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align; p.line_spacing = line_spacing
        for (t, sz, col, bold) in para:
            r = p.add_run(); r.text = t
            r.font.size = Pt(sz); r.font.name = FONT
            r.font.color.rgb = col; r.font.bold = bold
    return tb


def title(slide, text_):
    txt(slide, 0.83, 0.29, 10.5, 1.0, [[(text_, 30, NAVY, True)]])


def kicker(slide, text_):
    txt(slide, 0.83, 0.02, 10.9, 0.3, [[(text_, 12, ACCENT, True)]])


def lead(slide, text_, y=1.25):
    txt(slide, 0.83, y, 11.6, 0.5, [[(text_, 13.5, MUTE_C, False)]])


def delete_slide(prs, index):
    xml_slides = prs.slides._sldIdLst
    slides = list(xml_slides)
    rId = slides[index].rId
    prs.part.drop_rel(rId)
    xml_slides.remove(slides[index])


def by_name(slide, name):
    for shape in slide.shapes:
        if shape.name == name:
            return shape
    return None


def picture_scaled(slide, path, x, y, max_w, max_h):
    """Adds a picture capped to (max_w, max_h) inches, preserving aspect ratio."""
    from PIL import Image as PILImage
    im = PILImage.open(path)
    ar = im.size[0] / im.size[1]
    w, h = max_w, max_w / ar
    if h > max_h:
        h = max_h; w = max_h * ar
    return slide.shapes.add_picture(path, Inches(x), Inches(y), width=Inches(w), height=Inches(h)), w, h


# ------------------------------------------------------------------ load Gate 1 deck as the base
prs = Presentation(TEMPLATE)

BLANK = None
for layout in prs.slide_masters[0].slide_layouts:
    if layout.name in ("Vuota", "Blank"):
        BLANK = layout
        break


def new_slide():
    return prs.slides.add_slide(BLANK)


# ------------------------------------------------------------------ slide 1: cover (kept as-is)

# ------------------------------------------------------------------ slide 2: index, rewritten for Gate 3
idx_slide = prs.slides[1]
heading = by_name(idx_slide, "TextBox 36")
if heading is not None:
    tf = heading.text_frame
    for i, ptext in enumerate(("AEGIS ", "- Gate 3 - Team 4")):
        if i >= len(tf.paragraphs):
            p = tf.add_paragraph()
        else:
            p = tf.paragraphs[i]
        if not p.runs:
            p.add_run()
        p.runs[0].text = ptext
        for extra in p.runs[1:]:
            extra.text = ""

for shape in idx_slide.shapes:
    if shape.has_text_frame and "Fauzan" in shape.text_frame.text:
        for p in shape.text_frame.paragraphs:
            for r in p.runs:
                if "–" in r.text or "—" in r.text:
                    r.text = r.text.replace("–", ":").replace("—", ":")

for shape in list(idx_slide.shapes):
    if shape.has_text_frame and shape.name.startswith("TextBox"):
        t = shape.text_frame.text.strip()
        is_number = t.isdigit() and len(t) == 2
        is_old_topic = t in (
            "What We're Building", "The 5G Network, In Plain Terms",
            "The Agents We Modelled", "How We Built The Prototype",
            "Technologies & Synthetic Data", "The Attacks We Simulated",
            "How AEGIS Decides", "Proof It Works",
            "Project Plan: Schedule, Feasibility & Business Case",
        )
        if is_number or is_old_topic:
            shape._element.getparent().remove(shape._element)

GATE3_TOPICS = [
    "Recap: The Problem & Solution",
    "Architecture",
    "The ML Layer: Isolation Forest",
    "Gate 3: Real Open5GS Core Validation",
    "Numerical Performance Analysis",
    "Proof of Concept & Remaining Work",
]
top0, step, num_w, num_l, txt_l, txt_w = 1150000, 600000, 460000, 6130000, 6700000, 5900000
for i, topicname in enumerate(GATE3_TOPICS):
    y = top0 + i * step
    numbox = idx_slide.shapes.add_textbox(num_l, y, num_w, 520000)
    tf = numbox.text_frame; tf.word_wrap = True
    r = tf.paragraphs[0].add_run(); r.text = f"{i+1:02d}"
    r.font.size = Pt(16); r.font.bold = True; r.font.name = FONT; r.font.color.rgb = ACCENT
    txtbox = idx_slide.shapes.add_textbox(txt_l, y, txt_w, 520000)
    tf = txtbox.text_frame; tf.word_wrap = True; tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    r = tf.paragraphs[0].add_run(); r.text = topicname
    r.font.size = Pt(14.5); r.font.bold = True; r.font.name = FONT; r.font.color.rgb = NAVY

# ------------------------------------------------------------------ delete all Gate-1-specific content slides
last_index = len(prs.slides) - 1
for i in range(last_index - 1, 1, -1):
    delete_slide(prs, i)

# ============================================================ PART 1/6 - Recap: Proposed Technical Solution
sl = new_slide()
kicker(sl, "PART 1 OF 6 - GATE 3: CONTINUING FROM GATE 2")
title(sl, "Recap: The Problem & Our Solution")
lead(sl, "A stolen credential still passes authentication. AEGIS answers the question "
        "authentication never asks: is this really our agent, behaving the way it always does?")
rect(sl, 0.83, 1.85, 11.6, 0.85, fill=NAVY, radius=0.1)
txt(sl, 1.15, 2.05, 11.0, 0.5, [[
    ("“Credentials prove what you have, AEGIS verifies how you behave.”", 16, RGBColor(0xFF, 0xFF, 0xFF), True)]],
    align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
txt(sl, 0.83, 2.95, 11.6, 0.4, [[
    ("OBSERVE  ->  FINGERPRINT  ->  SCORE  ->  DECIDE   (zero-trust gate dashboard)", 14, ACCENT2, True)]],
    align=PP_ALIGN.CENTER)
threats = [
    ("T1", "Identity spoofing", "wrong endpoint mix, timing, sequence for that identity"),
    ("T2", "Compromised agent", "drift from its own baseline: new endpoints, elevated errors"),
    ("T3", "Reconnaissance", "high cardinality of endpoints/targets, many 403/404s"),
    ("T4", "Volumetric abuse", "request-rate spike far above the agent's normal envelope"),
    ("T5", "Slow exfiltration", "abnormal payload sizes, off-hours persistence"),
    ("T6", "Sequence anomaly", "broken call-sequence pattern vs the agent's normal graph"),
    ("T7", "Scope creep", "calls to NFs never in the agent's baseline scope"),
]
tw, th_, tgx, tgy = 3.75, 1.05, 0.2, 0.15
for i, (tid, name, desc) in enumerate(threats):
    r, c = divmod(i, 4)
    x = 0.83 + c * (tw + tgx); y = 3.55 + r * (th_ + tgy)
    rect(sl, x, y, tw, th_, fill=CARD_C, line=HAIR_C, radius=0.09)
    rect(sl, x, y, 0.5, th_, fill=RGBColor(0xFD, 0xE8, 0xEA))
    txt(sl, x, y, 0.5, th_, [[(tid, 14, RED_C, True)]], align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    txt(sl, x + 0.62, y + 0.08, tw - 0.75, 0.3, [[(name, 13, NAVY, True)]])
    txt(sl, x + 0.62, y + 0.44, tw - 0.75, 0.58, [[(desc, 10.5, BODY_C, False)]])

# ============================================================ PART 2/6 - Architecture (diagram)
sl = new_slide()
kicker(sl, "PART 2 OF 6 - ARCHITECTURE")
title(sl, "System Architecture")
lead(sl, "A four-stage pipeline sitting in front of the network's control-plane APIs. "
        "Every request an agent sends to a Network Function is observed by the gate.")
if os.path.exists(ARCH_PNG):
    from PIL import Image as PILImage
    im = PILImage.open(ARCH_PNG)
    ar = im.size[0] / im.size[1]
    pic_h = 5.1
    pic_w = pic_h * ar
    sl.shapes.add_picture(ARCH_PNG, Inches(0.9), Inches(1.9), height=Inches(pic_h))
    notes_x = 0.9 + pic_w + 0.4
else:
    notes_x = 0.9
notes_w = max(11.6 - (notes_x - 0.83), 3.0)
txt(sl, notes_x, 1.95, notes_w, 5.2, [
    [("OBSERVE", 13.5, ACCENT2, True)],
    [("Parse request + metadata; assemble a per-agent stream.", 11.5, BODY_C, False)],
    [("", 6, BODY_C, False)],
    [("FINGERPRINT", 13.5, ACCENT2, True)],
    [("Windowed feature vector: volume, surface, method mix, sequence, "
      "errors, payload, temporal, cross-window breadth.", 11.5, BODY_C, False)],
    [("", 6, BODY_C, False)],
    [("SCORE", 13.5, ACCENT2, True)],
    [("Rules (scope, session state, rolling accumulators) + an Isolation "
      "Forest, fused: risk = max(rule, ml).", 11.5, BODY_C, False)],
    [("", 6, BODY_C, False)],
    [("DECIDE", 13.5, ACCENT2, True)],
    [("Thresholds map risk to PASS / STEP-UP / BLOCK, automatically.", 11.5, BODY_C, False)],
], line_spacing=1.2)

# ============================================================ PART 2/6 (cont) - Scoring & Decision (TRIMMED)
sl = new_slide()
kicker(sl, "PART 2 OF 6 - ARCHITECTURE")
title(sl, "Scoring: Rules + ML, Fused")
score_cards = [
    ("RULES", "Fast, explainable. Hard caps on scope, errors, payload, rate.", ACCENT),
    ("ML (ISOLATION FOREST)", "Unsupervised, trained on legit traffic only - no labelled attacks.", ACCENT2),
    ("FUSION", "risk = max(rule, ml). Rules veto hard violations; ML catches subtler drift.", GREEN_C),
]
cw = 3.78
for i, (h, d, c) in enumerate(score_cards):
    x = 0.83 + i * (cw + 0.15)
    rect(sl, x, 1.95, cw, 2.1, fill=CARD_C, line=HAIR_C, radius=0.07)
    rect(sl, x, 1.95, cw, 0.09, fill=c)
    txt(sl, x + 0.22, 2.25, cw - 0.44, 0.5, [[(h, 14.5, NAVY, True)]])
    txt(sl, x + 0.22, 2.8, cw - 0.44, 1.15, [[(d, 12, BODY_C, False)]], line_spacing=1.25)
legend = [("PASS", GREEN_C, "Behaves like its known-good self"),
          ("STEP-UP", AMBER_C, "Challenge: re-authenticate, approve, throttle"),
          ("BLOCK", RED_C, "Quarantine the identity, alert the SOC")]
for i, (lab, c, d) in enumerate(legend):
    x = 0.83 + i * (3.85 + 0.15)
    rect(sl, x, 4.45, 3.85, 1.15, fill=PANEL_C, radius=0.1)
    rect(sl, x, 4.45, 0.09, 1.15, fill=c)
    txt(sl, x + 0.28, 4.65, 3.4, 0.4, [[(lab, 14.5, c, True)]])
    txt(sl, x + 0.28, 5.08, 3.4, 0.45, [[(d, 11.5, BODY_C, False)]], line_spacing=1.15)
txt(sl, 0.83, 5.95, 11.6, 0.6, [[
    ("Thresholds are tuned for a realistic ~1.5% false-positive budget, not maximum "
     "sensitivity - detection is intentionally not uniform across threats (Part 5).", 12.5, MUTE_C, False)]],
    line_spacing=1.2)

# ============================================================ PART 3/6 - Isolation Forest, Part 1: how it works
sl = new_slide()
kicker(sl, "PART 3 OF 6 - THE ML LAYER: ISOLATION FOREST")
title(sl, "How Isolation Forest Works")
lead(sl, "The model that does the heavy lifting in the ML layer - no labelled attacks, "
        "just a simple idea applied many times over.")
if os.path.exists(IFOREST_PNG):
    pic, pw, ph = picture_scaled(sl, IFOREST_PNG, 0.83, 1.9, 11.6, 3.95)
    next_y = 1.9 + ph + 0.2
else:
    next_y = 2.2
pts = [
    "An Isolation Forest builds many random binary-split trees over the feature "
    "space: at each node, pick a random feature and a random split value.",
    "An anomaly sits apart from the bulk of normal traffic, so a handful of random "
    "splits isolates it (short path length). A normal point is surrounded by "
    "similar points and takes many more splits to isolate on its own.",
    "Anomaly score = how short the average isolation path is across the whole "
    "forest, normalised to roughly 0..1. Short path -> high score -> looks like an attack.",
]
runs = [[("-  ", 12.5, ACCENT, True), (p, 12.5, BODY_C, False)] for p in pts]
txt(sl, 0.83, next_y, 11.6, 1.8, runs, line_spacing=1.25)

# ============================================================ PART 3/6 (cont) - Isolation Forest, Part 2: training & calibration
sl = new_slide()
kicker(sl, "PART 3 OF 6 - THE ML LAYER: ISOLATION FOREST")
title(sl, "Training, Features & Calibration")
lead(sl, "Unsupervised by design - the gate has to recognise drift from agents it has "
        "never seen attacked, not just the attacks in our synthetic catalogue.")
txt(sl, 0.83, 1.9, 5.55, 0.4, [[("10 INPUT FEATURES PER WINDOW", 13, ACCENT2, True)]])
features = [
    "n - request count", "distinct_nf - Network Functions touched",
    "distinct_ep - distinct endpoints", "err_rate - fraction of 4xx/5xx",
    "get_frac / write_frac - method mix", "mean_resp / max_resp - payload size",
    "hour - time of day", "distinct_target_roll - rolling subscriber/record breadth",
]
runs = [[("-  ", 11.5, ACCENT, True), (f, 11.5, BODY_C, False)] for f in features]
txt(sl, 0.83, 2.35, 5.55, 3.5, runs, line_spacing=1.3)

txt(sl, 6.75, 1.9, 5.7, 0.4, [[("WHY IT FITS A ZERO-TRUST GATE", 13, ACCENT, True)]])
why = [
    ("No labelled attacks needed", "Fit only on each agent's own legit traffic - an "
     "unlabelled 'normal', which is all a new agent's own history ever gives you."),
    ("Per-agent identity, one-hot", "The model learns each agent's own normal, so "
     "impersonation (T1) shows up as unlike-itself, not unlike-everyone-else."),
    ("300 trees, contamination 0.02", "scikit-learn IsolationForest(n_estimators=300, "
     "contamination=0.02); engages once an agent has >=10 requests in the trailing "
     "5 minutes."),
    ("Calibrated on held-out legit data", "Scores are mapped to 0..1 using a legit "
     "split the model never trained on - in-sample calibration would understate the "
     "tail and blow up the real-world false-positive rate."),
]
wy, wh = 2.35, 0.95
for i, (h, d) in enumerate(why):
    y = wy + i * (wh + 0.08)
    rect(sl, 6.75, y, 5.7, wh, fill=CARD_C, line=HAIR_C, radius=0.08)
    rect(sl, 6.75, y, 0.07, wh, fill=ACCENT2)
    txt(sl, 7.0, y + 0.08, 5.3, 0.3, [[(h, 11.5, NAVY, True)]])
    txt(sl, 7.0, y + 0.37, 5.3, wh - 0.4, [[(d, 9.8, BODY_C, False)]], line_spacing=1.15)

# ============================================================ PART 4/6 - Gate 3: Real Open5GS Core Validation (NEW)
sl = new_slide()
kicker(sl, "PART 4 OF 6 - GATE 3: REAL OPEN5GS CORE VALIDATION")
title(sl, "Beyond Synthetic: A Real 5G Core")
lead(sl, "Gate 2's roadmap item - validate against real M2M / API-gateway traffic, "
        "not just the synthetic generator. Done for Gate 3.")
if os.path.exists(OPEN5GS_PNG):
    pic, pw, ph = picture_scaled(sl, OPEN5GS_PNG, 0.83, 1.9, 5.7, 4.0)
    right_x = 0.83 + 5.7 + 0.35
else:
    right_x = 0.83
right_w = max(11.6 - (right_x - 0.83), 4.5)
steps = [
    ("Stood up a real core", "A full Open5GS SA deployment (14 NFs) plus a UERANSIM "
     "gNB/UE - a genuine registration and PDU session, not simulated logs."),
    ("Captured real SBI traffic", "tcpdump on the AMF's network namespace, parsed "
     "HTTP/2 request/response pairs into the same schema AEGIS's detector expects."),
    ("Schema matches, independently confirmed", "Real captured paths "
     "(/nudm-sdm/v2/..., /nsmf-pdusession/v1/sm-contexts) line up with the synthetic "
     "generator's naming convention almost exactly."),
    ("Detection code validated on real traffic", "A real agent making direct SBI "
     "calls, phase 1 narrow-scope baseline then phase 2 a fast all-NF burst, run "
     "through AEGIS's own unmodified detect functions: 2 of 2 real attack-like "
     "windows correctly flagged BLOCK."),
]
sy, sh_ = 1.9, 0.95
for i, (h, d) in enumerate(steps):
    y = sy + i * (sh_ + 0.1)
    rect(sl, right_x, y, right_w, sh_, fill=CARD_C, line=HAIR_C, radius=0.07)
    rect(sl, right_x, y, 0.07, sh_, fill=GREEN_C)
    txt(sl, right_x + 0.2, y + 0.08, right_w - 0.4, 0.3, [[(h, 12, NAVY, True)]])
    txt(sl, right_x + 0.2, y + 0.38, right_w - 0.4, sh_ - 0.42, [[(d, 10, BODY_C, False)]], line_spacing=1.15)
txt(sl, 0.83, 6.1, 11.6, 0.6, [[
    ("Honest scope: a short, single-agent capture - rules-dominant, not a replacement "
     "for the full synthetic evaluation's statistical rigor on the next slides.", 11.5, MUTE_C, False)]],
    line_spacing=1.2)

# ============================================================ PART 5/6 - Numerical Analysis: Benchmarking
sl = new_slide()
kicker(sl, "PART 5 OF 6 - NUMERICAL PERFORMANCE ANALYSIS")
title(sl, "Benchmarked Against Real Detection Systems")
lead(sl, "A PoC that reports 100% detection on everything is a red flag, not a strength. "
        "We researched how comparable systems perform in practice.")
research = [
    "Published NIDS studies show volumetric/DoS attacks detected at 97-99.8% - the "
    "easiest signal - while reconnaissance and other rare classes detect far lower.",
    "Industry SOC data (Microsoft/Omdia, SANS 2025) puts typical false-positive rates "
    "at 46-83%; even elite, well-tuned SOCs sit under 10%.",
    "UEBA research shows a hard precision/recall tradeoff: near-perfect recall models "
    "sacrifice precision to as low as 0.54. Low-and-slow exfiltration is consistently "
    "the hardest category.",
]
runs = [[("-  ", 13, ACCENT, True), (r, 13, BODY_C, False)] for r in research]
txt(sl, 0.83, 1.85, 11.6, 1.9, runs, line_spacing=1.3)
metrics = [("ROC-AUC", "0.979"), ("Recall", "91.6%"), ("False positives", "1.44%"), ("Precision", "0.64")]
for i, (lab, val) in enumerate(metrics):
    x = 0.83 + i * (2.85 + 0.15)
    rect(sl, x, 4.1, 2.85, 1.35, fill=PANEL_C, radius=0.12)
    txt(sl, x, 4.3, 2.85, 0.6, [[(val, 28, ACCENT2, True)]], align=PP_ALIGN.CENTER)
    txt(sl, x, 4.95, 2.85, 0.4, [[(lab.upper(), 11, MUTE_C, True)]], align=PP_ALIGN.CENTER)
txt(sl, 0.83, 5.7, 11.6, 0.9, [[
    ("These sit comfortably inside the literature above - and improved further since "
     "Gate 2 after widening the agent/attack-variant catalogue and hardening the two "
     "weakest threats.", 12.5, BODY_C, False)]], line_spacing=1.25)

# ============================================================ PART 5/6 (cont) - Per-threat detection
sl = new_slide()
kicker(sl, "PART 5 OF 6 - NUMERICAL PERFORMANCE ANALYSIS")
title(sl, "Detection Rate Per Threat - Deliberately Not Uniform")
if os.path.exists(EVAL_PNG):
    pic, pw, ph = picture_scaled(sl, EVAL_PNG, 0.83, 1.55, 11.6, 4.7)
    next_y = 1.55 + ph + 0.25
else:
    next_y = 1.9
txt(sl, 0.83, next_y, 11.6, 1.0, [[
    ("T1, T4 and T7 (and their harder v2 / slow-ramp variants) sit at 100% - loud, "
     "deterministic policy or rate signals. T2, T3, T5 and T6 span a genuine 72-95% "
     "range that tracks the strength of their underlying behavioural signal: "
     "diffuse, low-and-slow attacks are always the hardest to catch.", 12, BODY_C, False)]],
    line_spacing=1.25)

# ============================================================ PART 5/6 (cont) - ROC operating point
sl = new_slide()
kicker(sl, "PART 5 OF 6 - NUMERICAL PERFORMANCE ANALYSIS")
title(sl, "ROC Operating-Point Analysis")
lead(sl, "The chosen threshold is justified, not just picked - a full sensitivity sweep "
        "against the statistically optimal point.")
if os.path.exists(ROC_PNG):
    pic, pw, ph = picture_scaled(sl, ROC_PNG, 1.15, 1.95, 11.0, 4.5)
    next_y = 1.95 + ph + 0.2
else:
    next_y = 2.2
txt(sl, 0.83, next_y, 11.6, 0.8, [[
    ("Youden's J optimum (threshold ~0.01) would catch slightly more attacks at a far "
     "higher false-positive rate. Our chosen threshold (0.6) trades a little recall "
     "for 1.44% false positives - a deliberate, defensible choice.", 12, BODY_C, False)]],
    line_spacing=1.25)

# ============================================================ PART 6/6 - Expected PoC: Live demo script
sl = new_slide()
kicker(sl, "PART 6 OF 6 - PROOF OF CONCEPT & REMAINING WORK")
title(sl, "Live Demo Script")
lead(sl, "AEGIS already has a working, demonstrable baseline - this is not a forward-looking "
        "description, it's a walkthrough of what exists today.")
steps = [
    ("1. Agent overview", "All agents, live risk scores, all sitting in PASS."),
    ("2. Trigger a threat", "Select a scored window; watch the dashboard flip to BLOCK live."),
    ("3. Incident inspector", "Plain-language why: which feature tripped, and which layer caught it."),
    ("4. Cycle through difficulty", "A spoofing case, a quiet drift case, a low-and-slow case."),
    ("5. Show the numbers", "ROC-AUC, recall, false-positive rate, per-threat table, live."),
    ("6. Real core, real traffic", "The same detect code, run on a real Open5GS capture."),
]
sw, sh_ = 3.78, 1.6
for i, (h, d) in enumerate(steps):
    r, c = divmod(i, 3)
    x = 0.83 + c * (sw + 0.15); y = 1.95 + r * (sh_ + 0.15)
    rect(sl, x, y, sw, sh_, fill=CARD_C, line=HAIR_C, radius=0.07)
    rect(sl, x, y, sw, 0.07, fill=ACCENT)
    txt(sl, x + 0.18, y + 0.16, sw - 0.36, 0.4, [[(h, 13, NAVY, True)]])
    txt(sl, x + 0.18, y + 0.58, sw - 0.36, 0.95, [[(d, 11, BODY_C, False)]], line_spacing=1.2)

# ============================================================ PART 6/6 (cont) - PoC readiness
sl = new_slide()
kicker(sl, "PART 6 OF 6 - PROOF OF CONCEPT & REMAINING WORK")
title(sl, "PoC Readiness at Gate 3")
ready = [
    ("Synthetic traffic generator (7 agents, 10 attack variants)", True),
    ("Fingerprinting, rules, ML, fusion, decision policy", True),
    ("Live dashboard with incident inspector", True),
    ("Numerical evaluation (ROC-AUC, recall, precision, FPR)", True),
    ("ROC operating-point / threshold sensitivity analysis", True),
    ("Real Open5GS core: traffic captured, schema and detection validated", True),
    ("Final video demonstration, recorded and edited", False),
]
ry, rh = 1.9, 0.63
for i, (label, done) in enumerate(ready):
    y = ry + i * (rh + 0.08)
    rect(sl, 0.83, y, 11.6, rh, fill=CARD_C, line=HAIR_C, radius=0.08)
    c = GREEN_C if done else AMBER_C
    stat = "DONE" if done else "IN PROGRESS"
    rect(sl, 0.83, y, 1.7, rh, fill=c, radius=0.08)
    txt(sl, 0.83, y, 1.7, rh, [[(stat, 10.5, RGBColor(0xFF, 0xFF, 0xFF), True)]],
        align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    txt(sl, 2.75, y, 9.4, rh, [[(label, 12, NAVY, False)]], anchor=MSO_ANCHOR.MIDDLE)

# ============================================================ PART 6/6 (cont) - Remaining work
sl = new_slide()
kicker(sl, "PART 6 OF 6 - PROOF OF CONCEPT & REMAINING WORK")
title(sl, "Remaining Work Before Final Submission")
next_items = [
    "Record and edit the final video demonstration from the live demo script - "
    "storyboard and supporting animations are already built.",
    "Keep stress-testing generalisation: further agent and attack-variant coverage "
    "beyond the current 7-agent / 10-variant catalogue.",
    "Continue hardening the two remaining softer spots, T2 compromise and T6 "
    "sequence, with the same rolling-signal approach used for T3 and T5.",
]
ny, nh = 2.1, 1.2
for i, item in enumerate(next_items):
    y = ny + i * (nh + 0.15)
    rect(sl, 0.83, y, 11.6, nh, fill=PANEL_C, radius=0.08)
    rect(sl, 0.83, y, 0.09, nh, fill=ACCENT2)
    txt(sl, 1.15, y + 0.15, 10.9, nh - 0.3, [[(item, 13, BODY_C, False)]],
        anchor=MSO_ANCHOR.MIDDLE, line_spacing=1.25)

# ------------------------------------------------------------------ reorder: new slides go right after the index
sldIdLst = prs.slides._sldIdLst
all_ids = list(sldIdLst)
n_new = 12
new_ids = all_ids[-n_new:]
for el in new_ids:
    sldIdLst.remove(el)
for i, el in enumerate(new_ids):
    sldIdLst.insert(2 + i, el)

# ============================================================ presenter notes (30-min talk, 4 speakers)
NOTES = [
# 1 cover
"FAUZAN - 0:00-0:30\n"
"Good morning, we're Team 4. Gate 2 showed the working system and real numbers - "
"today, Gate 3, we bring in a real 5G core, a deeper look at the ML model, and the "
"final proof of concept.",

# 2 index
"FAUZAN - 0:30-1:00\n"
"Six parts: a quick recap of the problem and solution, the architecture, a deep "
"dive on the Isolation Forest model, our real Open5GS core validation, the full "
"numerical performance analysis, and the proof of concept with what's left to do.",

# 3 Part 1 - recap
"FAUZAN - 1:00-3:00\n"
"Quick recap for anyone new in the room: credentials prove what you have, AEGIS "
"verifies how you behave. A stolen token still passes normal authentication. Our "
"pipeline: observe, fingerprint, score, decide - against seven behavioural threats, "
"all built and all detected.",

# 4 Part 2 - architecture diagram
"ADITHYA - 3:00-5:30\n"
"The architecture end to end. Observe parses every request into a normalised "
"event. Fingerprint turns a sliding window into a behavioural feature vector. "
"Score combines deterministic rules with an unsupervised ML model. Decide turns "
"that risk number into a gate action.",

# 5 Part 2 cont - scoring
"ADITHYA - 5:30-7:00\n"
"Two layers, fused: rules are fast and explainable, the ML layer - an Isolation "
"Forest - catches subtler drift. Fusion takes the maximum of the two, so rules "
"keep veto power on hard violations.",

# 6 Part 3 - Isolation Forest mechanics
"GHAZANFAR - 7:00-10:00\n"
"This is where we go deeper than Gate 2: how does the Isolation Forest actually "
"work? It builds many random split trees. An anomaly sits apart from normal "
"traffic, so a handful of random splits isolates it - a short path. A normal "
"point takes many more splits. Short average path, high anomaly score.",

# 7 Part 3 cont - training & calibration
"GHAZANFAR - 10:00-13:00\n"
"Ten features feed the model per window: request volume, NF and endpoint "
"breadth, error rate, method mix, payload size, time of day, and a rolling "
"target-ID breadth signal. It's unsupervised, trained only on each agent's own "
"legit traffic with agent identity one-hot encoded, so impersonation shows up as "
"unlike-itself. Scores are calibrated on a held-out legit split so the real-world "
"false-positive rate isn't understated.",

# 8 Part 4 - real Open5GS validation
"LLAGAMI - 13:00-16:30\n"
"Gate 2 left one item outside our control: real traffic. For Gate 3 we built our "
"own real 5G core - Open5GS, fourteen network functions, plus a real UERANSIM UE "
"registering and establishing a PDU session. We captured the real signalling, "
"confirmed its schema matches our synthetic generator almost exactly, and ran it "
"through AEGIS's own unmodified detection code: two out of two real attack-like "
"windows were correctly blocked.",

# 9 Part 5 - benchmarking
"LLAGAMI - 16:30-19:00\n"
"Before finalising our numbers we researched how real detection systems perform. "
"Published studies show volumetric attacks caught at 97 to 99.8 percent but rare "
"classes far lower. Industry SOC data puts typical false-positive rates between "
"46 and 83 percent. Our numbers, tuned for a realistic budget: 0.979 ROC-AUC, "
"91.6% recall, 1.44% false positives - improved since Gate 2 after widening our "
"agent and attack catalogue.",

# 10 Part 5 cont - per-threat
"FAUZAN - 19:00-21:00\n"
"Detection per threat, deliberately not uniform. The loud, deterministic threats "
"sit at 100%. The subtler four span 72 to 95%, tracking exactly how hard each "
"behavioural signature is to see - low-and-slow exfiltration remains the hardest "
"case, consistent with the literature.",

# 11 Part 5 cont - ROC
"FAUZAN - 21:00-22:30\n"
"We didn't just pick a threshold, we justified it. The statistically optimal "
"point catches slightly more attacks at a far higher false-positive cost. Our "
"chosen threshold is the right tradeoff for a system that has to run continuously "
"without alert fatigue.",

# 12 Part 6 - demo script
"ADITHYA - 22:30-25:00\n"
"Here's what the proof of concept demonstrates, live: all agents at a glance, a "
"triggered threat flipping the gate to BLOCK in real time, the incident inspector "
"explaining why in plain language, easy and hard cases side by side, the numbers "
"live, and now, the same detection code running on a real core capture.",

# 13 Part 6 cont - readiness
"ADITHYA - 25:00-27:00\n"
"Where we stand: the generator, the full pipeline, the dashboard, the numerical "
"evaluation, the threshold analysis, and now the real Open5GS validation are all "
"done. The one open item is the final edited video.",

# 14 Part 6 cont - remaining work
"GHAZANFAR - 27:00-29:30\n"
"What's left before final submission: record and edit the demo video from a "
"storyboard that's already built, keep expanding agent and attack coverage to "
"stress-test generalisation, and keep hardening the two remaining softer threats "
"the same way we already improved T3 and T5.",

# 15 thank you
"FAUZAN - 29:30-30:00\n"
"To close: Gate 1 was the plan, Gate 2 the working system, and Gate 3 adds a real "
"5G core and a model we can explain end to end. Thank you - happy to take "
"questions.",
]

for slide, note in zip(prs.slides, NOTES):
    slide.notes_slide.notes_text_frame.text = note

os.makedirs(OUT_DIR, exist_ok=True)
prs.save(OUT)
print(f"Saved {OUT}  ({len(prs.slides._sldIdLst)} slides)")
